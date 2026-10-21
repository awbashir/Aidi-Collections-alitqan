"""Team performance dashboard (services/team_dashboard.py, api/team_dashboard.py; 2026-10-20)."""
from __future__ import annotations

import datetime as dt
import io
import json
import uuid

import pytest

from services import team_dashboard
from services.calling import monitor
from test_calling_helpers import (CL, add_contact, cl, make_case, new_contract, new_supervisor, new_user,  # noqa: F401
                                  pgc, run)

DAY = "2026-09-03"


@pytest.fixture(scope="module")
def td_app(cl):
    """A small app with only the dashboard routes."""
    import os

    from flask import Flask
    from sqlalchemy.pool import NullPool

    import api.team_dashboard
    from core.serialization import ORJSONProvider
    from extensions import db, jwt
    from test_calling_helpers import MAIN_DB, db_url

    app = Flask("cawp_test_team_dashboard")
    app.config.update(TESTING=True, SQLALCHEMY_DATABASE_URI=db_url(MAIN_DB, "psycopg2"),
                      SQLALCHEMY_ENGINE_OPTIONS={"poolclass": NullPool}, SQLALCHEMY_TRACK_MODIFICATIONS=False,
                      JWT_SECRET_KEY=os.environ["JWT_SECRET_KEY"])
    app.json = ORJSONProvider(app)
    db.init_app(app)
    jwt.init_app(app)
    app.register_blueprint(api.team_dashboard.bp)
    return app


def token(app, user):
    from flask_jwt_extended import create_access_token
    with app.app_context():
        return {"Authorization": f"Bearer {create_access_token(identity=str(user))}"}


def scalar(pgc, sql, *params):
    with pgc.cursor() as cur:
        cur.execute(sql, params)
        return cur.fetchone()[0]


def fresh_ext(pgc, first: str) -> str:
    while True:
        ext = f"{first}{uuid.uuid4().int % 900 + 100}"
        if not scalar(pgc, "SELECT count(*) FROM user_extensions WHERE extension = %s AND released_at IS NULL", ext):
            return ext


def own(pgc, user, ext):
    run(pgc, "INSERT INTO user_extensions (extension, user_id, assign_reason) VALUES (%s, %s, 'test')", ext, user)


def status(pgc, name, parent=None):
    sid = str(uuid.uuid4())
    run(pgc, 'INSERT INTO "CommentStatuses" (id, name, contract_id, "CommentReachability_id", is_active, parent_id, color) '
             "VALUES (%s, %s, NULL, %s, true, %s, '#123456')", sid, f"{name} {sid[:6]}", CL.reach, parent)
    return sid


def note(pgc, case, author, main, sub=None, *, hour=10, text="Called", deleted=False):
    at = f"{DAY} {hour:02d}:15:00+03"
    run(pgc, """INSERT INTO "Comments" (clint_id, user_id, "contactMethod_id", comment, comment_status_id,
                                         comment_sub_status_id, create_at, update_at, is_deleted)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
        case, author, CL.method, text, main, sub, at, at, deleted)


def world(pgc):
    """member (in sup's team) and other (outside); one case each; statuses S1 (+ sub S1a) and S2."""
    member, other = new_user(pgc, name="Member"), new_user(pgc, name="Other")
    sup = new_supervisor(pgc, users=[member])
    cid, portfolio = new_contract(pgc, "td contract")
    _, other_portfolio = new_contract(pgc, "td other contract")
    mine, theirs = make_case(pgc, member, portfolio), make_case(pgc, other, other_portfolio)
    s1, s2 = status(pgc, "TD one"), status(pgc, "TD two")
    s1a = status(pgc, "TD one-a", parent=s1)
    note(pgc, mine, member, s1, s1a, hour=9)
    note(pgc, mine, member, s1, hour=10)
    note(pgc, theirs, member, s1, s1a, hour=10)                 # a case outside the supervisor's scope
    note(pgc, mine, member, s2, hour=11)
    note(pgc, mine, member, s1, text="حالة سابقة (من النظام السابق)\nold", hour=10)      # imported: excluded
    note(pgc, mine, member, s2, hour=10, deleted=True)                                  # deleted: excluded
    note(pgc, theirs, other, s2, hour=12)
    run(pgc, 'INSERT INTO "paymentDates" (clint_id, user_id, payment_date, amount, create_at) '
             "VALUES (%s, %s, %s, 250.5, %s)", mine, member, f"{DAY} 12:00:00+03", f"{DAY} 12:00:00+03")
    e1, e2 = fresh_ext(pgc, "7"), fresh_ext(pgc, "8")
    own(pgc, member, e1)
    own(pgc, other, e2)
    at = lambda h, m=0: dt.datetime(2026, 9, 3, h, m)  # noqa: E731
    rows = [
        {"calldate": at(9), "channel": f"PJSIP/{e1}-1", "dstchannel": "SIP/ITC-1", "disposition": "ANSWERED",
         "billsec": 70, "duration": 80, "linkedid": "T1"},
        {"calldate": at(9, 30), "channel": f"PJSIP/{e1}-2", "dstchannel": "SIP/ITC-2", "disposition": "NO ANSWER",
         "billsec": 0, "duration": 20, "linkedid": "T2"},
        {"calldate": at(10), "channel": "Local/FMGL;2", "dstchannel": f"PJSIP/{e1}-3", "disposition": "ANSWERED",
         "billsec": 30, "duration": 35, "linkedid": "T3"},
        {"calldate": at(11), "channel": f"PJSIP/{e2}-4", "dstchannel": "SIP/ITC-4", "disposition": "ANSWERED",
         "billsec": 15, "duration": 20, "linkedid": "T4"},
    ]
    return {"member": member, "other": other, "sup": sup, "contract": cid, "mine": mine, "theirs": theirs,
            "s1": s1, "s1a": s1a, "s2": s2, "e1": e1, "e2": e2, "rows": rows}


def call(app, fn, *args, **kw):
    from extensions import db
    with app.app_context():
        try:
            return fn(*args, **kw)
        finally:
            db.session.remove()


def test_aggregates_notes_calls_and_promises_excluding_imported_and_deleted(td_app, pgc):
    w = world(pgc)
    args = {"from": DAY, "user_id": f"{w['member']},{w['other']}"}
    d = call(td_app, team_dashboard.dashboard, args, user_id=CL.super, rows=w["rows"])
    assert d["scope"] == "all" and d["granularity"] == "hour" and d["calls"]["available"] is True
    users = {u["user_id"]: u for u in d["users"]}
    assert set(users) == {w["member"], w["other"]}
    m = users[w["member"]]
    assert m["notes"] == 4 and m["clients"] == 2                     # imported + deleted notes not counted
    assert m["by_status"] == {w["s1"]: 3, w["s2"]: 1}
    assert m["top_status"]["id"] == w["s1"] and m["top_status"]["count"] == 3
    assert m["calls_made"] == 2 and m["answered"] == 1 and m["not_answered"] == 1 and m["inbound_answered"] == 1
    assert m["talk_seconds"] == 100 and m["answer_rate"] == 0.5 and m["extensions"] == [w["e1"]]
    assert m["promises"] == 1 and m["promises_amount"] == "250.50"
    assert users[w["other"]]["notes"] == 1 and users[w["other"]]["calls_made"] == 1
    t = d["totals"]
    assert t["notes"] == 5 and t["clients"] == 2 and t["calls_made"] == 3 and t["answered"] == 2
    assert t["promises"] == 1 and t["answer_rate"] == round(2 / 3, 4)
    donut = {s["key"]: s["count"] for s in d["statuses"]}
    assert donut == {w["s1"]: 3, w["s2"]: 2}
    assert len(d["series"]) == 24
    ten = next(b for b in d["series"] if b["bucket"] == f"{DAY} 10:00")
    assert ten["total"] == 2 and ten["by_status"] == {w["s1"]: 2}                 # Riyadh hours
    # the contract filter narrows the notes and promises, never the calls
    c = call(td_app, team_dashboard.dashboard, {**args, "contract_id": w["contract"]}, user_id=CL.super, rows=w["rows"])
    cu = {u["user_id"]: u for u in c["users"]}
    assert cu[w["member"]]["notes"] == 3 and cu[w["other"]]["notes"] == 0 and cu[w["other"]]["calls_made"] == 1
    assert c["calls"]["contract_filtered"] is False


def test_sub_status_toggle(td_app, pgc):
    w = world(pgc)
    d = call(td_app, team_dashboard.dashboard, {"from": DAY, "user_id": w["member"], "group": "sub_status"},
             user_id=CL.super, rows=w["rows"])
    m = d["users"][0]
    assert d["group"] == "sub_status"
    assert m["by_status"] == {w["s1a"]: 2, f"none:{w['s1']}": 1, f"none:{w['s2']}": 1}
    legend = {s["key"]: s for s in d["statuses"]}
    assert legend[w["s1a"]]["parent_id"] == w["s1"] and legend[w["s1a"]]["parent_name"].startswith("TD one")
    assert legend[f"none:{w['s1']}"]["no_sub"] is True and legend[f"none:{w['s1']}"]["id"] is None
    assert m["top_status"]["id"] == w["s1"]                                        # always the main status
    with pytest.raises(Exception) as bad:
        call(td_app, team_dashboard.dashboard, {"from": DAY, "group": "nope"}, user_id=CL.super, rows=[])
    assert bad.value.status == 422


def test_supervisor_sees_only_the_team_and_cases_in_scope(td_app, pgc):
    from core.api_response import ApiError
    w = world(pgc)
    d = call(td_app, team_dashboard.dashboard, {"from": DAY}, user_id=w["sup"], rows=w["rows"])
    assert d["scope"] == "team"
    users = {u["user_id"]: u for u in d["users"]}
    assert set(users) == {w["member"], w["sup"]}                                  # team + self, other excluded
    assert users[w["member"]]["notes"] == 3                                        # the note on `theirs` is out of scope
    assert users[w["member"]]["calls_made"] == 2 and users[w["sup"]]["notes"] == 0
    assert d["totals"]["calls_made"] == 2 and d["totals"]["notes"] == 3
    assert {o["id"] for o in d["options"]["users"]} == {w["member"], w["sup"]}
    with pytest.raises(ApiError) as refused:
        call(td_app, team_dashboard.dashboard, {"from": DAY, "user_id": w["other"]}, user_id=w["sup"], rows=w["rows"])
    assert refused.value.status == 403 and refused.value.code == "not_in_scope"
    one = call(td_app, team_dashboard.dashboard, {"from": DAY, "user_id": w["member"]}, user_id=w["sup"], rows=w["rows"])
    assert [u["user_id"] for u in one["users"]] == [w["member"]]


def test_calls_unknown_when_the_pbx_records_are_unavailable(td_app, pgc, monkeypatch):
    from core.api_response import ApiError
    w = world(pgc)

    def down(*a, **k):
        raise ApiError(503, "cdr_unavailable", "the PBX call records are unavailable")
    monkeypatch.setattr(monitor, "_cdr_rows", down)
    d = call(td_app, team_dashboard.dashboard, {"from": DAY, "user_id": w["member"]}, user_id=CL.super)
    assert d["calls"] == {"available": False, "error": "cdr_unavailable", "truncated": False, "contract_filtered": False}
    m = d["users"][0]
    assert m["calls_made"] is None and m["answer_rate"] is None and d["totals"]["talk_seconds"] is None
    assert m["notes"] == 4


def run_role(pgc, perms: dict) -> str:
    rid = str(uuid.uuid4())
    run(pgc, "INSERT INTO roles (id, name, permission) VALUES (%s, %s, %s)", rid, f"td role {rid[:8]}", json.dumps(perms))
    return rid


def test_permission_and_validation(td_app, pgc, monkeypatch):
    monkeypatch.setattr(monitor, "_cdr_rows", lambda lo, hi, ext=None, **k: [])
    client = td_app.test_client()
    collector = new_user(pgc, CL.collector_role)
    assert client.get("/v2/dashboards/team", headers=token(td_app, collector)).status_code == 403
    assert client.get("/v2/dashboards/team/export", headers=token(td_app, collector)).status_code == 403
    watcher = new_user(pgc, run_role(pgc, {"calls_supervise": True}))
    r = client.get("/v2/dashboards/team", query_string={"from": DAY}, headers=token(td_app, watcher))
    assert r.status_code == 200, r.data[:300]
    assert r.get_json()["data"]["scope"] == "own"
    assert [u["user_id"] for u in r.get_json()["data"]["users"]] == [watcher]
    for bad in ({"from": "2026-01-01", "to": "2026-03-01"}, {"user_id": "x"}, {"contract_id": str(uuid.uuid4())},
                {"granularity": "week"}):
        resp = client.get("/v2/dashboards/team", query_string=bad, headers=token(td_app, CL.super))
        assert resp.status_code == 422, bad


def test_export_is_scoped_xlsx_and_audited(td_app, pgc, monkeypatch):
    import openpyxl
    w = world(pgc)
    monkeypatch.setattr(monitor, "_cdr_rows", lambda lo, hi, ext=None, **k: w["rows"])
    client = td_app.test_client()
    r = client.get("/v2/dashboards/team/export", query_string={"from": DAY}, headers=token(td_app, w["sup"]))
    assert r.status_code == 200 and r.headers["Cache-Control"] == "private, no-store", r.data[:300]
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert [s.title for s in wb.worksheets] == ["أداء الموظفين", "الحالات حسب الموظف", "التحديثات حسب الوقت", "معلومات"]
    users = wb.worksheets[0]
    assert users.cell(1, 1).value == "الموظف" and users.cell(1, 11).value == "تحديثات الحالة"
    names = [users.cell(i, 1).value for i in range(2, users.max_row + 1)]
    assert names[-1] == "الإجمالي" and len(names) == 3                              # member, sup, totals
    row = next(i for i in range(2, users.max_row) if users.cell(i, 2).value == w["e1"])
    assert users.cell(row, 3).value == 2 and users.cell(row, 4).value == 1 and users.cell(row, 11).value == 3
    assert users.cell(users.max_row, 11).value == 3
    r = client.get("/v2/dashboards/team/export", query_string={"from": DAY, "user_id": w["other"]},
                   headers=token(td_app, w["sup"]))
    assert r.status_code == 403
    r = client.get("/v2/dashboards/team/export", query_string={"from": DAY, "user_id": w["member"], "group": "sub_status"},
                   headers=token(td_app, CL.super))
    assert r.status_code == 200
    st = openpyxl.load_workbook(io.BytesIO(r.data)).worksheets[1]
    heads = [st.cell(1, c).value for c in range(3, st.max_column + 1)]
    assert any("(بدون حالة فرعية)" in h for h in heads) and any(" / " in h for h in heads)
    n = scalar(pgc, "SELECT count(*) FROM audit_events WHERE action = 'team_dashboard.export' AND actor_user_id IN (%s, %s)",
               w["sup"], CL.super)
    assert n >= 2


# --------------------------------------------------------------------------- one user's client updates
def note_at(pgc, case, author, main, sub=None, *, at, text="Called", contact=None):
    run(pgc, """INSERT INTO "Comments" (clint_id, user_id, "contactMethod_id", comment, comment_status_id,
                                         comment_sub_status_id, create_at, update_at, "contactNumber")
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
        case, author, CL.method, text, main, sub, f"{DAY} {at}:00+03", f"{DAY} {at}:00+03", contact)


def test_user_updates_pair_each_update_with_the_users_call_to_that_client(td_app, pgc):
    member, other = new_user(pgc, name="Updater"), new_user(pgc, name="Someone else")
    _, portfolio = new_contract(pgc, "td updates contract")
    a, b, c = (make_case(pgc, member, portfolio) for _ in range(3))
    add_contact(pgc, a, "966501110001")
    add_contact(pgc, b, "966502220002")
    add_contact(pgc, c, "966503330003")
    s1, s2 = status(pgc, "TD promise"), status(pgc, "TD no answer")
    s1a = status(pgc, "TD promise-a", parent=s1)
    ext = fresh_ext(pgc, "6")
    own(pgc, member, ext)
    note_at(pgc, a, member, s1, s1a, at="10:05")                       # a's answered call at 10:00
    note_at(pgc, b, member, s2, at="11:00")                            # b's unanswered call at 10:58
    note_at(pgc, c, member, s2, at="12:00")                            # c: only a call 5 h earlier -> unknown
    note_at(pgc, a, member, s2, at="13:30", contact="0501110001")      # by the note's contacted number
    note_at(pgc, a, other, s1, at="10:06")                             # someone else's update: not listed
    at = lambda h, m=0: dt.datetime(2026, 9, 3, h, m)  # noqa: E731
    rows = [
        {"calldate": at(10), "channel": f"PJSIP/{ext}-1", "dstchannel": "SIP/ITC-1", "disposition": "ANSWERED",
         "billsec": 95, "duration": 110, "linkedid": "U1", "uniqueid": "U1", "src": ext, "dst": "0501110001"},
        {"calldate": at(10, 58), "channel": f"PJSIP/{ext}-2", "dstchannel": "SIP/ITC-2", "disposition": "NO ANSWER",
         "billsec": 0, "duration": 30, "linkedid": "U2", "uniqueid": "U2", "src": ext, "dst": "+966502220002"},
        {"calldate": at(7), "channel": f"PJSIP/{ext}-3", "dstchannel": "SIP/ITC-3", "disposition": "ANSWERED",
         "billsec": 40, "duration": 45, "linkedid": "U3", "uniqueid": "U3", "src": ext, "dst": "0503330003"},
        {"calldate": at(13, 25), "channel": f"PJSIP/{ext}-4", "dstchannel": "SIP/ITC-4", "disposition": "BUSY",
         "billsec": 0, "duration": 8, "linkedid": "U4", "uniqueid": "U4", "src": ext, "dst": "0501110001"},
    ]
    d = call(td_app, team_dashboard.user_updates, {"from": DAY, "user_id": member}, user_id=CL.super, rows=rows)
    assert d["user"]["id"] == member and d["user"]["extensions"] == [ext]
    assert (d["updates"], d["clients"], d["total"]) == (4, 3, 4) and d["calls"]["available"] is True
    by_time = {i["at"][11:16]: i for i in d["items"]}
    assert [i["at"][11:16] for i in d["items"]] == ["10:30", "09:00", "08:00", "07:05"]      # UTC, newest first
    first = by_time["07:05"]                                            # 10:05 Riyadh
    assert first["client"]["case_id"] == a and first["status"]["id"] == s1 and first["sub_status"]["id"] == s1a
    assert first["call"]["answered"] is True and first["call"]["talk_seconds"] == 95
    assert by_time["08:00"]["call"]["answered"] is False and by_time["08:00"]["call"]["result"] == "NO ANSWER"
    assert by_time["09:00"]["call"] is None                             # no call within 3 h: unknown
    assert by_time["10:30"]["call"]["result"] == "BUSY"
    page = call(td_app, team_dashboard.user_updates, {"from": DAY, "user_id": member, "page_size": "2", "page": "2"},
                user_id=CL.super, rows=rows)
    assert [i["at"][11:16] for i in page["items"]] == ["08:00", "07:05"] and page["total"] == 4


def test_user_updates_scope_and_validation(td_app, pgc, monkeypatch):
    from core.api_response import ApiError
    w = world(pgc)
    with pytest.raises(ApiError) as e:
        call(td_app, team_dashboard.user_updates, {"from": DAY}, user_id=CL.super, rows=[])
    assert e.value.status == 422
    with pytest.raises(ApiError) as e:
        call(td_app, team_dashboard.user_updates, {"from": DAY, "user_id": f"{w['member']},{w['other']}"},
             user_id=CL.super, rows=[])
    assert e.value.status == 422
    with pytest.raises(ApiError) as e:                                  # outside the supervisor's team
        call(td_app, team_dashboard.user_updates, {"from": DAY, "user_id": w["other"]}, user_id=w["sup"], rows=[])
    assert e.value.status == 403
    d = call(td_app, team_dashboard.user_updates, {"from": DAY, "user_id": w["member"]}, user_id=w["sup"], rows=[])
    assert w["theirs"] not in {i["client"]["case_id"] for i in d["items"]}       # case outside the scope
    assert d["updates"] == 3

    def down(*a, **k):
        raise ApiError(503, "cdr_unavailable", "the PBX call records are unavailable")
    monkeypatch.setattr(monitor, "_cdr_rows", down)
    d = call(td_app, team_dashboard.user_updates, {"from": DAY, "user_id": w["member"]}, user_id=CL.super)
    assert d["calls"]["available"] is False and all(i["call"] is None for i in d["items"])
    r = td_app.test_client().get("/v2/dashboards/team/user-updates", query_string={"from": DAY, "user_id": w["member"]},
                                 headers=token(td_app, new_user(pgc)))
    assert r.status_code == 403


def test_a_supervisor_linked_only_to_a_contract_sees_that_contracts_users(td_app, pgc):
    cid, portfolio = new_contract(pgc, "td sup contract only")
    other_cid, other_portfolio = new_contract(pgc, "td sup contract other")
    inside, outside = new_user(pgc, name="Inside"), new_user(pgc, name="Outside")
    run(pgc, "UPDATE users SET contract_id = %s WHERE id = %s", cid, inside)
    run(pgc, "UPDATE users SET contract_id = %s WHERE id = %s", other_cid, outside)
    sup = new_supervisor(pgc, contracts=[cid])
    s1 = status(pgc, "TD contract scope")
    note(pgc, make_case(pgc, inside, portfolio), inside, s1)
    note(pgc, make_case(pgc, outside, other_portfolio), outside, s1)
    d = call(td_app, team_dashboard.dashboard, {"from": DAY}, user_id=sup, rows=[])
    ids = {u["user_id"] for u in d["users"]}
    assert d["scope"] == "team" and inside in ids and outside not in ids
    assert call(td_app, team_dashboard.user_updates, {"from": DAY, "user_id": inside}, user_id=sup, rows=[])["updates"] == 1
    from core.api_response import ApiError
    with pytest.raises(ApiError) as e:
        call(td_app, team_dashboard.user_updates, {"from": DAY, "user_id": outside}, user_id=sup, rows=[])
    assert e.value.status == 403
