"""Portfolio upload (services/portfolio_import.py, api/portfolio_import.py, migration 0070).

Mapping + preview counts, one portfolio per upload in one transaction, failure leaves nothing,
idempotent re-commit, duplicate portfolio names, active cases skipped, archived cases re-linked
(renamed old record, lineage, history visible, old money not in the new balance), collector
assignment by username, contact de-duplication, permissions, Arabic header guesses.

The background job is run inline: `api.portfolio_import._start` is replaced by a direct call to
`services.portfolio_import.run_import` (same function the thread pool runs)."""
from __future__ import annotations

import io
import os
import uuid
from decimal import Decimal
from pathlib import Path

import pytest

from conftest import MAIN_DB, db_url, run_sql_file
from test_bulk_framework_support import bulk_storage, make_csv, make_xlsx, q, q1  # noqa: F401

SEED = Path(__file__).resolve().parent / "fixtures" / "seed_portfolio_import.sql"
MOBILE = "00000000-0000-4000-8000-000000000706"      # seed.sql 'Mobile' (capitalised)
CONTACT_METHOD = "00000000-0000-4000-8000-000000000701"
COMMENT_STATUS = "00000000-0000-4000-8000-000000000703"
COLLECTION_TYPE = "00000000-0000-4000-8000-000000000704"
COLLECTION_METHOD = "00000000-0000-4000-8000-000000000705"


class P:
    admin = "c7c7c7c7-0000-4000-8000-000000000501"
    col = "c7c7c7c7-0000-4000-8000-000000000502"
    viewer = "c7c7c7c7-0000-4000-8000-000000000503"
    col_a = "c7c7c7c7-0000-4000-8000-000000000504"
    off = "c7c7c7c7-0000-4000-8000-000000000505"
    contract = "c7c7c7c7-0000-4000-8000-000000000301"
    contract_off = "c7c7c7c7-0000-4000-8000-000000000302"
    old_portfolio = "c7c7c7c7-0000-4000-8000-000000000401"
    email = "c7c7c7c7-0000-4000-8000-000000000706"
    reason = "c7c7c7c7-0000-4000-8000-000000000801"


# --------------------------------------------------------------------------- fixtures
@pytest.fixture(scope="session")
def pi_db(databases):
    run_sql_file(MAIN_DB, SEED)
    return databases


@pytest.fixture(scope="session")
def pi_app(pi_db, bulk_storage):  # noqa: F811
    from flask import Flask
    from sqlalchemy.pool import NullPool

    from extensions import db, jwt

    app = Flask("cawp_portfolio_import_test")
    app.config.update(
        TESTING=True,
        SQLALCHEMY_DATABASE_URI=db_url(MAIN_DB, "psycopg2"),
        SQLALCHEMY_ENGINE_OPTIONS={"poolclass": NullPool},
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        JWT_SECRET_KEY=os.environ["JWT_SECRET_KEY"],
    )
    db.init_app(app)
    jwt.init_app(app)
    import api.money
    import api.portfolio_import
    app.register_blueprint(api.portfolio_import.bp)
    app.register_blueprint(api.money.bp)
    from core import lookups
    lookups.invalidate()
    return app


@pytest.fixture(autouse=True)
def inline_jobs(request, monkeypatch):
    started = []
    if request.node.get_closest_marker("no_db"):
        return started
    import api.portfolio_import as route
    from services import portfolio_import as svc

    def start(import_id, actor):
        started.append(import_id)
        svc.run_import(import_id)

    monkeypatch.setattr(route, "_start", start)
    return started


class Client:
    def __init__(self, app):
        self.app = app
        self.client = app.test_client()
        self._tokens = {}

    def headers(self, user, key=None):
        if user not in self._tokens:
            from flask_jwt_extended import create_access_token
            with self.app.app_context():
                self._tokens[user] = create_access_token(identity=user)
        h = {"Authorization": f"Bearer {self._tokens[user]}"}
        if key is not False:
            h["Idempotency-Key"] = key or str(uuid.uuid4())
        return h

    def get(self, url, user=P.admin):
        return self.client.get(url, headers=self.headers(user, key=False))

    def post(self, url, body, user=P.admin, key=None):
        return self.client.post(url, json=body, headers=self.headers(user, key))

    def upload(self, content, name="portfolio.xlsx", user=P.admin):
        return self.client.post("/v2/portfolio-imports/uploads", data={"file": (io.BytesIO(content), name)},
                                content_type="multipart/form-data", headers=self.headers(user))


@pytest.fixture
def pc(pi_app):
    return Client(pi_app)


@pytest.fixture
def db(pi_db):
    import psycopg2
    conn = psycopg2.connect(db_url(MAIN_DB))
    conn.autocommit = True
    yield conn
    conn.close()


# --------------------------------------------------------------------------- helpers
HEAD = ["name", "identityNumber", "accountNumber", "amount", "mobile_1", "user_name"]
MAP = {"fields": {"name": "name", "identity_number": "identityNumber", "account_number": "accountNumber",
                  "amount": "amount", "collector_username": "user_name"},
       "contacts": [{"column": "mobile_1", "contact_type_id": MOBILE}]}


def tag() -> str:
    return uuid.uuid4().hex[:8].upper()


def make_case(conn, account, archived=False, owner=None, amount=500, identity=None):
    cid = str(uuid.uuid4())
    q(conn, """INSERT INTO clints (id, name, "identityNumber", "accountNumber", amount, attribute, user_id,
                                   portfolios_id, is_archive)
               VALUES (%s, %s, %s, %s, %s, '{}', %s, %s, %s)""",
      (cid, "Old " + account, identity or ("9" + cid[:9].replace("-", "0")), account, amount, owner,
       P.old_portfolio, archived))
    if archived:
        aid = str(uuid.uuid4())
        q(conn, """INSERT INTO "clintArchive" (id, clint_id, user_id, archive_reason_id) VALUES (%s, %s, %s, %s)""",
          (aid, cid, P.admin, P.reason))
        q(conn, """UPDATE clints SET "clintArchive_id" = %s WHERE id = %s""", (aid, cid))
    return cid


def uploaded(pc, rows, name="portfolio.xlsx"):
    content = make_csv(rows) if name.endswith(".csv") else make_xlsx(rows)
    r = pc.upload(content, name)
    assert r.status_code == 201, r.get_json()
    return r.get_json()["data"]


def body(up, name, mapping=None, contract=P.contract):
    return {"upload_id": up["id"], "contract_id": contract, "portfolio_name": name, "mapping": mapping or MAP}


def preview(pc, b, user=P.admin, expect=200):
    r = pc.post("/v2/portfolio-imports/preview", b, user)
    assert r.status_code == expect, r.get_json()
    return r.get_json()["data"] if expect == 200 else r.get_json()


def commit(pc, b, user=P.admin, key=None, expect=202):
    r = pc.post("/v2/portfolio-imports/commit", b, user, key=key)
    assert r.status_code == expect, r.get_json()
    return r


def status(pc, import_id, user=P.admin):
    r = pc.get(f"/v2/portfolio-imports/{import_id}", user)
    assert r.status_code == 200, r.get_json()
    return r.get_json()["data"]


def portfolios_named(conn, name, contract=P.contract):
    return q(conn, "SELECT id::text FROM portfolios WHERE contract_id = %s AND lower(btrim(name)) = lower(btrim(%s))",
             (contract, name))


# --------------------------------------------------------------------------- unit
@pytest.mark.no_db
def test_amount_parsing_accepts_arabic_digits_and_separators():
    from services.portfolio_import import parse_amount
    assert parse_amount("١٬٢٣٤٫٥٠") == (Decimal("1234.50"), None)
    assert parse_amount("1,234.5") == (Decimal("1234.50"), None)
    assert parse_amount(" 0 ") == (Decimal("0.00"), None)
    assert parse_amount("1234.5600000001") == (Decimal("1234.56"), None)
    assert parse_amount("-5")[1] == "negative_amount"
    assert parse_amount("(5)")[1] == "negative_amount"
    assert parse_amount("12.345")[1] == "amount_too_precise"
    assert parse_amount("abc")[1] == "invalid_amount"
    assert parse_amount("1.2.3")[1] == "invalid_amount"


@pytest.mark.no_db
def test_header_guess_arabic_and_legacy_english():
    from services.portfolio_import import guess_mapping
    types = [{"id": "t-mobile", "name": "Mobile"}, {"id": "t-phone", "name": "phone"},
             {"id": "t-email", "name": "email"}]
    g = guess_mapping(["الاسم", "رقم الهوية", "رقم الحساب", "المبلغ", "اسم المستخدم", "جوال 1", "الهاتف",
                       "البريد الإلكتروني", "ملاحظات"], types)
    assert g["fields"] == {"name": "الاسم", "identity_number": "رقم الهوية", "account_number": "رقم الحساب",
                           "amount": "المبلغ", "collector_username": "اسم المستخدم"}
    assert {(c["column"], c["contact_type_id"]) for c in g["contacts"]} == {
        ("جوال 1", "t-mobile"), ("الهاتف", "t-phone"), ("البريد الإلكتروني", "t-email")}
    g = guess_mapping(["name", "identityNumber", "accountNumber", "amount", "mobile_1", "mobile_2", "phone_1",
                       "email_1", "user_name", "city"], types)
    assert g["fields"] == {"name": "name", "identity_number": "identityNumber", "account_number": "accountNumber",
                           "amount": "amount", "collector_username": "user_name"}
    assert [c["column"] for c in g["contacts"]] == ["mobile_1", "mobile_2", "phone_1", "email_1"]
    assert g["contacts"][0]["contact_type_id"] == "t-mobile"     # 'Mobile' matched case-insensitively
    g = guess_mapping(["الإسم", "رقم الهويه", "رقم الحساب", "المديونية", "username"], types)
    assert set(g["fields"]) == {"name", "identity_number", "account_number", "amount", "collector_username"}


# --------------------------------------------------------------------------- upload / mapping
def test_upload_inspects_blank_headers_and_guesses(pc, db):
    t = tag()
    up = uploaded(pc, [["الاسم", "رقم الهوية", None, "رقم الحساب", "المبلغ", "جوال"],
                       ["عميل", "1000", "note in blank column", f"AR-{t}", "١٠٠", "0555"]])
    sheet = up["sheets"][0]
    assert sheet["headers"] == ["الاسم", "رقم الهوية", "Column C", "رقم الحساب", "المبلغ", "جوال"]
    assert sheet["sample"][0][3] == f"AR-{t}"                  # nothing shifted left
    assert up["purpose"] == "portfolio_import" and up["guess_sheet"] == "Sheet1"
    assert up["guess"]["fields"] == {"name": "الاسم", "identity_number": "رقم الهوية",
                                     "account_number": "رقم الحساب", "amount": "المبلغ"}
    assert [c["column"] for c in up["guess"]["contacts"]] == ["جوال"]
    mobile_types = {r[0] for r in q(db, """SELECT id::text FROM "ClintContactTypes"
                                           WHERE lower(name) IN ('mobile', 'جوال')""")}
    assert up["guess"]["contacts"][0]["contact_type_id"] in mobile_types
    r = pc.get(f"/v2/portfolio-imports/fields?upload_id={up['id']}")
    assert r.status_code == 200
    data = r.get_json()["data"]
    assert [f["key"] for f in data["fields"]] == ["name", "identity_number", "account_number", "amount",
                                                  "collector_username"]
    assert data["guess"] == up["guess"] and "Column C" in data["headers"]
    assert any(ct["id"] == P.email for ct in data["contact_types"])


def test_mapping_validation(pc):
    up = uploaded(pc, [HEAD, ["n", "1", f"V-{tag()}", "10", "", ""]])
    bad = {"fields": {"name": "name", "identity_number": "identityNumber", "account_number": "accountNumber"},
           "contacts": []}
    r = pc.post("/v2/portfolio-imports/preview", body(up, "x", bad))
    assert r.status_code == 422 and r.get_json()["code"] == "invalid_mapping"
    assert r.get_json()["error"]["mapping"] == {"amount": "required"}
    bad = {"fields": {**MAP["fields"], "amount": "nope", "name": "identityNumber"},
           "contacts": [{"column": "mobile_1", "contact_type_id": str(uuid.uuid4())}]}
    err = pc.post("/v2/portfolio-imports/preview", body(up, "x", bad)).get_json()["error"]["mapping"]
    assert err["amount"] == "unknown_column" and err["identity_number"] == "column_used_twice"
    assert err["contacts[0]"] == "unknown_contact_type"
    r = pc.post("/v2/portfolio-imports/preview", body(up, "x", contract=P.contract_off))
    assert r.status_code == 422 and r.get_json()["error"] == {"contract_id": "inactive"}


# --------------------------------------------------------------------------- preview
def test_preview_counts_without_writes(pc, db):
    t = tag()
    make_case(db, f"ACT-{t}")
    make_case(db, f"ARC-{t}", archived=True, identity="5550001")
    rows = [HEAD,
            ["New One", "111", f"N1-{t}", "1000", "0500000001", ""],
            ["New Two", "112", f"N2-{t}", "١٬٢٣٤٫٥٠", "", " PI.COLLECTOR.A "],
            ["Active", "113", f"ACT-{t}", "300", "", ""],
            ["Archived", "5550001", f"ARC-{t}", "700", "", "nobody.here"],
            [None, "114", f"N3-{t}", "10", "", ""],
            ["Bad amount", "115", f"N4-{t}", "abc", "", ""],
            ["Dup", "116", f"N1-{t}", "5", "", ""],
            ["Inactive col", "117", f"N5-{t}", "5", "", "pi.off"]]
    up = uploaded(pc, rows)
    before = (q1(db, "SELECT count(*) FROM portfolios")[0], q1(db, "SELECT count(*) FROM clints")[0])
    data = preview(pc, body(up, f"Preview {t}"))
    c = data["counts"]
    assert (c["total"], c["created"], c["created_from_archive"], c["skipped_active_exists"], c["rejected"]) \
        == (8, 3, 1, 1, 3)
    assert c["to_create"] == 4 and c["assigned"] == 1 and c["collector_not_matched"] == 2
    assert c["warnings"] == {"unknown_collector": 1, "inactive_collector": 1}
    assert c["contacts_new"] == 1
    by_row = {p["row_number"]: p for p in data["problems"]}
    assert by_row[4]["outcome"] == "skipped_active_exists"
    assert by_row[5]["outcome"] == "created_from_archive"
    assert by_row[5]["warnings"] == [{"code": "unknown_collector", "value": "nobody.here"}]
    assert by_row[6]["code"] == "missing_name" and by_row[7]["code"] == "invalid_amount"
    assert by_row[8]["code"] == "duplicate_in_file" and "2" in by_row[8]["detail"]
    assert by_row[9]["warnings"][0]["code"] == "inactive_collector"
    assert 2 not in by_row and 3 not in by_row                    # clean new rows are not problems
    assert data["can_commit"] is True and data["blocking"] == []
    after = (q1(db, "SELECT count(*) FROM portfolios")[0], q1(db, "SELECT count(*) FROM clints")[0])
    assert after == before
    assert q1(db, "SELECT count(*) FROM portfolio_imports WHERE upload_id = %s", (up["id"],))[0] == 0


# --------------------------------------------------------------------------- commit
def test_commit_creates_one_portfolio_with_all_rows_in_one_go(pc, db, inline_jobs):
    t = tag()
    rows = [HEAD + [None, "city"],
            ["New One", "211", f"C1-{t}", "1000", "0500000101", "pi.collector.a", "extra", "Riyadh"],
            ["New Two", "212", f"C2-{t}", "١٬٢٣٤٫٥٠", "", "", None, "Jeddah"],
            ["Unknown", "213", f"C3-{t}", "10", "", "ghost", None, None],
            ["Bad", "214", f"C4-{t}", "-1", "", "", None, None]]
    up = uploaded(pc, rows)
    name = f"Import {t}"
    r = commit(pc, body(up, name))
    imp = r.get_json()["data"]
    assert imp["status"] == "queued" and inline_jobs == [imp["id"]]
    done = status(pc, imp["id"])
    assert done["status"] == "succeeded", done
    assert done["counts"]["created"] == 3 and done["counts"]["rejected"] == 1
    pids = portfolios_named(db, name)
    assert len(pids) == 1 and done["portfolio_id"] == pids[0][0] and done["portfolio"]["name"] == name
    cases = {r[0]: r for r in q(db, """SELECT "accountNumber", name, "identityNumber", amount, user_id::text,
                                              attribute, id::text, is_archive
                                         FROM clints WHERE portfolios_id = %s""", (pids[0][0],))}
    assert set(cases) == {f"C1-{t}", f"C2-{t}", f"C3-{t}"}
    assert cases[f"C2-{t}"][3] == 1234.5
    assert cases[f"C1-{t}"][4] == P.col_a and cases[f"C3-{t}"][4] is None and cases[f"C2-{t}"][4] is None
    attr = cases[f"C1-{t}"][5]
    assert attr["city"] == "Riyadh" and attr["Column G"] == "extra" and attr["accountNumber"] == f"C1-{t}"
    # assignment written like every assign path: clints.user_id + one "SepcifyUsers" row
    assert q(db, 'SELECT user_id::text, create_by::text FROM "SepcifyUsers" WHERE clint_id = %s',
             (cases[f"C1-{t}"][6],)) == [(P.col_a, P.admin)]
    assert q1(db, 'SELECT count(*) FROM "SepcifyUsers" WHERE clint_id = %s', (cases[f"C3-{t}"][6],))[0] == 0
    assert q(db, """SELECT contact, contact_status FROM "ClintContacts" WHERE "identityNumber" = '211'""") \
        == [("0500000101", "new")]
    # report rows: every source row, with outcome and warnings
    report = {r[0]: r for r in q(db, """SELECT row_number, outcome, code, warnings FROM portfolio_import_rows
                                        WHERE import_id = %s""", (imp["id"],))}
    assert len(report) == 4 and report[4][1] == "created" and report[4][3][0]["code"] == "unknown_collector"
    assert report[5][1] == "rejected" and report[5][2] == "negative_amount"
    r = pc.get(f"/v2/portfolio-imports/{imp['id']}/rows?outcome=rejected")
    assert [x["row_number"] for x in r.get_json()["data"]["items"]] == [5]
    r = pc.get(f"/v2/portfolio-imports/{imp['id']}/report")
    assert r.status_code == 200 and r.mimetype.endswith("spreadsheetml.sheet")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    ws = wb.worksheets[0]
    assert ws.cell(1, 1).value == "رقم الصف" and ws.cell(1, 3).value == "النتيجة"
    assert ws.max_row == 5 and ws.cell(5, 3).value == "مرفوض"
    assert q1(db, """SELECT count(*) FROM audit_events WHERE entity_type = 'portfolio_import' AND entity_id = %s
                     AND action = 'portfolio_import.succeed'""", (imp["id"],))[0] == 1


def test_failure_midway_leaves_nothing_and_retry_uses_the_same_import(pc, db, monkeypatch):
    from services import portfolio_import as svc
    t = tag()
    arc = make_case(db, f"FA-{t}", archived=True)
    up = uploaded(pc, [HEAD, ["A", "311", f"F1-{t}", "10", "0500000301", ""],
                       ["B", "312", f"FA-{t}", "20", "", ""]])
    name = f"Fail {t}"

    def boom(*_a, **_k):
        raise RuntimeError("disk on fire")
    monkeypatch.setattr(svc, "_insert_report", boom)
    imp = commit(pc, body(up, name)).get_json()["data"]
    st = status(pc, imp["id"])
    assert st["status"] == "failed" and st["error"]["code"] == "import_failed" and st["portfolio_id"] is None
    assert portfolios_named(db, name) == []
    assert q1(db, """SELECT count(*) FROM clints WHERE "accountNumber" IN (%s, %s)""", (f"F1-{t}", f"FA-{t}-أرشيف1"))[0] == 0
    assert q1(db, 'SELECT "accountNumber", is_archive FROM clints WHERE id = %s', (arc,)) == (f"FA-{t}", True)
    assert q1(db, """SELECT count(*) FROM "ClintContacts" WHERE "identityNumber" = '311'""")[0] == 0
    assert q1(db, "SELECT count(*) FROM case_lineage WHERE previous_case_id = %s", (arc,))[0] == 0
    # retry: same import row, exactly one portfolio
    monkeypatch.undo()
    import api.portfolio_import as route
    monkeypatch.setattr(route, "_start", lambda i, a: svc.run_import(i))
    again = commit(pc, body(up, name)).get_json()["data"]
    assert again["id"] == imp["id"] and again["attempts"] == 2
    st = status(pc, imp["id"])
    assert st["status"] == "succeeded" and st["error"] is None
    assert len(portfolios_named(db, name)) == 1
    assert q1(db, "SELECT count(*) FROM portfolio_imports WHERE upload_id = %s", (up["id"],))[0] == 1


def test_recommit_of_a_succeeded_import_is_refused_or_replayed(pc, db):
    t = tag()
    up = uploaded(pc, [HEAD, ["A", "411", f"R1-{t}", "10", "", ""]])
    name = f"Once {t}"
    key = str(uuid.uuid4())
    first = commit(pc, body(up, name), key=key)
    replay = commit(pc, body(up, name), key=key)                     # double click / lost response
    assert replay.headers.get("Idempotent-Replayed") == "true" and replay.data == first.data
    r = commit(pc, body(up, name), expect=409)                       # a second, new submit
    assert r.get_json()["code"] == "already_imported"
    r = commit(pc, body(up, f"Other name {t}"), expect=409)         # even under another name
    assert r.get_json()["code"] == "already_imported"
    assert len(portfolios_named(db, name)) == 1 and portfolios_named(db, f"Other name {t}") == []
    data = preview(pc, body(up, name))
    assert data["can_commit"] is False and {b["code"] for b in data["blocking"]} >= {"already_imported"}


def test_running_import_cannot_be_started_twice(pc, db):
    t = tag()
    up = uploaded(pc, [HEAD, ["A", "421", f"RN-{t}", "10", "", ""]])
    q(db, """INSERT INTO portfolio_imports (upload_id, contract_id, portfolio_name, status, created_by)
             VALUES (%s, %s, %s, 'running', %s)""", (up["id"], P.contract, f"Run {t}", P.admin))
    r = commit(pc, body(up, f"Run {t}"), expect=409)
    assert r.get_json()["code"] == "import_in_progress"
    assert portfolios_named(db, f"Run {t}") == []


def test_duplicate_portfolio_name_is_refused(pc, db):
    t = tag()
    q(db, "INSERT INTO portfolios (name, contract_id) VALUES (%s, %s)", (f" Dup Name {t} ", P.contract))
    up = uploaded(pc, [HEAD, ["A", "511", f"D1-{t}", "10", "", ""]])
    data = preview(pc, body(up, f"dup name {t}"))
    assert data["can_commit"] is False and data["blocking"][0]["code"] == "portfolio_name_exists"
    r = commit(pc, body(up, f"Dup  Name {t}"), expect=409)
    assert r.get_json()["code"] == "portfolio_name_exists"
    assert len(portfolios_named(db, f"Dup Name {t}")) == 1
    assert q1(db, "SELECT count(*) FROM portfolio_imports WHERE upload_id = %s", (up["id"],))[0] == 0
    assert q1(db, """SELECT count(*) FROM clints WHERE "accountNumber" = %s""", (f"D1-{t}",))[0] == 0
    # the same name in ANOTHER contract is fine
    q(db, "UPDATE contracts SET is_active = true WHERE id = %s", (P.contract_off,))
    try:
        assert preview(pc, body(up, f"Dup Name {t}", contract=P.contract_off))["can_commit"] is True
    finally:
        q(db, "UPDATE contracts SET is_active = false WHERE id = %s", (P.contract_off,))


def test_active_existing_is_skipped_and_nothing_to_import_creates_no_portfolio(pc, db):
    t = tag()
    owner_case = make_case(db, f"S1-{t}", owner=P.col, amount=900)
    snap = q1(db, 'SELECT name, amount, user_id::text, portfolios_id::text, revision FROM clints WHERE id = %s',
              (owner_case,))
    up = uploaded(pc, [HEAD, ["Changed", "611", f"S1-{t}", "1", "", "pi.collector.a"]])
    data = preview(pc, body(up, f"Skip {t}"))
    assert data["counts"]["skipped_active_exists"] == 1 and data["counts"]["to_create"] == 0
    assert {b["code"] for b in data["blocking"]} == {"nothing_to_import"}
    imp = commit(pc, body(up, f"Skip {t}")).get_json()["data"]
    st = status(pc, imp["id"])
    assert st["status"] == "failed" and st["error"]["code"] == "nothing_to_import"
    assert st["counts"]["skipped_active_exists"] == 1
    assert portfolios_named(db, f"Skip {t}") == []
    assert q1(db, 'SELECT name, amount, user_id::text, portfolios_id::text, revision FROM clints WHERE id = %s',
              (owner_case,)) == snap


def test_archived_existing_becomes_a_new_record_linked_to_its_history(pc, db):
    t = tag()
    acc = f"AR-{t}"
    old = make_case(db, acc, archived=True, owner=P.col, amount=5000, identity="7770001")
    make_case(db, f"{acc}-أرشيف1")                                # suffix 1 already taken
    # history of the archived record: a callback note (opens a follow-up), payments, a promise
    q(db, """INSERT INTO "Comments" (clint_id, user_id, "contactMethod_id", comment, comment_status_id,
                                     next_contact_date)
             VALUES (%s, %s, %s, 'old promise to pay', %s, now() + interval '2 days')""",
      (old, P.col, CONTACT_METHOD, COMMENT_STATUS))
    for amount, st in ((100, True), (50, None), (25, False)):
        q(db, """INSERT INTO "Payments" (clint_id, user_id, collection_type_id, collection_method_id, amount,
                                         payment_date, approval_status)
                 VALUES (%s, %s, %s, %s, %s, now(), %s)""", (old, P.col, COLLECTION_TYPE, COLLECTION_METHOD,
                                                              amount, st))
    q(db, """INSERT INTO "paymentDates" (clint_id, user_id, payment_date, amount)
             VALUES (%s, %s, now() + interval '5 days', 300)""", (old, P.col))
    old_payments = q1(db, 'SELECT count(*) FROM "Payments" WHERE clint_id = %s', (old,))[0]

    up = uploaded(pc, [HEAD, ["Returning", "7770002", acc, "2000", "", ""]])
    data = preview(pc, body(up, f"Arch {t}"))
    assert data["counts"]["created_from_archive"] == 1
    assert data["problems"][0]["warnings"] == [{"code": "identity_changed", "previous": "7770001"}]
    imp = commit(pc, body(up, f"Arch {t}")).get_json()["data"]
    assert status(pc, imp["id"])["status"] == "succeeded"

    new = q1(db, """SELECT id::text, amount, is_archive FROM clints WHERE "accountNumber" = %s""", (acc,))
    assert new[0] != old and new[1] == 2000 and new[2] is False
    assert q1(db, 'SELECT "accountNumber", is_archive FROM clints WHERE id = %s', (old,)) == (f"{acc}-أرشيف2", True)
    assert q1(db, """SELECT previous_case_id::text, previous_account_number, renamed_to, import_id::text
                     FROM case_lineage WHERE case_id = %s""", (new[0],)) == (old, acc, f"{acc}-أرشيف2", imp["id"])
    # history stays on the old record; nothing was copied
    assert q1(db, 'SELECT count(*) FROM "Payments" WHERE clint_id = %s', (old,))[0] == old_payments == 3
    assert q1(db, 'SELECT count(*) FROM "Payments" WHERE clint_id = %s', (new[0],))[0] == 0
    assert q1(db, 'SELECT count(*) FROM "Comments" WHERE clint_id = %s', (new[0],))[0] == 0
    # the new balance ignores the previous record's money
    bal = pc.get(f"/v2/cases/{new[0]}/balance").get_json()["data"]
    assert bal["original"] == "2000.00" and bal["approved"] == "0.00" and bal["outstanding"] == "2000.00"
    # previous records: read-only history with each payment's state
    r = pc.get(f"/v2/cases/{new[0]}/previous-records")
    assert r.status_code == 200, r.get_json()
    prev = r.get_json()["data"]
    assert prev["has_previous"] is True and len(prev["records"]) == 1
    rec = prev["records"][0]
    assert rec["case_id"] == old and rec["account_number"] == f"{acc}-أرشيف2"
    assert rec["original_account_number"] == acc and rec["is_archived"] is True
    assert rec["archive"]["reason"] == "PI closed by client"
    assert rec["portfolio"]["id"] == P.old_portfolio
    assert [n["text"] for n in rec["notes"]] == ["old promise to pay"]
    assert sorted(p["status"] for p in rec["payments"]) == ["approved", "pending", "rejected"]
    assert rec["payment_counts"] == {"approved": 1, "pending": 1, "rejected": 1}
    assert [p["amount"] for p in rec["promises"]] == ["300.00"]
    assert len(rec["follow_ups"]) == 1 and rec["follow_ups"][0]["status"] == "open"
    events = {r[0]: r[1] for r in q(db, """SELECT action, case_id::text FROM audit_events
                                           WHERE source = 'portfolio_import' AND request_id = %s
                                             AND entity_type = 'case'""", (imp["id"],))}
    assert events == {"case.archive_relinked": old, "case.created_from_archive": new[0]}

    # a second re-import after the NEW record is archived links new -> newer, chain of two
    q(db, "UPDATE clints SET is_archive = true WHERE id = %s", (new[0],))
    up2 = uploaded(pc, [HEAD, ["Again", "7770002", acc, "100", "", ""]])
    imp2 = commit(pc, body(up2, f"Arch2 {t}")).get_json()["data"]
    assert status(pc, imp2["id"])["status"] == "succeeded"
    newest = q1(db, """SELECT id::text FROM clints WHERE "accountNumber" = %s""", (acc,))[0]
    assert q1(db, 'SELECT "accountNumber" FROM clints WHERE id = %s', (new[0],))[0] == f"{acc}-أرشيف3"
    chain = pc.get(f"/v2/cases/{newest}/previous-records").get_json()["data"]["records"]
    assert [(c["depth"], c["case_id"]) for c in chain] == [(1, new[0]), (2, old)]
    # an archived record that was already replaced cannot be linked again (typed suffix account)
    up3 = uploaded(pc, [HEAD, ["Typed suffix", "7770002", f"{acc}-أرشيف2", "1", "", ""]])
    data = preview(pc, body(up3, f"Arch3 {t}"))
    assert data["counts"]["rejected"] == 1 and data["problems"][0]["code"] == "archived_already_linked"


def test_previous_records_follow_the_new_case_scope(pc, db):
    t = tag()
    acc = f"SC-{t}"
    make_case(db, acc, archived=True)
    up = uploaded(pc, [HEAD, ["Scoped", "811", acc, "10", "", ""]])
    commit(pc, body(up, f"Scope {t}"))
    new = q1(db, """SELECT id::text FROM clints WHERE "accountNumber" = %s""", (acc,))[0]
    # pi.col is a collector: the new case is not theirs -> 404, no existence leak
    assert pc.get(f"/v2/cases/{new}/previous-records", P.col).status_code == 404
    q(db, "UPDATE clints SET user_id = %s WHERE id = %s", (P.col, new))
    r = pc.get(f"/v2/cases/{new}/previous-records", P.col)
    assert r.status_code == 200
    rec = r.get_json()["data"]["records"][0]
    assert rec["notes"] == [] and rec["payments"] is None and rec["promises"] is None   # no money/promise keys
    # a case without lineage
    plain = make_case(db, f"PL-{t}", owner=P.col)
    r = pc.get(f"/v2/cases/{plain}/previous-records", P.col)
    assert r.get_json()["data"] == {"case_id": plain, "has_previous": False, "records": [],
                                    "permitted": r.get_json()["data"]["permitted"]}


def test_contacts_are_deduplicated(pc, db):
    t = tag()
    q(db, """INSERT INTO "ClintContacts" (contact_type_id, "identityNumber", contact_status, contact)
             VALUES (%s, %s, 'active', '0555111222')""", (MOBILE, f"ID{t}"))
    head = ["name", "identityNumber", "accountNumber", "amount", "mobile_1", "mobile_2", "email_1"]
    up = uploaded(pc, [head, ["X", f"ID{t}", f"K1-{t}", "1", "0555111222", "0555333444", "x@example.invalid"],
                       ["Y", f"IE{t}", f"K2-{t}", "1", "0555999000", "0555999000", None]])
    mapping = {"fields": {"name": "name", "identity_number": "identityNumber", "account_number": "accountNumber",
                          "amount": "amount"},
               "contacts": [{"column": "mobile_1", "contact_type_id": MOBILE},
                            {"column": "mobile_2", "contact_type_id": MOBILE},
                            {"column": "email_1", "contact_type_id": P.email}]}
    data = preview(pc, body(up, f"Contacts {t}", mapping))
    assert data["counts"]["contacts_new"] == 3 and data["counts"]["contacts_existing"] == 2
    commit(pc, body(up, f"Contacts {t}", mapping))
    rows = sorted(q(db, """SELECT "identityNumber", contact, contact_type_id::text FROM "ClintContacts"
                           WHERE "identityNumber" IN (%s, %s)""", (f"ID{t}", f"IE{t}")))
    assert rows == sorted([(f"ID{t}", "0555111222", MOBILE), (f"ID{t}", "0555333444", MOBILE),
                           (f"ID{t}", "x@example.invalid", P.email), (f"IE{t}", "0555999000", MOBILE)])


def test_permission_gate_and_visibility(pc, db):
    content = make_xlsx([HEAD, ["A", "911", f"G1-{tag()}", "1", "", ""]])
    r = pc.upload(content, "p.xlsx", user=P.viewer)
    assert r.status_code == 403 and r.get_json()["code"] == "forbidden"
    assert pc.get("/v2/portfolio-imports/fields", P.viewer).status_code == 403
    up = uploaded(pc, [HEAD, ["A", "912", f"G2-{tag()}", "1", "", ""]])
    assert pc.post("/v2/portfolio-imports/preview", body(up, "x"), P.viewer).status_code == 403
    assert pc.post("/v2/portfolio-imports/commit", body(up, "x"), P.viewer).status_code == 403
    # another importer cannot use someone else's upload, nor see the import
    r = pc.post("/v2/portfolio-imports/preview", body(up, "x"), P.col)
    assert r.status_code == 404 and r.get_json()["code"] == "upload_not_found"
    t = tag()
    imp = commit(pc, body(up, f"Gate {t}")).get_json()["data"]
    assert pc.get(f"/v2/portfolio-imports/{imp['id']}", P.col).status_code == 404
    assert pc.get(f"/v2/portfolio-imports/{imp['id']}/report", P.col).status_code == 404
    assert pc.get(f"/v2/cases/{uuid.uuid4()}/previous-records", P.viewer).status_code == 404
    # an upload from the generic bulk store cannot be imported as a portfolio
    q(db, "UPDATE bulk_uploads SET purpose = 'lock' WHERE id = %s", (up["id"],))
    r = pc.post("/v2/portfolio-imports/preview", body(up, "x"))
    assert r.status_code == 409 and r.get_json()["code"] == "upload_purpose_mismatch"


def test_csv_upload_and_formula_rows(pc, db):
    t = tag()
    up = uploaded(pc, [HEAD, ["A", "951", f"CS1-{t}", "=1+1", "", ""], ["B", "952", f"CS2-{t}", "7", "", ""]],
                  name="portfolio.csv")
    data = preview(pc, body(up, f"Csv {t}"))
    assert data["counts"]["rejected"] == 1 and data["problems"][0]["code"] == "formula_not_allowed"
    assert data["counts"]["created"] == 1
