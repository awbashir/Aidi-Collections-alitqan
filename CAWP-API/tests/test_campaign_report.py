"""Campaign report (customer request 2026-10-21): aggregation on a seeded campaign run through the
real dispatch and wrap-up paths, the money rules (approved only; pending / rejected / reversed /
deleted / before the first call / after the window never count), scope and permission, and the
XLSX download (Arabic headers, audited).

Run: tools/dev/server_check.sh campreport api tests/test_campaign_report.py
"""
from __future__ import annotations

import datetime as dt
import io
import uuid

import pytest

from test_calling_campaign_helpers import (CL, CP, act, add_contact, cl, cp, fences, key, make_case,  # noqa: F401
                                           new_campaign, note_body, number, pgc, q1, qall, ready, run, tick,
                                           to_state, transport, world)
from test_calling_campaign_helpers import dispatched as _dispatched
from test_calling_helpers import invalidate_caches

CTYPE, CMETHOD = "00000000-0000-4000-8000-000000000704", "00000000-0000-4000-8000-000000000705"


@pytest.fixture(scope="session")
def rp_app(cp):
    import os

    from flask import Flask
    from sqlalchemy.pool import NullPool

    from conftest import MAIN_DB, db_url
    from core.serialization import ORJSONProvider
    from extensions import db, jwt

    app = Flask("cawp_test_campaign_report")
    app.config.update(
        TESTING=True,
        SQLALCHEMY_DATABASE_URI=db_url(MAIN_DB, "psycopg2"),
        SQLALCHEMY_ENGINE_OPTIONS={"poolclass": NullPool},
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        JWT_SECRET_KEY=os.environ["JWT_SECRET_KEY"],
    )
    app.json = ORJSONProvider(app)
    db.init_app(app)
    jwt.init_app(app)
    import api.calling
    import api.calling_campaign_report
    import api.calling_campaigns
    app.register_blueprint(api.calling.bp)
    app.register_blueprint(api.calling_campaigns.bp)
    app.register_blueprint(api.calling_campaign_report.bp)
    return app


@pytest.fixture
def client(rp_app):
    invalidate_caches()
    return rp_app.test_client()


@pytest.fixture
def auth(rp_app):
    from flask_jwt_extended import create_access_token

    def make(user_id, k=None):
        with rp_app.app_context():
            token = create_access_token(identity=str(user_id))
        h = {"Authorization": f"Bearer {token}"}
        if k is not None:
            h["Idempotency-Key"] = k
        return h
    return make


def riyadh_today() -> dt.date:
    from core import riyadh
    return riyadh.today()


def at(day: dt.date, hour=10) -> str:
    return f"{day.isoformat()} {hour:02d}:15:00+03"


def pay(pgc, case, author, amount, day, approved=True, deleted=False):
    pid = str(uuid.uuid4())
    run(pgc, """INSERT INTO "Payments" (id, clint_id, user_id, collection_type_id, collection_method_id, amount,
                                       payment_date, approval_status, is_deleted)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
        pid, case, author, CTYPE, CMETHOD, amount, at(day), approved, deleted)
    return pid


def statuses(pgc):
    reach = q1(pgc, 'SELECT "CommentReachability_id" FROM "CommentStatuses" WHERE id = %s', CL.status)[0]
    main, sub = str(uuid.uuid4()), str(uuid.uuid4())
    run(pgc, """INSERT INTO "CommentStatuses" (id, name, contract_id, "CommentReachability_id", is_active, parent_id)
                VALUES (%s, %s, NULL, %s, true, NULL), (%s, %s, NULL, %s, true, %s)""",
        main, f"CR main {main[:6]}", reach, sub, f"CR sub {sub[:6]}", reach, main)
    return main, sub


def queue_sent(pgc, attempt):
    run(pgc, "UPDATE call_dispatch_queue SET state = 'sending', claimed_at = now(), claimed_by = 't' "
             "WHERE attempt_id = %s AND state = 'queued'", attempt)
    run(pgc, "UPDATE call_dispatch_queue SET state = 'submitted', finished_at = now(), result_code = 'submitted', "
             "ami_response = 'Success' WHERE attempt_id = %s AND state = 'sending'", attempt)
    run(pgc, "UPDATE call_attempts SET pbx_epoch = 'e', pbx_channel_id = 'cawp-' || id::text WHERE id = %s "
             "AND pbx_channel_id IS NULL", attempt)


def talked(pgc, attempt, seconds):
    """Sent, bridged `seconds` ago, ended now: talk time = seconds exactly."""
    queue_sent(pgc, attempt)
    to_state(pgc, attempt, ["dispatched", "agent_answered", "customer_dialing", "customer_answered"])
    run(pgc, "UPDATE call_attempts SET state = 'bridged', bridged_at = clock_timestamp() - make_interval(secs => %s) "
             "WHERE id = %s", seconds, attempt)
    run(pgc, "UPDATE call_attempts SET state = 'ended', ended_at = bridged_at + make_interval(secs => %s), "
             "reason_code = 'customer_hung_up' WHERE id = %s", seconds, attempt)


def unanswered(pgc, attempt):
    queue_sent(pgc, attempt)
    to_state(pgc, attempt, ["dispatched", "agent_answered", "customer_dialing"])
    run(pgc, "UPDATE call_attempts SET state = 'ended', ended_at = now(), reason_code = 'no_answer' WHERE id = %s",
        attempt)


def wrap(client, auth, pgc, user, attempt, outcome, status, sub=None, promise=None):
    rev = q1(pgc, "SELECT revision FROM call_attempts WHERE id = %s", attempt)[0]
    note = note_body(f"report {outcome}", comment_status_id=status)
    if sub:
        note["comment_sub_status_id"] = sub
    body = {"expected_revision": rev, "outcome": outcome, "note": note}
    if promise:
        body["promise"] = promise
    r = client.post(f"/v2/calls/{attempt}/wrapup", headers=auth(user, key()), json=body)
    assert r.status_code == 200, r.get_json()


def by_user(results):
    return {g["user_id"]: g for g in _dispatched(results)}


@pytest.fixture
def seeded(rp_app, client, auth, world, pgc):
    """c1 owns cases A, B; c2 owns C, D (all callable); E (c1) has no contact (never called,
    nothing recorded: pending); F (c2) has an exclusion recorded by the dispatcher (excluded).
    Round 1: c1 reaches a case (90 s, promise 300), c2 does not (no answer). Round 2: c1 reaches
    the other (30 s, connected), c2 does not (busy). Round 3: c2 gets a call that is still queued
    when the campaign is stopped (cancelled, never sent); c1 gets none: both of c1's customers
    answered, and a customer who answered is never called again by the campaign (2026-10-24)."""
    w = world
    main, sub = statuses(pgc)
    for owner in (w.c1, w.c1, w.c2, w.c2):
        case = make_case(pgc, owner, w.portfolio)
        add_contact(pgc, case, number(4), default=True)
    pending = make_case(pgc, w.c1, w.portfolio)
    excluded = make_case(pgc, w.c2, w.portfolio)
    add_contact(pgc, excluded, number(4), default=True)
    cid = new_campaign(client, auth, w, settings={"wrapup_delay_seconds": 0, "cooldown_minutes": 0,
                                                    "max_attempts_per_day": 50})
    run(pgc, """INSERT INTO campaign_pool_state (campaign_id, case_id, excluded_reason, excluded_at, recheck_after)
                VALUES (%s, %s, 'case_locked', now(), now() + interval '1 day')""", cid, excluded)
    ready(pgc, cid, w.c1)
    ready(pgc, cid, w.c2)
    due = (riyadh_today() + dt.timedelta(days=5)).isoformat()

    r1 = by_user(tick(rp_app, cid, w))
    assert set(r1) == {w.c1, w.c2}, r1
    talked(pgc, r1[w.c1]["attempt_id"], 90)
    wrap(client, auth, pgc, w.c1, r1[w.c1]["attempt_id"], "promise_discussed", main,
         promise={"amount": "300.00", "due_date": due})
    unanswered(pgc, r1[w.c2]["attempt_id"])
    wrap(client, auth, pgc, w.c2, r1[w.c2]["attempt_id"], "no_answer", main, sub)

    r2 = by_user(tick(rp_app, cid, w))
    assert set(r2) == {w.c1, w.c2}, r2
    assert r2[w.c1]["case_id"] != r1[w.c1]["case_id"] and r2[w.c2]["case_id"] != r1[w.c2]["case_id"]
    talked(pgc, r2[w.c1]["attempt_id"], 30)
    wrap(client, auth, pgc, w.c1, r2[w.c1]["attempt_id"], "connected", main, sub)
    unanswered(pgc, r2[w.c2]["attempt_id"])
    wrap(client, auth, pgc, w.c2, r2[w.c2]["attempt_id"], "busy", CL.status)

    r3 = by_user(tick(rp_app, cid, w))
    assert set(r3) == {w.c2}, r3
    return {"cid": cid, "w": w, "main": main, "sub": sub, "due": due, "pending": pending, "excluded": excluded,
            "promise_case": r1[w.c1]["case_id"], "promise_attempt": r1[w.c1]["attempt_id"],
            "unreached_case": r1[w.c2]["case_id"]}


def get_report(client, auth, user, cid, **params):
    r = client.get(f"/v2/campaigns/{cid}/report", headers=auth(user), query_string=params)
    return r.status_code, r.get_json()


def test_report_aggregates_a_stopped_campaign(client, auth, pgc, seeded):
    s = seeded
    w, cid = s["w"], s["cid"]
    today = riyadh_today()
    p_case, u_case = s["promise_case"], s["unreached_case"]
    pay(pgc, p_case, w.c1, 100, today)                                  # counts
    pay(pgc, p_case, w.c1, 50, today, approved=None)                    # pending: shown apart, never counted
    pay(pgc, p_case, w.c1, 20, today, approved=False)                   # rejected
    rev = pay(pgc, p_case, w.c1, 40, today)                             # reversed below
    run(pgc, "INSERT INTO payment_reversals (payment_id, case_id, amount, reason, reversed_by) "
             "VALUES (%s, %s, 40, 'test', %s)", rev, p_case, w.c1)
    pay(pgc, p_case, w.c1, 10, today, deleted=True)                     # deleted
    pay(pgc, p_case, w.c1, 70, today - dt.timedelta(days=1))            # before the first campaign call
    pay(pgc, p_case, w.c1, 80, today + dt.timedelta(days=1))            # after the window (today, running end)
    pay(pgc, u_case, w.c2, 60, today)                                   # called, not reached: counts
    pay(pgc, s["pending"], w.c1, 25, today)                             # never called: not counted

    st, body = act(client, auth, w.supervisor, cid, "stop")
    assert st == 200, body
    assert sorted(x[0] for x in qall(pgc, "SELECT a.state FROM campaign_attempts ca JOIN call_attempts a "
                                          "ON a.id = ca.attempt_id WHERE ca.campaign_id = %s", cid)) == \
        ["cancelled", "ended", "ended", "ended", "ended"]

    st, body = get_report(client, auth, w.supervisor, cid)
    assert st == 200, body
    d = body["data"]
    assert d["final"] is True and d["campaign"]["status"] == "stopped"
    assert d["window"]["started_at"] and d["window"]["stopped_at"] and d["window"]["duration_seconds"] >= 0
    assert [c["id"] for c in d["pool"]["contracts"]] == [w.contract]
    assert d["clients"] == {"total": 6, "pool_now": 6, "called": 4, "reached": 2, "not_reached": 2, "excluded": 1,
                            "pending": 1, "excluded_by_reason": {"case_locked": 1}}
    a = d["attempts"]
    assert (a["total"], a["sent"], a["not_sent"], a["reached"], a["pending_wrapup"]) == (5, 4, 1, 2, 0)
    assert a["answer_rate"] == 0.5
    assert (a["talk_seconds"], a["talk_known"], a["talk_unknown"], a["avg_talk_seconds"]) == (120, 2, 0, 60)
    assert d["outcomes"] == {"connected": 1, "promise_discussed": 1, "dispute": 0, "no_answer": 1, "busy": 1,
                             "wrong_number": 0, "do_not_call": 0}
    got = {x["key"]: x["count"] for x in d["statuses"]}
    assert got == {s["main"]: 1, f"{s['main']}/{s['sub']}": 2, CL.status: 1}
    subrow = next(x for x in d["statuses"] if x["sub_status_id"] == s["sub"])
    assert subrow["status_name"].startswith("CR main") and subrow["sub_status_name"].startswith("CR sub")
    pr = d["promises"]
    assert (pr["count"], pr["amount"], pr["cancelled"], pr["first_due"], pr["last_due"]) == \
        (1, "300.00", 0, s["due"], s["due"])
    assert pr["by_due_date"] == [{"date": s["due"], "count": 1, "amount": "300.00"}]
    pm = d["payments"]
    assert pm["label"] == "collected_after_contact"
    assert (pm["count"], pm["amount"], pm["cases"], pm["reached_cases_amount"]) == (2, "160.00", 2, "100.00")
    assert pm["pending_not_counted"] == {"count": 1, "amount": "50.00"}
    assert pm["payment_days"] == 7 and pm["until"] == today.isoformat()
    ag = {x["user"]["id"]: x for x in d["agents"]}
    assert set(ag) == {w.c1, w.c2}
    c1, c2 = ag[w.c1], ag[w.c2]
    assert (c1["attempts"], c1["sent"], c1["reached"], c1["talk_seconds"], c1["avg_talk_seconds"]) == (2, 2, 2, 120, 60)
    assert (c1["promises"], c1["promises_amount"], c1["answer_rate"]) == (1, "300.00", 1.0)
    assert c1["outcomes"] == {"promise_discussed": 1, "connected": 1}
    assert c1["statuses"] == {s["main"]: 1, f"{s['main']}/{s['sub']}": 1}
    assert (c2["attempts"], c2["sent"], c2["reached"], c2["talk_seconds"], c2["avg_talk_seconds"]) == (3, 2, 0, 0, None)
    assert (c2["promises"], c2["promises_amount"], c2["answer_rate"]) == (0, "0.00", 0.0)
    assert c2["outcomes"] == {"no_answer": 1, "busy": 1}
    assert d["agents_took_part"] == 2
    assert sum(h["attempts"] for h in d["hours"]) == 4 and sum(h["reached"] for h in d["hours"]) == 2
    assert len(d["hours"]) == 24

    # a longer payment window takes in the payment dated tomorrow only once that day has come
    st, body = get_report(client, auth, w.supervisor, cid, payment_days=30)
    assert st == 200 and body["data"]["payments"]["amount"] == "160.00"
    st, body = get_report(client, auth, w.supervisor, cid, payment_days=91)
    assert st == 422 and body["error"] == {"payment_days": "invalid"}

    # auto-assignments (campaign rules, 0080) when that table exists
    if q1(pgc, "SELECT to_regclass('campaign_case_assignments')")[0] is None:
        assert d["assignments"]["available"] is False
    else:
        run(pgc, """INSERT INTO campaign_case_assignments (campaign_id, case_id, attempt_id, user_id, cause)
                    VALUES (%s, %s, %s, %s, 'outcome')""", cid, p_case, s["promise_attempt"], w.c1)
        st, body = get_report(client, auth, w.supervisor, cid)
        asg = body["data"]["assignments"]
        assert asg["available"] is True and asg["count"] == 1 and asg["by_user"][0]["user"]["id"] == w.c1


def test_report_so_far_while_running_and_scope(client, auth, pgc, world):
    w = world
    case = make_case(pgc, w.c1, w.portfolio)
    add_contact(pgc, case, number(4), default=True)
    cid = new_campaign(client, auth, w)
    st, body = get_report(client, auth, w.supervisor, cid)
    assert st == 200, body
    d = body["data"]
    assert d["final"] is False and d["campaign"]["status"] == "running" and d["window"]["stopped_at"] is None
    assert d["clients"]["total"] == d["clients"]["pending"] == 1 and d["clients"]["called"] == 0
    assert d["attempts"]["total"] == 0 and d["attempts"]["answer_rate"] is None
    assert d["attempts"]["avg_talk_seconds"] is None and d["payments"]["amount"] == "0.00"
    assert {x["user"]["id"] for x in d["agents"]} == {w.c1, w.c2}          # the roster, even with no calls
    assert get_report(client, auth, CP.super, cid)[0] == 200                # SuperAdmin sees every campaign
    st, body = get_report(client, auth, w.other_supervisor, cid)            # another team: not visible
    assert st == 404
    assert client.get(f"/v2/campaigns/{cid}/report/export", headers=auth(w.other_supervisor)).status_code == 404
    st, body = get_report(client, auth, w.c1, cid)                          # a member without campaigns_manage
    assert st == 403
    assert client.get(f"/v2/campaigns/{cid}/report/export", headers=auth(w.c1)).status_code == 403
    assert get_report(client, auth, w.supervisor, uuid.uuid4())[0] == 404


def test_report_export_headers_and_audit(client, auth, pgc, seeded):
    from openpyxl import load_workbook
    s = seeded
    w, cid = s["w"], s["cid"]
    pay(pgc, s["promise_case"], w.c1, 100, riyadh_today())
    r = client.get(f"/v2/campaigns/{cid}/report/export", headers=auth(w.supervisor))
    assert r.status_code == 200, r.get_data(as_text=True)[:500]
    assert r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "campaign-report-" in r.headers["Content-Disposition"]
    assert r.headers["Cache-Control"] == "private, no-store"
    wb = load_workbook(io.BytesIO(r.data), read_only=True)
    assert wb.sheetnames[:4] == ["ملخص الحملة", "الحالات المسجلة", "أداء الموظفين", "العملاء"]
    assert "المستبعدون" in wb.sheetnames and "الوعود" in wb.sheetnames and "حسب الساعة" in wb.sheetnames
    summary = {row[0]: row[1] for row in wb["ملخص الحملة"].iter_rows(min_row=2, values_only=True)}
    assert summary["نوع التقرير"].startswith("حتى الآن")                    # still running
    assert summary["تم الاتصال بهم"] == 4 and summary["تم الوصول إليهم (محادثة)"] == 2
    assert summary["المحصّل بعد التواصل (معتمد فقط)"] == 100
    assert summary["مبلغ الوعود"] == 300
    clients = list(wb["العملاء"].iter_rows(values_only=True))
    assert clients[0] == ("رقم الحساب", "اسم العميل", "الأرقام التي جُرّبت", "المحاولات", "أُرسلت", "تم الوصول",
                          "أول مكالمة", "آخر مكالمة", "آخر نتيجة", "آخر حالة (رئيسية › فرعية)", "مبلغ الوعد",
                          "تاريخ الوعد", "حالة الوعد", "المحصّل بعد التواصل (معتمد)", "المحصّل المسند", "الموظفون")
    assert len(clients) == 5                                                 # the 4 called clients
    promised = next(row for row in clients[1:] if row[10] is not None)
    assert promised[10] == 300 and promised[13] == 100 and promised[5] == "نعم" and promised[2].startswith("+9665")
    agents = list(wb["أداء الموظفين"].iter_rows(values_only=True))
    assert agents[0][:6] == ("الموظف", "التحويلة", "المحاولات", "أُرسلت", "محادثات", "نسبة الرد %")
    assert len(agents) == 3
    ev = q1(pgc, "SELECT actor_user_id::text, after->>'clients', after->>'payment_days' FROM audit_events "
                 "WHERE action = 'campaign.report.export' AND entity_id = %s", cid)
    assert ev == (w.supervisor, "4", "7")
