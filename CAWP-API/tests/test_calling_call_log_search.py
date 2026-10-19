"""Call list (customer request 2026-10-19): search by the client's number (normalised: digits, the
last 9), the XLSX download with the list's filters and scope (audited), and the recording field
gating (services/calling/queries.py, number_search.py, call_log_export.py; api/calling.py)."""
from __future__ import annotations

import io
import json
import uuid

import pytest

from test_calling_helpers import (CL, add_contact, cl, cl_app, cl_client, hdr, insert_attempt,  # noqa: F401
                                  new_supervisor, new_user, pgc, q1, run, team)


def get(client, h, url, **params):
    r = client.get(url, headers=h, query_string=params)
    return r.status_code, r.get_json()


def ids(body):
    return {i["id"] for i in body["data"]["items"]}


def numbered(pgc, team):
    """Three calls: by destination, by contact value (no destination), and another number."""
    local = f"05{uuid.uuid4().int % 10**8:08d}"                          # 05xxxxxxxx
    by_dest = insert_attempt(pgc, team["case1"], team["c1"], state="ended", bridged=True, extension="7101",
                             destination="+966" + local[1:])
    contact = add_contact(pgc, team["case2"], local)
    by_contact = insert_attempt(pgc, team["case2"], team["c2"], state="ended", contact_id=contact)
    other = insert_attempt(pgc, team["case1"], team["c1"], state="failed", destination="+966500000001")
    return local, by_dest, by_contact, other


def test_number_search_is_normalised(cl_client, hdr, team, pgc):
    local, by_dest, by_contact, other = numbered(pgc, team)
    h = hdr(team["supervisor"])
    for typed in (local, "966" + local[1:], "+966" + local[1:], "+966 " + local[1:4] + " " + local[4:],
                  local[1:], "٠" + "".join("٠١٢٣٤٥٦٧٨٩"[int(c)] for c in local[1:])):
        st, body = get(cl_client, h, "/v2/calls", number=typed)
        assert st == 200 and ids(body) == {by_dest, by_contact}, typed
        assert sum(body["data"]["counts_by_state"].values()) == 2
    st, body = get(cl_client, h, "/v2/calls", number=local[-6:])                    # 4-8 digits: partial
    assert {by_dest, by_contact} <= ids(body) and other not in ids(body)
    st, mine = get(cl_client, hdr(team["c1"]), "/v2/calls/mine", number=local)
    assert ids(mine) == {by_dest}
    st, bad = get(cl_client, h, "/v2/calls", number="12")
    assert (st, bad["error"]) == (422, {"number": "invalid"})
    # the numbers need clints_details, so does searching by one
    limited = new_user(pgc, CL.limited_role)
    st, body = get(cl_client, hdr(limited), "/v2/calls", number=local)
    assert st == 403 and body["error"] == {"number": "clints_details_required"}
    assert get(cl_client, hdr(limited), "/v2/calls")[0] == 200


def _sheet(data):
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
    head = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = [{h: ws.cell(r, c + 1).value for c, h in enumerate(head)} for r in range(2, ws.max_row + 1)]
    return head, rows


def test_export_mirrors_filters_and_scope_and_is_audited(cl_client, hdr, team, pgc):
    from services.calling import call_log_export
    local, by_dest, by_contact, other = numbered(pgc, team)
    r = cl_client.get("/v2/calls/export", headers=hdr(team["supervisor"]), query_string={"number": local})
    assert r.status_code == 200 and r.headers["Cache-Control"] == "private, no-store", r.data[:300]
    head, rows = _sheet(r.data)
    assert tuple(head) == call_log_export.HEADERS
    by = {x["معرف المكالمة"]: x for x in rows}
    assert set(by) == {by_dest, by_contact}
    assert by[by_dest]["الرقم"] == "+966" + local[1:] and by[by_contact]["الرقم"] == local
    assert by[by_dest]["تم الرد"] == "نعم" and by[by_contact]["تم الرد"] == "لا"
    assert by[by_dest]["الحالة"] == "انتهت" and by[by_dest]["التحويلة"] == "7101"
    assert by[by_dest]["مدة المحادثة (ثانية)"] is not None and by[by_dest]["العميل"] == "CL case"
    assert by[by_dest]["مسجلة"] == "غير معروف"                                   # CDR disabled: unknown, not "no"
    n = q1(pgc, "SELECT count(*) FROM audit_events WHERE action = 'calls.export' AND actor_user_id = %s",
           team["supervisor"])[0]
    assert n == 1
    after = q1(pgc, "SELECT after FROM audit_events WHERE action = 'calls.export' AND actor_user_id = %s",
               team["supervisor"])[0]
    after = after if isinstance(after, dict) else json.loads(after)
    assert after["kind"] == "team" and after["filters"] == {"number": local} and after["rows"] == 2
    # the list's filters: state + the collector's own calls
    r = cl_client.get("/v2/calls/export", headers=hdr(team["supervisor"]),
                      query_string={"state": "failed", "collector_id": team["c1"]})
    assert {x["معرف المكالمة"] for x in _sheet(r.data)[1]} == {other}
    r = cl_client.get("/v2/calls/mine/export", headers=hdr(team["c1"]))
    mine = {x["معرف المكالمة"] for x in _sheet(r.data)[1]}
    assert {by_dest, other} <= mine and by_contact not in mine                    # c2's call is not "mine"
    r = cl_client.get("/v2/calls/export", headers=hdr(team["outsider"]))
    assert not ({by_dest, by_contact, other} & {x["معرف المكالمة"] for x in _sheet(r.data)[1]})   # outside scope
    # numbers hidden without clints_details; no calling key at all = 403
    limited = new_user(pgc, CL.limited_role)
    own = insert_attempt(pgc, team["case1"], limited, state="ended", destination="+966500000002")
    r = cl_client.get("/v2/calls/mine/export", headers=hdr(limited))
    assert r.status_code == 200
    row = next(x for x in _sheet(r.data)[1] if x["معرف المكالمة"] == own)
    assert row["الرقم"] is None
    assert cl_client.get("/v2/calls/export", headers=hdr(new_user(pgc, CL.none_role))).status_code == 403
    assert cl_client.get("/v2/calls/export", headers=hdr(team["supervisor"]),
                         query_string={"number": "1"}).status_code == 422


def test_recording_field_needs_a_recording_key(cl_client, hdr, team, pgc):
    local, by_dest, by_contact, other = numbered(pgc, team)
    st, body = get(cl_client, hdr(team["supervisor"]), "/v2/calls", number=local)
    assert body["data"]["recordings_status"] == "no_key"
    assert all(i["recording"] is None for i in body["data"]["items"])
    listener = str(uuid.uuid4())
    run(pgc, "INSERT INTO roles (id, name, permission) VALUES (%s, %s, %s)", listener, f"cl listener {listener[:8]}",
        json.dumps({"calls": True, "calls_make": True, "clints_details": True, "clints_callsRecords_play": True}))
    sup = new_supervisor(pgc, role=listener, users=[team["c1"], team["c2"]])
    st, body = get(cl_client, hdr(sup), "/v2/calls", number=local)
    assert st == 200 and body["data"]["recordings_status"] == "disabled"           # the tests run with CDR_ENABLED=0
    assert all(i["recording"] is None for i in body["data"]["items"])
    # a case's calls keep their contract (no recording key added)
    st, body = get(cl_client, hdr(team["c1"]), f"/v2/cases/{team['case1']}/calls")
    assert "recordings_status" not in body["data"] and "recording" not in body["data"]["items"][0]


@pytest.mark.parametrize("typed,stored,hit", [
    ("0551234567", "+966551234567", True), ("966551234567", "0551234567", True), ("+966551234567", "551234567", True),
    ("0551234567", "+966551234568", False), ("1234", "0551234567", True), ("551234567", "12345", False),
])
def test_number_needle_and_match(typed, stored, hit):
    from services.calling import number_search as ns
    assert ns.matches(stored, ns.needle(typed)) is hit
