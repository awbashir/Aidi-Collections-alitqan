"""Stream R-B: stable paging (fully tied rows, every sort, both directions), full export =
full preview (CSV, every report), typed and Arabic RTL XLSX, formula safety, and the
recorder/owner rules applied again by the export worker."""
from __future__ import annotations

import datetime as dt

import pytest

from test_reports_v2_b_support import *  # noqa: F401,F403
from test_reports_v2_b_support import (ALL_B, B, key_of, make_case, make_contract, make_note, make_payment,
                                       make_portfolio, make_promise, q, read_csv, reassign, run_export)

N_CASES = 60


@pytest.fixture(scope="module")
def tied(repb_db):
    """60 unarchived col1 cases in one contract-A portfolio, all tied: same creation time, debt,
    promise (due date, amount, entry time), note time and callback; 30 of them with a pending
    payment. 25 tied portfolios in a fresh contract for RPT-12."""
    import psycopg2

    from conftest import MAIN_DB, db_url
    conn = psycopg2.connect(db_url(MAIN_DB))
    conn.autocommit = True
    pf = make_portfolio(conn, B.contract_a, name="RB tied")
    cases = [make_case(conn, pf, owner=B.col1, amount=1000, name="RB tied case") for _ in range(N_CASES)]
    q(conn, "UPDATE clints SET create_at = '2026-01-01T09:00:00Z' WHERE portfolios_id = %s", (pf,))
    for i, c in enumerate(cases):
        make_promise(conn, c, "5.00", "2026-10-01", recorder=B.col1, created="2026-02-01T09:00:00Z")
        make_note(conn, c, text="tied", created="2026-03-01T09:00:00Z", callback="2026-10-02T06:00:00Z")
        if i % 2 == 0:
            make_payment(conn, c, "7.00", None, recorder=B.col1, when="2026-03-10T21:00:00Z",
                         created="2026-03-11T09:00:00Z")
    contract = make_contract(conn, "RB tied contract")
    pfs = [make_portfolio(conn, contract, name="RB tied portfolio") for _ in range(25)]
    q(conn, "UPDATE portfolios SET create_at = '2026-01-05T09:00:00Z' WHERE contract_id = %s", (contract,))
    conn.close()
    return {"portfolio": pf, "contract": contract, "cases": cases, "portfolios": pfs}


def params_for(rid, tied):
    return {"contract": tied["contract"]} if rid == "RPT-12" else {"portfolio": tied["portfolio"]}


EXPECTED_ROWS = {"RPT-09": 60, "RPT-10": 60, "RPT-11": 60, "RPT-12": 25, "RPT-13": 30, "RPT-14": 30,
                 "RPT-15": 60, "RPT-16": 60}


def user_for(rid):
    # RPT-12's tied portfolios are empty and in an unsupervised contract: only SuperAdmin sees them
    return B.col1 if rid in ("RPT-10", "RPT-11", "RPT-16") else B.super if rid == "RPT-12" else B.sup


@pytest.mark.parametrize("rid", ALL_B)
def test_paging_is_stable_for_every_sort(rb, tied, rid):
    from core.reports import registry
    report = registry.get(rid)
    key = key_of(rid)
    params = params_for(rid, tied)
    user = user_for(rid)
    for s in report.sorts:
        for sort in (s.key, "-" + s.key):
            whole = rb.rows(rid, user, {**params, "sort": sort, "page_size": 200})
            assert whole["total"] == EXPECTED_ROWS[rid] and len(whole["items"]) == EXPECTED_ROWS[rid], (rid, sort)
            paged = []
            for page in range(1, (EXPECTED_ROWS[rid] + 19) // 20 + 1):
                d = rb.rows(rid, user, {**params, "sort": sort, "page_size": 20, "page": page})
                assert d["totals"] == whole["totals"], (rid, sort, page)
                paged.extend(i[key] for i in d["items"])
            assert paged == [i[key] for i in whole["items"]], (rid, sort)
            assert len(set(paged)) == len(paged)
    # every row ties on the default sort key here, so the id tie-breaker orders them, in the
    # sort's direction
    ids = [i[key] for i in rb.rows(rid, user, {**params, "page_size": 200})["items"]]
    assert ids == sorted(ids, reverse=report.default_sort.startswith("-"))
    # an empty page past the end still carries the full totals
    past = rb.rows(rid, user, {**params, "page_size": 20, "page": 99})
    assert past["items"] == [] and past["total"] == EXPECTED_ROWS[rid]


@pytest.mark.parametrize("rid", ALL_B)
def test_full_csv_export_equals_the_full_preview(rb, rep_app, db, tied, rid):
    params = params_for(rid, tied)
    user = user_for(rid)
    items, first = rb.all_rows(rid, user, params)
    assert len(items) == EXPECTED_ROWS[rid]
    exp = rb.export(rid, user, params, "csv")
    assert run_export(rep_app, exp["id"]) == "succeeded"
    done = rb.export_get(exp["id"], user)
    assert done["row_count"] == len(items) and done["totals"] == first["totals"]
    header, rows = read_csv(rb.download(exp["id"], user).data)
    cols = done["definition"]["columns"]
    assert header == cols
    expected = [["" if i[c] is None else ("true" if i[c] is True else "false" if i[c] is False else str(i[c]))
                 for c in cols] for i in items]
    assert rows == expected


def test_xlsx_values_are_typed_and_equal_the_preview(rb, rep_app, db):
    from openpyxl import load_workbook

    from test_reports_v2_support import artifact_path
    pf = make_portfolio(db, B.contract_a)
    case = make_case(db, pf, owner=B.col1, amount=1234.5)
    make_promise(db, case, "250.75", "2026-03-31T21:00:00Z")        # civil 2026-04-01
    x = rb.export("RPT-09", B.sup, {"portfolio": pf}, "xlsx")
    run_export(rep_app, x["id"])
    wb = load_workbook(artifact_path(db, x["id"]))
    assert wb.sheetnames == ["Report", "Summary"]
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    row = dict(zip(hdr, [c.value for c in ws[2]]))
    assert row["State"] == "Overdue" and row["Promised amount"] == 250.75 and row["Due date"] == dt.datetime(2026, 4, 1)
    assert row["Original debt (per case, not summed)"] == 1234.5 and row["Imported, fulfilment unknown"] == "No"
    assert ws.cell(row=2, column=hdr.index("Promised amount") + 1).number_format == "#,##0.00"
    summary = {r[0].value: r[1].value for r in wb.worksheets[1].iter_rows()}
    assert summary["Promised, all states (obligation, not recovery)"] == 250.75 and summary["Rows"] == 1


@pytest.mark.parametrize("rid,user,sheet,headers,labels", [
    ("RPT-09", B.sup, "مواعيد السداد", ["رقم الموعد", "الحالة", "موعد السداد", "مبلغ السداد"], {"الحالة": "متأخر"}),
    ("RPT-10", B.col1, "مواعيدي", ["رقم الموعد", "الحالة"], {"الحالة": "متأخر"}),
    ("RPT-15", B.sup, "تقرير الاهمال", ["معرف العميل", "رقم الحساب", "اسم العميل"], {"أخر متابعة منذ (يوم)": None}),
    ("RPT-16", B.col1, "المتابعات", ["رقم المتابعة", "تاريخ الاتصال القادم"],
     {"حالة المتابعة": "غير متتبعة (قبل نظام المتابعات)"}),
    ("RPT-13", B.sup, "العملاء النشطين", ["معرف العميل"], {"عدد الدفعات": 1}),
    ("RPT-14", B.sup, "العملاء الغير نشطين", ["معرف العميل"], {"مبلغ المديونية": 900}),
])
def test_arabic_xlsx_is_rtl_with_arabic_headers_and_values(rb, rep_app, db, rid, user, sheet, headers, labels):
    from openpyxl import load_workbook

    from test_reports_v2_support import artifact_path
    pf = make_portfolio(db, B.contract_a)
    case = make_case(db, pf, owner=B.col1, amount=900, name="عميل عربي")
    if rid in ("RPT-09", "RPT-10"):
        make_promise(db, case, "10.00", "2026-01-15", recorder=B.col1)
    if rid == "RPT-16":
        make_note(db, case, callback="2025-01-05T06:00:00Z", created="2025-01-01T09:00:00Z", track=False)
    if rid == "RPT-13":
        make_payment(db, case, "3.00", None, recorder=B.col1)
    x = rb.export(rid, user, {"portfolio": pf}, "xlsx_ar")
    assert run_export(rep_app, x["id"]) == "succeeded"
    done = rb.export_get(x["id"], user)
    assert done["artifact"]["file_name"].endswith("-ar.xlsx")
    wb = load_workbook(artifact_path(db, x["id"]))
    assert wb.sheetnames == [sheet, "الملخص"]
    for ws in wb.worksheets:
        assert ws.sheet_view.rightToLeft is True
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    assert hdr[:len(headers)] == headers
    row = dict(zip(hdr, [c.value for c in ws[2]]))
    for k, v in labels.items():
        assert row[k] == v, (k, row[k])
    assert row.get("اسم العميل") == "عميل عربي"
    summary = {r[0].value: r[1].value for r in wb.worksheets[1].iter_rows()}
    assert summary["عدد السجلات"] == 1 and summary["التقرير"].startswith(rid)
    assert summary["الجهة"] == "فنيه الاتقان لتحصيل الديون"


def test_arabic_xlsx_for_portfolios(rb, rep_app, db):
    from openpyxl import load_workbook

    from test_reports_v2_support import artifact_path
    contract = make_contract(db, "RB عقد")
    pf = make_portfolio(db, contract, name="RB محفظة")
    make_case(db, pf, owner=B.col2, amount=321.5)
    x = rb.export("RPT-12", B.super, {"contract": contract}, "xlsx_ar")
    run_export(rep_app, x["id"])
    wb = load_workbook(artifact_path(db, x["id"]))
    assert wb.sheetnames == ["تقرير المحافظ", "الملخص"]
    ws = wb.worksheets[0]
    row = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
    assert row["المحفظة"] == "RB محفظة" and row["العقد"] == "RB عقد" and row["عدد العملاء"] == 1
    assert row["مبلغ المديونية"] == 321.5 and row["المتبقي"] == 321.5 and row["المحفظة محذوفة"] == "لا"


def test_note_text_is_formula_safe_in_csv_and_xlsx(rb, rep_app, db):
    from openpyxl import load_workbook

    from test_reports_v2_support import artifact_path
    pf = make_portfolio(db, B.contract_a)
    case = make_case(db, pf, owner=B.col1, name="+SUM(A1)")
    make_note(db, case, text="=HYPERLINK(\"http://x\")", callback="2026-10-05T06:00:00Z")
    for rid, user, col in (("RPT-15", B.sup, "last_note_text"), ("RPT-16", B.col1, "note_text")):
        x = rb.export(rid, user, {"portfolio": pf}, "csv")
        run_export(rep_app, x["id"])
        header, rows = read_csv(rb.download(x["id"], user).data)
        cell = dict(zip(header, rows[0]))
        assert cell[col] == "'=HYPERLINK(\"http://x\")" and cell["customer_name"] == "'+SUM(A1)"
        preview = rb.rows(rid, user, {"portfolio": pf})["items"][0]
        assert preview[col] == "=HYPERLINK(\"http://x\")"                   # the preview is the raw value
        x = rb.export(rid, user, {"portfolio": pf}, "xlsx")
        run_export(rep_app, x["id"])
        path = artifact_path(db, x["id"])
        wb = load_workbook(path)
        ws = wb.worksheets[0]
        row = [c.value for c in ws[2]]
        assert "=HYPERLINK(\"http://x\")" in row                            # a typed string, not a formula
        import zipfile
        with zipfile.ZipFile(path) as z:
            assert "<f>" not in z.read("xl/worksheets/sheet1.xml").decode("utf-8")


def test_the_worker_applies_the_recorder_and_owner_rules(rb, rep_app, db):
    pb = make_portfolio(db, B.contract_b)
    theirs = make_case(db, pb, owner=B.col2)
    p = make_promise(db, theirs, "8.00", "2026-10-01", recorder=B.col1)       # col1's work on col2's case
    make_note(db, theirs, user=B.col1, callback="2026-10-03T06:00:00Z")
    x10 = rb.export("RPT-10", B.col1, {"portfolio": pb}, "csv")
    x11 = rb.export("RPT-11", B.col1, {"portfolio": pb}, "csv")
    x16 = rb.export("RPT-16", B.col2, {"portfolio": pb}, "csv")
    reassign(db, theirs, B.col1)                                     # after the request, before generation
    for x, user in ((x10, B.col1), (x11, B.col1), (x16, B.col2)):
        assert run_export(rep_app, x["id"]) == "succeeded"
    rows10 = read_csv(rb.download(x10["id"], B.col1).data)[1]
    assert [r[0] for r in rows10] == [p]                              # the recorder's row either way
    assert rb.export_get(x11["id"], B.col1)["row_count"] == 1        # col1 owns the case at generation
    assert rb.export_get(x16["id"], B.col2)["row_count"] == 0        # col2 no longer owns it
