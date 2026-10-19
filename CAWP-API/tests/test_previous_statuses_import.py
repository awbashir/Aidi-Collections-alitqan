"""Previous statuses import (services/previous_status_import.py, api/previous_status_import.py,
migration 0074). SYNTHETIC files only.

Header detection, `---` rows skipped, unknown accounts, archived / locked cases, deleted statuses
never matched, the contract's own status preferred, missing sub-statuses created (and not when the
option is off), the note written with main + sub + the chosen create_at, the last-comment trigger
(an older imported note never replaces a newer note; it becomes the latest when there is none),
idempotent re-runs, the XLSX report, and the permission gate.

The background job is run inline: `api.previous_status_import._start` is replaced by a direct call
to `services.previous_status_import.run_import` (the function the thread pool runs)."""
from __future__ import annotations

import datetime as _dt
import io
import os
import uuid
from pathlib import Path

import pytest

from conftest import MAIN_DB, db_url, run_sql_file
from test_bulk_framework_support import bulk_storage, make_csv, make_xlsx, q, q1  # noqa: F401

SEED = Path(__file__).resolve().parent / "fixtures" / "seed_previous_status_import.sql"
CONTACT_METHOD = "00000000-0000-4000-8000-000000000701"      # seed.sql 'Phone'
HEAD = ["رقم الحساب", "الحالة الرئيسية", "الحالة الفرعية"]


class P:
    admin = "e9e9e9e9-0000-4000-8000-000000000501"
    sup = "e9e9e9e9-0000-4000-8000-000000000502"
    nokey = "e9e9e9e9-0000-4000-8000-000000000503"
    portfolio = "e9e9e9e9-0000-4000-8000-000000000401"
    other_portfolio = "e9e9e9e9-0000-4000-8000-000000000402"
    reach = "e9e9e9e9-0000-4000-8000-000000000601"
    promise = "e9e9e9e9-0000-4000-8000-000000000701"          # وعد بالسداد PS
    promise_reached = "e9e9e9e9-0000-4000-8000-000000000702"  # تم التواصل PS (sub)
    evades = "e9e9e9e9-0000-4000-8000-000000000703"           # يتهرب من الاتصال PS
    evades_whatsapp_deleted = "e9e9e9e9-0000-4000-8000-000000000704"
    contract_own = "e9e9e9e9-0000-4000-8000-000000000706"     # حالة العقد PS (PS contract)
    contract_general = "e9e9e9e9-0000-4000-8000-000000000707"


PROMISE, REACHED, EVADES, NO_ANSWER, CONTRACT = ("وعد بالسداد PS", "تم التواصل PS", "يتهرب من الاتصال PS",
                                                 "لا يرد PS", "حالة العقد PS")


# --------------------------------------------------------------------------- fixtures
@pytest.fixture(scope="session")
def ps_db(databases):
    run_sql_file(MAIN_DB, SEED)
    return databases


@pytest.fixture(scope="session")
def ps_app(ps_db, bulk_storage):  # noqa: F811
    from flask import Flask
    from sqlalchemy.pool import NullPool

    from extensions import db, jwt

    app = Flask("cawp_previous_status_import_test")
    app.config.update(
        TESTING=True,
        SQLALCHEMY_DATABASE_URI=db_url(MAIN_DB, "psycopg2"),
        SQLALCHEMY_ENGINE_OPTIONS={"poolclass": NullPool},
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        JWT_SECRET_KEY=os.environ["JWT_SECRET_KEY"],
    )
    db.init_app(app)
    jwt.init_app(app)
    import api.previous_status_import
    app.register_blueprint(api.previous_status_import.bp)
    from core import lookups
    lookups.invalidate()
    return app


@pytest.fixture(autouse=True)
def inline_jobs(request, monkeypatch):
    started = []
    if request.node.get_closest_marker("no_db"):
        return started
    import api.previous_status_import as route
    from services import previous_status_import as svc

    def start(import_id, actor):
        started.append(import_id)
        svc.run_import(import_id)

    monkeypatch.setattr(route, "_start", start)
    return started


class Client:
    def __init__(self, app):
        self.app = app
        self.client = app.test_client()
        self._tokens = {}

    def headers(self, user, key=None):
        if user not in self._tokens:
            from flask_jwt_extended import create_access_token
            with self.app.app_context():
                self._tokens[user] = create_access_token(identity=user)
        h = {"Authorization": f"Bearer {self._tokens[user]}"}
        if key is not False:
            h["Idempotency-Key"] = key or str(uuid.uuid4())
        return h

    def get(self, url, user=P.admin):
        return self.client.get(url, headers=self.headers(user, key=False))

    def post(self, url, body, user=P.admin, key=None):
        return self.client.post(url, json=body, headers=self.headers(user, key))

    def upload(self, content, name="statuses.xlsx", user=P.admin):
        return self.client.post("/v2/previous-status-imports/uploads",
                                data={"file": (io.BytesIO(content), name)},
                                content_type="multipart/form-data", headers=self.headers(user))


@pytest.fixture
def pc(ps_app):
    return Client(ps_app)


@pytest.fixture
def db(ps_db):
    import psycopg2
    conn = psycopg2.connect(db_url(MAIN_DB))
    conn.autocommit = True
    yield conn
    conn.close()


# --------------------------------------------------------------------------- helpers
def tag() -> str:
    return uuid.uuid4().hex[:8].upper()


def make_case(conn, account, archived=False, locked=False, portfolio=P.portfolio):
    cid = str(uuid.uuid4())
    q(conn, """INSERT INTO clints (id, name, "identityNumber", "accountNumber", amount, attribute, user_id,
                                   portfolios_id, is_archive, is_locked)
               VALUES (%s, %s, %s, %s, 100, '{}', NULL, %s, %s, %s)""",
      (cid, "PS " + account.strip(), "8" + cid[:9].replace("-", "0"), account, portfolio, archived, locked))
    return cid


def add_note(conn, case_id, status=P.promise, at=None):
    nid = str(uuid.uuid4())
    q(conn, """INSERT INTO "Comments" (id, clint_id, user_id, "contactMethod_id", comment, comment_status_id, create_at)
               VALUES (%s, %s, %s, %s, 'collector note', %s, coalesce(%s, now()))""",
      (nid, case_id, P.admin, CONTACT_METHOD, status, at))
    return nid


def uploaded(pc, rows, name="statuses.xlsx", user=P.admin):
    content = make_csv(rows) if name.endswith(".csv") else make_xlsx(rows, sheet="XlsxTable")
    r = pc.upload(content, name, user)
    assert r.status_code == 201, r.get_json()
    return r.get_json()["data"]


def preview(pc, up, expect=200, **opts):
    r = pc.post("/v2/previous-status-imports/preview", {"upload_id": up["id"], **opts})
    assert r.status_code == expect, r.get_json()
    return r.get_json()["data"] if expect == 200 else r.get_json()


def apply(pc, up, expect=202, noted_on="2026-10-11", key=None, **opts):
    body = {"upload_id": up["id"], "noted_on": noted_on, "contact_method_id": CONTACT_METHOD, **opts}
    r = pc.post("/v2/previous-status-imports/apply", body, key=key)
    assert r.status_code == expect, r.get_json()
    return r.get_json()


def status(pc, import_id):
    r = pc.get(f"/v2/previous-status-imports/{import_id}")
    assert r.status_code == 200, r.get_json()
    return r.get_json()["data"]


def notes_of(conn, case_id):
    return q(conn, """SELECT id::text, comment_status_id::text, comment_sub_status_id::text, comment, user_id::text,
                             "contactMethod_id"::text, comment_reachability_id::text, next_contact_date, create_at
                        FROM "Comments" WHERE clint_id = %s AND is_deleted = false ORDER BY create_at""", (case_id,))


def last_comment(conn, case_id):
    return q1(conn, "SELECT last_comment_id::text, last_comment_at FROM clints WHERE id = %s", (case_id,))


RIYADH_1011 = _dt.datetime(2026, 10, 10, 21, 0, tzinfo=_dt.timezone.utc)     # 2026-10-11 00:00 Asia/Riyadh


# --------------------------------------------------------------------------- unit
@pytest.mark.no_db
def test_header_guess_and_value_rules():
    from services.previous_status_import import guess_mapping, norm_name, status_value
    assert guess_mapping(["رقم الحساب", "الحالة الرئيسية", "الحالة الفرعية"]) == {
        "account_number": "رقم الحساب", "main_status": "الحالة الرئيسية", "sub_status": "الحالة الفرعية"}
    assert guess_mapping(["Account Number", "Status", "Sub Status", "x"]) == {
        "account_number": "Account Number", "main_status": "Status", "sub_status": "Sub Status"}
    assert guess_mapping(["الحساب", "الحالة"]) == {"account_number": "الحساب", "main_status": "الحالة"}
    assert status_value("---") is None and status_value("  ") is None and status_value(None) is None
    assert status_value(" وعد   بالسداد ") == "وعد بالسداد"
    assert norm_name("وعد بالسداد‏") == "وعد بالسداد"
    assert norm_name("حالة") != norm_name("حاله")                  # letters are kept exactly


# --------------------------------------------------------------------------- upload / options
def test_upload_detects_the_three_headers_and_options(pc):
    up = uploaded(pc, [HEAD, ["084832", PROMISE, REACHED]])
    assert up["purpose"] == "previous_statuses_import" and up["guess_sheet"] == "XlsxTable"
    assert up["guess"] == {"account_number": "رقم الحساب", "main_status": "الحالة الرئيسية",
                           "sub_status": "الحالة الفرعية"}
    r = pc.get(f"/v2/previous-status-imports/options?upload_id={up['id']}")
    assert r.status_code == 200, r.get_json()
    data = r.get_json()["data"]
    assert data["guess"] == up["guess"] and data["headers"] == HEAD
    # a call-like method (the shared test DB may hold several, e.g. other suites' 'Phone' rows)
    from services.previous_status_import import default_contact_method
    assert data["default_contact_method_id"] == default_contact_method(data["contact_methods"])
    assert data["default_contact_method_id"] in {m["id"] for m in data["contact_methods"]}
    assert data["default_noted_on"] == "2026-10-11" and data["marker"] == "حالة سابقة (من النظام السابق)"
    # other headers: mapped explicitly
    t = tag()
    csv_up = uploaded(pc, [["acc", "st", "sub"], [f"M-{t}", PROMISE, ""]], name="other.csv")
    assert csv_up["guess"] == {}
    bad = pc.post("/v2/previous-status-imports/preview", {"upload_id": csv_up["id"]})
    assert bad.status_code == 422 and bad.get_json()["code"] == "invalid_mapping"
    data = preview(pc, csv_up, mapping={"account_number": "acc", "main_status": "st", "sub_status": "sub"})
    assert data["counts"]["account_not_found"] == 1


# --------------------------------------------------------------------------- preview
def test_preview_counts_without_writes(pc, db):
    t = tag()
    ready = make_case(db, f"0{t}")                       # leading zero kept (text)
    trimmed = make_case(db, f" T-{t} ")                   # stored with spaces around it
    make_case(db, f"AR-{t}", archived=True)
    make_case(db, f"LK-{t}", locked=True)
    make_case(db, f"NM-{t}")
    make_case(db, f"NS-{t}")
    own = make_case(db, f"OWN-{t}")
    other = make_case(db, f"OTH-{t}", portfolio=P.other_portfolio)
    rows = [HEAD,
            [f"0{t}", PROMISE, REACHED],
            [f"T-{t}", f"  {PROMISE} ", "---"],             # trimmed account, status with spaces, no sub
            [f"GONE-{t}", PROMISE, REACHED],
            [f"AR-{t}", PROMISE, REACHED],
            [f"LK-{t}", PROMISE, REACHED],
            [f"NM-{t}", NO_ANSWER, "---"],                   # only a DELETED main has this name
            [f"NS-{t}", EVADES, f"واتساب PS"],               # the sub-status is deleted -> missing
            [f"X-{t}", "---", "---"],
            [f"0{t}", EVADES, ""],                           # repeats row 2
            [f"OWN-{t}", CONTRACT, ""],
            [f"OTH-{t}", CONTRACT, ""]]
    up = uploaded(pc, rows)
    before = (q1(db, 'SELECT count(*) FROM "Comments"')[0], q1(db, 'SELECT count(*) FROM "CommentStatuses"')[0])
    data = preview(pc, up)
    c = data["counts"]
    assert (c["total"], c["ready"], c["no_status"], c["account_not_found"], c["case_archived"], c["case_locked"],
            c["status_not_found"], c["duplicate_in_file"]) == (11, 5, 1, 1, 1, 1, 1, 1)
    assert c["ready_creating_sub_status"] == 1 and c["sub_statuses_to_create"] == 1
    assert data["sub_statuses_to_create"] == [{"main_id": P.evades, "main_name": EVADES, "name": "واتساب PS",
                                               "rows": 1}]
    assert data["missing_main_statuses"] == [{"name": NO_ANSWER, "rows": 1}]
    by_row = {p["row_number"]: p for p in data["problems"]}
    assert by_row[4]["code"] == "account_not_found" and by_row[7]["code"] == "status_not_found"
    assert by_row[8]["creates_sub_status"] is True and by_row[10]["code"] == "duplicate_in_file"
    assert 2 not in by_row and 9 not in by_row             # clean ready rows and `---` rows are not listed
    assert data["can_apply"] is True
    after = (q1(db, 'SELECT count(*) FROM "Comments"')[0], q1(db, 'SELECT count(*) FROM "CommentStatuses"')[0])
    assert after == before
    # option off: the missing sub-status is reported instead
    off = preview(pc, up, create_missing_sub_statuses=False)
    assert off["counts"]["sub_status_not_found"] == 1 and off["counts"]["ready"] == 4
    assert off["missing_sub_statuses"] == [{"main_name": EVADES, "name": "واتساب PS", "rows": 1}]
    # archived cases on request
    assert preview(pc, up, include_archived=True)["counts"]["ready"] == 6
    assert ready and trimmed and own and other


# --------------------------------------------------------------------------- apply
def test_apply_writes_notes_with_main_sub_and_the_chosen_date(pc, db, inline_jobs):
    t = tag()
    newer = make_case(db, f"NW-{t}")
    fresh = make_case(db, f"FR-{t}")
    created_case = make_case(db, f"CR-{t}")
    own = make_case(db, f"OWN-{t}")
    collector_note = add_note(db, newer)                   # written today by a collector
    sub_name = f"واتساب {t}"
    up = uploaded(pc, [HEAD,
                       [f"NW-{t}", PROMISE, REACHED],
                       [f"FR-{t}", PROMISE, "---"],
                       [f"CR-{t}", EVADES, sub_name],
                       [f"OWN-{t}", CONTRACT, "---"],
                       [f"NONE-{t}", PROMISE, REACHED],
                       [f"Z-{t}", "---", "---"]])
    out = apply(pc, up, note="من ملف موبايلي")
    imp = out["data"]
    assert imp["status"] == "queued" and inline_jobs == [imp["id"]]
    done = status(pc, imp["id"])
    assert done["status"] == "succeeded", done
    assert done["counts"]["imported"] == 4 and done["counts"]["account_not_found"] == 1
    assert done["counts"]["no_status"] == 1 and done["counts"]["sub_statuses_created"] == 1

    # the created sub-status: under its main, with the main's contract and reachability
    created = done["created_sub_statuses"]
    assert len(created) == 1 and created[0]["name"] == sub_name and created[0]["parent_id"] == P.evades
    sub = q1(db, """SELECT parent_id::text, contract_id, "CommentReachability_id", deleted_at, is_active
                      FROM "CommentStatuses" WHERE id = %s""", (created[0]["id"],))
    assert sub == (P.evades, None, None, None, False)

    marker = "حالة سابقة (من النظام السابق)"
    n = notes_of(db, created_case)
    assert len(n) == 1
    _id, st, sb, text_, user, method, reach, callback, at = n[0]
    assert (st, sb, user, method, reach, callback) == (P.evades, created[0]["id"], P.admin, CONTACT_METHOD, None, None)
    assert text_ == marker + "\nمن ملف موبايلي" and at == RIYADH_1011
    n = notes_of(db, fresh)
    assert [(r[1], r[2], r[6]) for r in n] == [(P.promise, None, P.reach)]    # main's reachability
    assert [r[1] for r in notes_of(db, own)] == [P.contract_own]              # the contract's own status wins

    # last comment: a newer collector note stays the latest; without notes the import becomes it
    nn = notes_of(db, newer)
    assert len(nn) == 2 and nn[0][1:3] == (P.promise, P.promise_reached) and nn[0][8] == RIYADH_1011
    assert last_comment(db, newer)[0] == collector_note
    assert last_comment(db, fresh) == (notes_of(db, fresh)[0][0], RIYADH_1011)

    # report rows + XLSX of the skipped rows
    rows = {r[0]: r for r in q(db, """SELECT row_number, outcome, code, comment_id::text FROM previous_status_import_rows
                                       WHERE import_id = %s""", (imp["id"],))}
    assert len(rows) == 6 and rows[2][1] == "imported" and rows[2][3] == nn[0][0]
    assert rows[6][1:3] == ("skipped", "account_not_found") and rows[7][2] == "no_status"
    r = pc.get(f"/v2/previous-status-imports/{imp['id']}/report")
    assert r.status_code == 200 and r.mimetype.endswith("spreadsheetml.sheet")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    ws = wb.worksheets[0]
    assert ws.cell(1, 2).value == "رقم الحساب" and ws.max_row == 3
    assert {ws.cell(i, 5).value for i in (2, 3)} == {"account_not_found", "no_status"}
    assert [s.title for s in wb.worksheets] == ["الصفوف المتجاوزة", "الملخص", "الحالات الفرعية المنشأة"]
    assert q1(db, """SELECT count(*) FROM audit_events WHERE source = 'previous_status_import' AND action = 'note.create'
                     AND request_id = %s""", (imp["id"],))[0] == 4
    assert q1(db, """SELECT count(*) FROM audit_events WHERE entity_type = 'previous_status_import'
                     AND entity_id = %s AND action = 'previous_status_import.succeed'""", (imp["id"],))[0] == 1

    # the same upload never runs twice; another upload of the same file adds nothing
    again = apply(pc, up, expect=409)
    assert again["code"] == "already_imported"
    up2 = uploaded(pc, [HEAD,
                        [f"NW-{t}", PROMISE, REACHED],
                        [f"FR-{t}", PROMISE, "---"],
                        [f"CR-{t}", EVADES, sub_name],
                        [f"OWN-{t}", CONTRACT, "---"]])
    p2 = preview(pc, up2)
    assert p2["counts"]["already_imported"] == 4 and p2["counts"]["ready"] == 0
    assert p2["blocking"] == [{"code": "nothing_to_import"}] and p2["can_apply"] is False
    out2 = apply(pc, up2)["data"]
    st2 = status(pc, out2["id"])
    assert st2["status"] == "failed" and st2["error"]["code"] == "nothing_to_import"
    assert len(notes_of(db, newer)) == 2 and len(notes_of(db, created_case)) == 1
    assert q1(db, """SELECT count(*) FROM "CommentStatuses" WHERE parent_id = %s AND name = %s""",
              (P.evades, sub_name))[0] == 1


def test_apply_without_creating_sub_statuses_skips_those_rows(pc, db):
    t = tag()
    case = make_case(db, f"OFF-{t}")
    ok_case = make_case(db, f"OK-{t}")
    sub_name = f"جديدة {t}"
    up = uploaded(pc, [HEAD, [f"OFF-{t}", EVADES, sub_name], [f"OK-{t}", EVADES, ""]])
    imp = apply(pc, up, create_missing_sub_statuses=False, noted_on="2026-10-01")["data"]
    done = status(pc, imp["id"])
    assert done["status"] == "succeeded" and done["counts"]["imported"] == 1
    assert done["counts"]["sub_status_not_found"] == 1 and done["created_sub_statuses"] == []
    assert notes_of(db, case) == []
    assert q1(db, 'SELECT count(*) FROM "CommentStatuses" WHERE name = %s', (sub_name,))[0] == 0
    n = notes_of(db, ok_case)
    assert len(n) == 1 and n[0][8] == _dt.datetime(2026, 9, 30, 21, 0, tzinfo=_dt.timezone.utc)


def test_deleted_statuses_are_never_matched(pc, db):
    t = tag()
    case = make_case(db, f"DEL-{t}")
    other = make_case(db, f"DEL2-{t}")
    main_name, sub_name = f"محذوفة {t}", f"فرعية محذوفة {t}"
    q(db, """INSERT INTO "CommentStatuses" (name, color, is_active, deleted_at) VALUES (%s, '#FFFFFF', false, now())""",
      (main_name,))
    deleted_sub = q1(db, """INSERT INTO "CommentStatuses" (name, color, is_active, parent_id, deleted_at)
                            VALUES (%s, '#FFFFFF', false, %s, now()) RETURNING id::text""", (sub_name, P.evades))[0]
    up = uploaded(pc, [HEAD, [f"DEL-{t}", main_name, ""], [f"DEL2-{t}", EVADES, sub_name]])
    data = preview(pc, up, create_missing_sub_statuses=False)
    assert data["counts"]["status_not_found"] == 1 and data["counts"]["sub_status_not_found"] == 1
    assert data["missing_main_statuses"] == [{"name": main_name, "rows": 1}]
    imp = apply(pc, up)["data"]
    done = status(pc, imp["id"])
    assert done["status"] == "succeeded" and done["counts"]["imported"] == 1
    sub_id = done["created_sub_statuses"][0]["id"]
    assert sub_id != deleted_sub                                     # a new live sub, not the deleted one
    assert notes_of(db, other)[0][2] == sub_id and notes_of(db, case) == []
    # the seed's deleted sub "واتساب PS" is not matched either
    from services.previous_status_import import load_statuses
    from extensions import db as _db
    with pc.app.app_context():
        _mains, subs = load_statuses(_db.session)
        assert (P.evades, "واتساب PS") not in subs or all(
            s["id"] != P.evades_whatsapp_deleted for s in subs[(P.evades, "واتساب PS")])
        _db.session.rollback()


def test_noted_on_must_be_before_today(pc, db):
    from core import riyadh
    t = tag()
    make_case(db, f"DT-{t}")
    up = uploaded(pc, [HEAD, [f"DT-{t}", PROMISE, REACHED]])
    r = apply(pc, up, expect=422, noted_on=riyadh.today().isoformat())
    assert r["error"] == {"noted_on": "not_before_today"}
    r = apply(pc, up, expect=422, noted_on="11/10/2026")
    assert r["error"] == {"noted_on": "invalid"}


def test_permission_gate(pc, db):
    t = tag()
    content = make_xlsx([HEAD, [f"PG-{t}", PROMISE, REACHED]])
    r = pc.upload(content, user=P.nokey)
    assert r.status_code == 403 and r.get_json()["code"] == "forbidden"
    r = pc.upload(content, user=P.sup)
    assert r.status_code == 403 and r.get_json()["code"] == "super_admin_required"
    up = uploaded(pc, [HEAD, [f"PG-{t}", PROMISE, REACHED]])
    for user, code in ((P.nokey, "forbidden"), (P.sup, "super_admin_required")):
        r = pc.post("/v2/previous-status-imports/preview", {"upload_id": up["id"]}, user)
        assert r.status_code == 403 and r.get_json()["code"] == code
        r = pc.post("/v2/previous-status-imports/apply", {"upload_id": up["id"]}, user)
        assert r.status_code == 403
        assert pc.get("/v2/previous-status-imports/options", user).status_code == 403
