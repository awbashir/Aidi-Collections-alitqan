"""Call list recordings (customer request 2026-10-19) against the local CDR fixture: the list's
`recording` comes from the call's tagged CDR rows (userfield cawp:<attempt>, uniqueid cawp-<attempt>
or its 32-character truncation), in one CDR query, and the legacy recording route authorizes
those dated cawp paths (services/calling/call_recordings.py, services/calling_cdr.py)."""
from __future__ import annotations

import io

import pytest

from test_calling_cdr_support import *  # noqa: F401,F403
from test_calling_cdr_support import P4, c4_db, cdr, cdr_reachable, set_perm  # noqa: F401

DAY = {"from": P4.DAY, "to": P4.DAY}


@pytest.fixture(scope="module")
def calls_app(c4_db):
    import os

    from flask import Flask
    from sqlalchemy.pool import NullPool

    import api.calling
    from conftest import MAIN_DB, db_url
    from core.serialization import ORJSONProvider
    from extensions import db, jwt

    app = Flask("cawp_test_call_log_recordings")
    app.config.update(TESTING=True, SQLALCHEMY_DATABASE_URI=db_url(MAIN_DB, "psycopg2"),
                      SQLALCHEMY_ENGINE_OPTIONS={"poolclass": NullPool}, SQLALCHEMY_TRACK_MODIFICATIONS=False,
                      JWT_SECRET_KEY=os.environ["JWT_SECRET_KEY"])
    app.json = ORJSONProvider(app)
    db.init_app(app)
    jwt.init_app(app)
    app.register_blueprint(api.calling.bp)
    return app


@pytest.fixture
def cc(calls_app):
    from core import lookups

    from test_bulk_framework_support import BulkClient
    lookups.invalidate()
    return BulkClient(calls_app)


def mine(cc, user=P4.cola, **q):
    r = cc.get("/v2/calls/mine", user, query_string={**DAY, **q})
    assert r.status_code == 200, r.get_json()
    return r.get_json()["data"]


def test_list_recordings_from_the_tagged_cdr_rows(cc, cdr, calls_app):
    d = mine(cc)
    assert d["recordings_status"] == "available"
    by = {i["id"]: i["recording"] for i in d["items"]}
    assert by[P4.AT1] == {"has_recording": True, "file": P4.REC_AT1,
                          "url": "/cdr/recording?recordingPath=2026%2F09%2F22%2Fcawp-a1c4a001-0000-4000-8000-000000000001.wav"}
    assert by[P4.AT3]["file"] == P4.REC_AT3                                   # truncated uniqueid only
    assert by[P4.AT4]["file"] == P4.REC_AT4
    # the legacy recording route authorizes the dated cawp path the list hands out
    from services.calling_cdr import authorize_recording
    with calls_app.app_context():
        assert authorize_recording(P4.cola, by[P4.AT1]["file"])["allowed"] is True
        assert authorize_recording(P4.colb, by[P4.AT1]["file"])["allowed"] is False   # not colb's case or call


def test_list_recordings_need_a_key_and_unknown_is_not_none(cc, cdr, c4pg, monkeypatch):
    set_perm(c4pg, P4.role_col, "clints_callsRecords_play", False)
    try:
        d = mine(cc)
        assert d["recordings_status"] == "no_key" and all(i["recording"] is None for i in d["items"])
    finally:
        set_perm(c4pg, P4.role_col, "clints_callsRecords_play", True)
    from core.config import Config
    monkeypatch.setattr(Config, "CDR_HOST", "127.0.0.1")
    monkeypatch.setattr(Config, "CDR_CONNECT_TIMEOUT", 1)
    d = mine(cc)
    assert d["recordings_status"] == "unavailable" and all(i["recording"] is None for i in d["items"])


def test_export_line_and_recorded_from_the_cdr(cc, cdr):
    import openpyxl
    r = cc.get("/v2/calls/mine/export", P4.cola, query_string=DAY)
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).worksheets[0]
    head = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = {ws.cell(i, len(head)).value: {h: ws.cell(i, c + 1).value for c, h in enumerate(head)}
            for i in range(2, ws.max_row + 1)}
    assert rows[P4.AT1]["مسجلة"] == "نعم" and rows[P4.AT1]["الرقم"] == "+966561240001"
    assert rows[P4.AT1]["الحالة"] == "انتهت" and rows[P4.AT1]["التحويلة"] == "7301"
