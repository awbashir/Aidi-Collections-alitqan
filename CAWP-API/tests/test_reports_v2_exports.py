"""Phase 4 R0: private asynchronous report exports (API, worker, writers, retention)."""
from __future__ import annotations

import datetime as dt
import os
import uuid
from decimal import Decimal

import pytest

from test_reports_v2_support import *  # noqa: F401,F403
from test_reports_v2_support import (R, artifact_path, export_row, make_case, make_payment, make_portfolio, q, q1,
                                     read_csv, run_export, set_perm)


class Crash(BaseException):
    """Simulates a worker process dying mid-export (not caught by `except Exception`)."""


def _fill(db, n=105, owner=R.col1, pf=None, names=None):
    pf = pf or make_portfolio(db)
    cases = [make_case(db, pf, owner=owner, name=(names or {}).get(i)) for i in range(max(1, min(n, 7)))]
    for i in range(n):
        status = ("approved", "pending", "rejected", "approved", "reversed")[i % 5]
        make_payment(db, cases[i % len(cases)], f"{10 + i}.{i % 100:02d}", status,
                     when=f"2026-03-{1 + i % 28:02d}T21:00:00Z")
    return pf


def _preview_all(rc, report, user, params, sort=None):
    items, first = rc.all_rows(report, user, {**params, **({"sort": sort} if sort else {})})
    return items, first["totals"]


# --------------------------------------------------------------------------- API lifecycle
def test_create_is_idempotent_202_and_creator_private(rc, db):
    pf = make_portfolio(db)
    key = str(uuid.uuid4())
    body = {"format": "csv", "params": {"portfolio": pf, "status": "approved"}, "sort": "amount"}
    r1 = rc.post("/v2/reports/RPT-06/exports", R.sup, body, key=key)
    assert r1.status_code == 202, r1.get_json()
    d = r1.get_json()["data"]
    assert d["status"] == "queued" and d["format"] == "csv" and d["report"]["id"] == "RPT-06"
    assert d["definition"]["params"] == {"portfolio": pf, "status": "approved"}
    assert d["definition"]["sort"] == "amount" and "recorder_id" not in d["definition"]["columns"]
    assert d["actions"] == {"download": False} and d["artifact"] is None
    r2 = rc.post("/v2/reports/RPT-06/exports", R.sup, body, key=key)
    assert r2.status_code == 202 and r2.headers.get("Idempotent-Replayed") == "true"
    assert r2.get_json()["data"]["id"] == d["id"]
    assert q1(db, "SELECT count(*) FROM report_exports WHERE id = %s", (d["id"],))[0] == 1
    assert q1(db, "SELECT count(*) FROM audit_events WHERE action = 'report_export.create' AND entity_id = %s",
              (d["id"],))[0] == 1
    # other users never see it
    assert rc.get(f"/v2/report-exports/{d['id']}", R.super).status_code == 404
    assert rc.download(d["id"], R.super).status_code == 404
    mine = rc.get("/v2/report-exports", R.sup, {"report": "RPT-06"}).get_json()["data"]
    assert d["id"] in {i["id"] for i in mine["items"]}
    theirs = rc.get("/v2/report-exports", R.super).get_json()["data"]
    assert d["id"] not in {i["id"] for i in theirs["items"]}
    # not ready yet
    r = rc.download(d["id"], R.sup)
    assert r.status_code == 409 and r.get_json()["code"] == "export_not_ready"


def test_create_validation(rc):
    def post(body, user=R.sup, report="RPT-06"):
        return rc.post(f"/v2/reports/{report}/exports", user, body)

    r = post({"format": "pdf"})
    assert r.status_code == 422 and r.get_json()["error"] == {"format": "unknown_value"}
    r = post({"params": {"payment_date_from": "2026-02-30"}})
    assert r.status_code == 422 and r.get_json()["error"] == {"payment_date_from": "invalid"}
    r = post({"params": {"bogus": 1}})
    assert r.status_code == 422 and r.get_json()["error"] == {"bogus": "unknown_parameter"}
    r = post({"filters": {}})
    assert r.status_code == 422 and r.get_json()["error"] == {"filters": "unknown_parameter"}
    r = post({"sort": "nope"})
    assert r.status_code == 422 and r.get_json()["error"] == {"sort": "unknown_value"}
    r = post({}, user=R.col1)
    assert r.status_code == 403 and r.get_json()["error"] == {"required_any": ["collections_users"]}
    r = rc.client.post("/v2/reports/RPT-06/exports", json={}, headers=rc.headers(R.sup, key=False))
    assert r.status_code == 400 and r.get_json()["code"] == "idempotency_key_required"
    # list validation
    assert rc.get("/v2/report-exports", R.sup, {"status": "done"}).status_code == 422
    assert rc.get("/v2/report-exports", R.sup, {"x": "1"}).status_code == 422
    # paging params in an export are ignored (an export is always the full result)
    d = rc.export("RPT-06", R.sup, {"page": 3, "page_size": 5})
    assert d["definition"]["params"] == {}


# --------------------------------------------------------------------------- export == preview
def test_full_csv_export_equals_the_full_preview(rc, rep_app, db):
    pf = _fill(db, 110)
    params = {"portfolio": pf}
    items, totals = _preview_all(rc, "RPT-06", R.sup, params, sort="amount")
    assert len(items) == 110
    exp = rc.export("RPT-06", R.sup, params, "csv", sort="amount")
    assert run_export(rep_app, exp["id"]) == "succeeded"
    done = rc.export_get(exp["id"], R.sup)
    assert done["status"] == "succeeded" and done["row_count"] == 110 and done["totals"] == totals
    assert done["progress"] == {"rows_written": 110, "rows_expected": 110, "percent": 100.0}
    assert done["cutoff_at"] and done["artifact"]["available"] and done["actions"]["download"]
    assert [a["outcome"] for a in done["attempts"]] == ["completed"]
    r = rc.download(exp["id"], R.sup)
    assert r.status_code == 200
    assert r.headers["Content-Type"] == "text/csv; charset=utf-8"
    assert r.headers["Cache-Control"] == "private, no-store"
    assert f'filename={done["artifact"]["file_name"]}' in r.headers["Content-Disposition"]
    assert done["artifact"]["file_name"].startswith("RPT-06-") and done["artifact"]["file_name"].endswith(".csv")
    header, rows = read_csv(r.data)
    cols = done["definition"]["columns"]
    assert header == cols
    expected = [["" if i[c] is None else ("true" if i[c] is True else str(i[c])) for c in cols] for i in items]
    assert rows == expected                                  # every cell, in the preview's order
    import hashlib
    assert hashlib.sha256(r.data).hexdigest() == done["artifact"]["sha256"]
    assert q1(db, "SELECT count(*) FROM audit_events WHERE action = 'report_export.download' AND entity_id = %s",
              (exp["id"],))[0] == 1
    assert q1(db, "SELECT count(*) FROM audit_events WHERE action = 'report_export.publish' AND entity_id = %s",
              (exp["id"],))[0] == 1
    # the file is private under the storage root
    path = artifact_path(db, exp["id"])
    assert path.startswith(os.environ["BULK_STORAGE_DIR"]) and "/reports/" in path
    assert oct(os.stat(path).st_mode & 0o777) == "0o600"


def test_export_uses_a_fresh_snapshot(rc, rep_app, db):
    """A later export includes rows added after the preview (documented behavior)."""
    pf = make_portfolio(db)
    case = make_case(db, pf)
    make_payment(db, case, "1.00", "approved")
    preview = rc.rows("RPT-06", R.sup, {"portfolio": pf})
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf})
    make_payment(db, case, "2.00", "approved")
    run_export(rep_app, exp["id"])
    done = rc.export_get(exp["id"], R.sup)
    assert preview["total"] == 1 and done["row_count"] == 2 and done["totals"]["approved_amount"] == "3.00"


def test_batches_of_500_stream_with_progress(rc, rep_app, db, monkeypatch):
    pf = make_portfolio(db)
    cases = [make_case(db, pf) for _ in range(3)]
    q(db, """INSERT INTO "Payments" (clint_id, user_id, collection_type_id, collection_method_id, amount,
                                     payment_date, approval_status)
             SELECT (ARRAY[%s, %s, %s]::uuid[])[1 + g %% 3], %s, '00000000-0000-4000-8000-000000000704',
                    '00000000-0000-4000-8000-000000000705', 1.25, '2026-04-01T21:00:00Z', true
               FROM generate_series(1, 1203) g""", (*cases, R.col1))
    batches = []
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf}, "xlsx")
    assert run_export(rep_app, exp["id"], after_batch=batches.append) == "succeeded"
    assert batches == [500, 1000, 1203]
    done = rc.export_get(exp["id"], R.sup)
    assert done["row_count"] == 1203 and done["totals"]["approved_amount"] == "1503.75"


# --------------------------------------------------------------------------- formula safety
DANGEROUS = ["=HYPERLINK(\"http://x\",\"y\")", "+1+1", "-2+3", "@SUM(A1)", "\tTAB", "\rCR",
             "=cmd|' /C calc'!A0"]


def test_csv_and_xlsx_are_formula_safe(rc, rep_app, db):
    pf = make_portfolio(db)
    for i, bad in enumerate(DANGEROUS):
        case = make_case(db, pf, name=bad, account=f"=ACC{i}-{uuid.uuid4().hex[:6]}")
        make_payment(db, case, "7.00", "approved")
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf}, "csv", sort="customer_name")
    run_export(rep_app, exp["id"])
    header, rows = read_csv(rc.download(exp["id"], R.sup).data)
    name_i, acc_i, amt_i = header.index("customer_name"), header.index("account_number"), header.index("amount")
    names = sorted(r[name_i] for r in rows)
    assert names == sorted("'" + b for b in DANGEROUS)
    assert all(r[acc_i].startswith("'=ACC") for r in rows)
    assert {r[amt_i] for r in rows} == {"7.00"}                 # money is never altered
    # XLSX: typed strings, never formulas, exact text
    x = rc.export("RPT-06", R.sup, {"portfolio": pf}, "xlsx", sort="customer_name")
    run_export(rep_app, x["id"])
    from openpyxl import load_workbook
    wb = load_workbook(artifact_path(db, x["id"]))
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    ci = hdr.index("Customer") + 1
    cells = [ws.cell(row=r, column=ci) for r in range(2, ws.max_row + 1)]
    assert all(c.data_type == "s" for c in cells)
    assert sorted(c.value for c in cells) == sorted(DANGEROUS)          # exact text, CR included
    import zipfile
    xml = zipfile.ZipFile(artifact_path(db, x["id"])).read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in xml                                      # no formula anywhere


# --------------------------------------------------------------------------- XLSX + Arabic RTL XLSX
def test_xlsx_values_are_typed_and_equal_the_preview(rc, rep_app, db):
    pf = _fill(db, 12)
    items, totals = _preview_all(rc, "RPT-06", R.sup, {"portfolio": pf}, sort="amount")
    x = rc.export("RPT-06", R.sup, {"portfolio": pf}, "xlsx", sort="amount")
    run_export(rep_app, x["id"])
    r = rc.download(x["id"], R.sup)
    assert r.status_code == 200
    assert r.headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    from openpyxl import load_workbook
    path = artifact_path(db, x["id"])
    wb = load_workbook(path)
    assert wb.sheetnames == ["Report", "Summary"]
    ws = wb["Report"]
    assert ws.sheet_view.rightToLeft in (None, False) and ws.freeze_panes == "A2"
    hdr = [c.value for c in ws[1]]
    assert hdr[:4] == ["Payment ID", "Status", "Amount", "Payment date"]
    data = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert len(data) == 12
    for item, row in zip(items, data):
        assert row[0] == item["payment_id"]
        assert row[1] == {"pending": "Pending approval", "approved": "Approved", "rejected": "Rejected",
                          "reversed": "Reversed"}[item["status"]]
        assert Decimal(str(row[2])).quantize(Decimal("0.01")) == Decimal(item["amount"])
        assert row[3].date().isoformat() == item["payment_date"]
    amount_cell = ws.cell(row=2, column=3)
    assert amount_cell.data_type == "n" and amount_cell.number_format == "#,##0.00"
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows()}
    assert summary["Rows"] == 12
    assert Decimal(str(summary["Approved recovery"])).quantize(Decimal("0.01")) == Decimal(totals["approved_amount"])


def test_arabic_xlsx_is_rtl_with_arabic_headers_and_values(rc, rep_app, db):
    pf = make_portfolio(db)
    case = make_case(db, pf, name="عميل اختبار", amount=5000)
    make_payment(db, case, "1234.50", "approved", when="2026-03-31T21:00:00Z")    # civil 2026-04-01
    make_payment(db, case, "10.00", None)
    make_payment(db, case, "20.00", "reversed")
    x = rc.export("RPT-06", R.sup, {"portfolio": pf}, "xlsx_ar", sort="-amount")
    run_export(rep_app, x["id"])
    done = rc.export_get(x["id"], R.sup)
    assert done["artifact"]["file_name"].endswith("-ar.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(artifact_path(db, x["id"]))
    assert wb.sheetnames == ["تحصيلات المستخدمين", "الملخص"]
    for ws in wb.worksheets:
        assert ws.sheet_view.rightToLeft is True
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    assert hdr[:5] == ["رقم الدفعة", "الحالة", "المبلغ", "تاريخ الدفع", "تاريخ الادخال"]
    assert "اسم العميل" in hdr and "مبلغ المديونية" in hdr and "قام بالتعميد" in hdr
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    status = [r[1] for r in rows]
    assert status == ["معتمدة", "معكوسة", "بانتظار الاعتماد"]
    assert rows[0][2] == 1234.5 and ws.cell(row=2, column=3).number_format == "#,##0.00"
    assert rows[0][3] == dt.datetime(2026, 4, 1)
    assert rows[0][hdr.index("اسم العميل")] == "عميل اختبار"
    summary = {r[0].value: r[1].value for r in wb.worksheets[1].iter_rows()}
    assert summary["عدد السجلات"] == 3
    assert summary["المحصل المعتمد"] == 1234.5 and summary["المعكوس (ليس تحصيلاً)"] == 20
    assert summary["التقرير"].startswith("RPT-06")
    assert "المحفظة" in summary                                   # the applied filter, Arabic label
    # RPT-07/08 also offer it
    assert "xlsx_ar" in rc.get("/v2/reports/RPT-08", R.sup).get_json()["data"]["formats"]


# --------------------------------------------------------------------------- permission rechecks
def test_permission_lost_mid_export_publishes_nothing(rc, rep_app, db):
    pf = _fill(db, 1100, owner=R.col1)
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf})

    def revoke(done):
        if done == 500:
            set_perm(db, R.role_sup, "collections_users", False)

    try:
        assert run_export(rep_app, exp["id"], after_batch=revoke) == "failed"
    finally:
        set_perm(db, R.role_sup, "collections_users", True)
    row = export_row(db, exp["id"])
    assert row["status"] == "failed" and row["error_code"] == "permission_denied"
    assert row["stored_path"] is None and row["row_count"] is None
    assert q(db, "SELECT outcome FROM report_export_attempts WHERE export_id = %s", (exp["id"],)) == [("failed",)]
    r = rc.download(exp["id"], R.sup)
    assert r.status_code == 409 and r.get_json()["code"] == "export_failed"
    assert not [f for f in os.listdir(os.path.join(os.environ["BULK_STORAGE_DIR"], "tmp"))
                if f.endswith(".part") and os.path.getmtime(os.path.join(os.environ["BULK_STORAGE_DIR"], "tmp", f))
                > dt.datetime.now().timestamp() - 5]


def test_permission_missing_at_start_and_at_download(rc, rep_app, db):
    pf = _fill(db, 3)
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf})
    set_perm(db, R.role_sup, "collections_users", False)
    try:
        assert run_export(rep_app, exp["id"]) == "failed"
    finally:
        set_perm(db, R.role_sup, "collections_users", True)
    assert export_row(db, exp["id"])["error_code"] == "permission_denied"
    # a finished export: the download rechecks the key
    ok = rc.export("RPT-06", R.sup, {"portfolio": pf})
    run_export(rep_app, ok["id"])
    set_perm(db, R.role_sup, "collections_users", False)
    try:
        r = rc.download(ok["id"], R.sup)
        assert r.status_code == 403 and r.get_json()["code"] == "forbidden"
    finally:
        set_perm(db, R.role_sup, "collections_users", True)
    assert rc.download(ok["id"], R.sup).status_code == 200


def test_export_scope_is_the_requesters_current_scope(rc, rep_app, db):
    """The worker re-resolves the scope: a case moved out of the supervisor's contract before
    generation is not exported (RPT-06); RPT-07 follows the recorder, not the scope."""
    pf = make_portfolio(db, R.contract_a)
    other = make_portfolio(db, R.contract_b)
    stays = make_case(db, pf, owner=R.col1)
    moved = make_case(db, pf, owner=R.col1)
    make_payment(db, stays, "1.00", "approved", recorder=R.col1)
    make_payment(db, moved, "2.00", "approved", recorder=R.col1)
    both = {"portfolio": f"{pf},{other}"}
    exp = rc.export("RPT-06", R.sup, both)
    mine = rc.export("RPT-07", R.col1, both)
    q(db, "UPDATE clints SET portfolios_id = %s WHERE id = %s", (other, moved))
    run_export(rep_app, exp["id"])
    run_export(rep_app, mine["id"])
    done = rc.export_get(exp["id"], R.sup)
    assert done["row_count"] == 1 and done["totals"]["approved_amount"] == "1.00"
    own = rc.export_get(mine["id"], R.col1)
    assert own["row_count"] == 2 and own["totals"]["approved_amount"] == "3.00"


# --------------------------------------------------------------------------- crash / lease / stop
def test_worker_crash_then_lease_recovery_completes_once(rc, rep_app, db):
    pf = _fill(db, 1050)
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf})

    def crash(done):
        if done == 500:
            raise Crash("worker killed")

    with pytest.raises(Crash):
        run_export(rep_app, exp["id"], worker="w-crash", after_batch=crash)
    row = export_row(db, exp["id"])
    assert row["status"] == "running" and row["locked_by"] == "w-crash" and row["rows_written"] == 500
    # the lease is still live: nobody else may claim it
    with rep_app.app_context():
        from core.reports import exports
        assert exports.claim("w-other", export_id=exp["id"]) is None
    q(db, "UPDATE report_exports SET lease_until = now() - interval '1 second' WHERE id = %s", (exp["id"],))
    assert run_export(rep_app, exp["id"], worker="w-two") == "succeeded"
    row = export_row(db, exp["id"])
    assert row["status"] == "succeeded" and row["row_count"] == 1050 and row["failed_attempts"] == 1
    assert row["attempt"] == 2
    assert q(db, "SELECT attempt, outcome, worker_id, rows_written FROM report_export_attempts WHERE export_id = %s "
                 "ORDER BY attempt", (exp["id"],)) == [(1, "lease_expired", "w-crash", 500), (2, "completed", "w-two", 1050)]
    header, rows = read_csv(rc.download(exp["id"], R.sup).data)
    assert len(rows) == 1050 and len({r[0] for r in rows}) == 1050


def test_max_attempts_then_failed(rc, rep_app, db):
    pf = _fill(db, 2)
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf})
    q(db, "UPDATE report_exports SET max_attempts = 2 WHERE id = %s", (exp["id"],))
    for n in (1, 2):
        with rep_app.app_context():
            from core.reports import exports
            from extensions import db as _db
            exports.recover_expired()
            got = exports.claim(f"w{n}", export_id=exp["id"])
            _db.session.remove()
        assert got is not None
        q(db, "UPDATE report_exports SET lease_until = now() - interval '1 second' WHERE id = %s", (exp["id"],))
    with rep_app.app_context():
        from core.reports import exports
        exports.recover_expired()
    row = export_row(db, exp["id"])
    assert row["status"] == "failed" and row["error_code"] == "max_attempts_exceeded" and row["failed_attempts"] == 2


def test_lost_lease_publishes_nothing(rc, rep_app, db):
    pf = _fill(db, 600)
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf})

    def steal(done):
        if done == 500:     # another worker recovered the export after our lease expired
            q(db, "UPDATE report_exports SET locked_by = 'thief' WHERE id = %s", (exp["id"],))

    assert run_export(rep_app, exp["id"], worker="victim", after_batch=steal) == "lease_lost"
    row = export_row(db, exp["id"])
    assert row["locked_by"] == "thief" and row["stored_path"] is None and row["status"] == "running"


def test_graceful_stop_releases_as_retrying_without_counting(rc, rep_app, db):
    pf = _fill(db, 700)
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf})
    stop = {"now": False}

    def after(done):
        stop["now"] = True

    assert run_export(rep_app, exp["id"], should_stop=lambda: stop["now"], after_batch=after) == "retrying"
    row = export_row(db, exp["id"])
    assert row["status"] == "retrying" and row["failed_attempts"] == 0 and row["locked_by"] is None
    assert run_export(rep_app, exp["id"]) == "succeeded"
    assert [o for (o,) in q(db, "SELECT outcome FROM report_export_attempts WHERE export_id = %s ORDER BY attempt",
                            (exp["id"],))] == ["interrupted", "completed"]


def test_too_many_rows_and_unexpected_error(rc, rep_app, db, monkeypatch):
    pf = _fill(db, 5)
    monkeypatch.setenv("REPORT_EXPORT_MAX_ROWS", "4")
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf})
    assert run_export(rep_app, exp["id"]) == "failed"
    row = export_row(db, exp["id"])
    assert row["error_code"] == "too_many_rows" and "narrow the filters" in row["error_message"]
    monkeypatch.delenv("REPORT_EXPORT_MAX_ROWS")
    # an unexpected error is retried, then fails at max_attempts
    from core.reports import writers
    boom = rc.export("RPT-06", R.sup, {"portfolio": pf})
    q(db, "UPDATE report_exports SET max_attempts = 2 WHERE id = %s", (boom["id"],))
    monkeypatch.setattr(writers.CsvWriter, "write_rows", lambda self, rows: (_ for _ in ()).throw(RuntimeError("disk")))
    assert run_export(rep_app, boom["id"]) == "retrying"
    assert run_export(rep_app, boom["id"]) == "failed"
    row = export_row(db, boom["id"])
    assert row["error_code"] == "internal_error" and "disk" in row["error_message"] and row["failed_attempts"] == 2


# --------------------------------------------------------------------------- expiry / retention
def test_expiry_and_retention(rc, rep_app, db):
    pf = _fill(db, 3)
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf})
    run_export(rep_app, exp["id"])
    path = artifact_path(db, exp["id"])
    assert os.path.isfile(path)
    ttl = export_row(db, exp["id"])
    assert dt.timedelta(days=6, hours=23) < ttl["expires_at"] - ttl["finished_at"] <= dt.timedelta(days=7, seconds=5)
    q(db, "UPDATE report_exports SET expires_at = requested_at + interval '1 millisecond' WHERE id = %s", (exp["id"],))
    # expired before retention runs: shown as expired, download refused
    d = rc.export_get(exp["id"], R.sup)
    assert d["status"] == "expired" and d["stored_status"] == "succeeded" and d["actions"]["download"] is False
    r = rc.download(exp["id"], R.sup)
    assert r.status_code == 410 and r.get_json()["code"] == "export_expired"
    listed = rc.get("/v2/report-exports", R.sup, {"status": "expired"}).get_json()["data"]
    assert exp["id"] in {i["id"] for i in listed["items"]}
    with rep_app.app_context():
        from core.bulk import retention
        report = retention.run()
    assert report["purged_report_exports"] >= 1
    assert not os.path.exists(path)
    row = export_row(db, exp["id"])
    assert row["status"] == "expired" and row["purged_at"] is not None and row["row_count"] == 3
    assert q1(db, "SELECT count(*) FROM audit_events WHERE action = 'report_export.purge' AND entity_id = %s",
              (exp["id"],))[0] == 1
    with rep_app.app_context():
        from core.bulk import retention
        assert retention.run()["purged_report_exports"] == 0


# --------------------------------------------------------------------------- the worker loop
def test_worker_loop_claims_exports_and_alternates(rc, rep_app, db, rep_storage, monkeypatch):
    monkeypatch.setenv("BULK_WORKER_HEALTH_FILE", str(rep_storage / "worker.alive"))
    pf = _fill(db, 4)
    # finish any exports left queued by other tests first, so this one is next in line
    with rep_app.app_context():
        from core.reports import exports
        from extensions import db as _db
        while True:
            left = exports.claim("drain")
            if left is None:
                break
            exports.process(left, exports.Control(worker_id="drain"))
        _db.session.remove()
    exp = rc.export("RPT-06", R.sup, {"portfolio": pf}, "xlsx_ar")
    from worker.main import Worker
    w = Worker(rep_app, worker_id="loop-test")
    w.exports_first = True               # the turn after a bulk job: the export queue is tried first
    with rep_app.app_context():
        kind, work = w.next_work()
    assert kind == "export" and str(work["id"]) == exp["id"]
    assert w.run_export(work) == "succeeded"               # real heartbeat thread
    assert export_row(db, exp["id"])["status"] == "succeeded"
    assert w.current_job is None
    with rep_app.app_context():
        assert w.claim_export() is None                    # nothing left
