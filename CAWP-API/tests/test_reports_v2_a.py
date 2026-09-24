"""Phase 4 R-A: RPT-01 (users roster), RPT-02 (unarchived cases), RPT-03 (archived cases), RPT-04 (team
note activity) and RPT-05 (my notes) on the v2 report framework.

Covered: definitions and screen keys, Riyadh date boundaries, scope (case scope and the RPT-01 employee
scope), the recorder rule (RPT-05), the latest matching note across the full filtered set (RPT-04), exact
totals, paging stability, full export = full preview, the Arabic XLSX, formula safety, protected employee
fields, selectors, and the documented differences from the legacy endpoints (synthetic rows).

Every test creates its own portfolio/branch/users and filters on them, so the shared test database's
other rows never reach an assertion.
"""
from __future__ import annotations

import datetime as dt
import math
import os
import uuid
from decimal import Decimal
from pathlib import Path

import pytest

from conftest import MAIN_DB, run_sql_file
from test_reports_v2_support import *  # noqa: F401,F403
from test_reports_v2_support import (artifact_path, make_case, make_contract, make_payment, make_portfolio, q,
                                     read_csv, run_export)

SEED_A = Path(__file__).resolve().parent / "fixtures" / "seed_reports_a.sql"
NATIONALITY = "00000000-0000-4000-8000-000000000101"


class A:
    super = "ea0ea0ea-0000-4000-8000-000000000501"
    sup = "ea0ea0ea-0000-4000-8000-000000000502"
    col1 = "ea0ea0ea-0000-4000-8000-000000000503"
    col2 = "ea0ea0ea-0000-4000-8000-000000000504"
    col3 = "ea0ea0ea-0000-4000-8000-000000000505"
    col4 = "ea0ea0ea-0000-4000-8000-000000000506"
    tgt = "ea0ea0ea-0000-4000-8000-000000000507"
    idadmin = "ea0ea0ea-0000-4000-8000-000000000508"
    none = "ea0ea0ea-0000-4000-8000-000000000509"
    role_super = "ea0ea0ea-0000-4000-8000-000000000201"
    role_sup = "ea0ea0ea-0000-4000-8000-000000000202"
    role_col = "ea0ea0ea-0000-4000-8000-000000000203"
    role_tgt = "ea0ea0ea-0000-4000-8000-000000000204"
    role_none = "ea0ea0ea-0000-4000-8000-000000000205"
    role_idadmin = "ea0ea0ea-0000-4000-8000-000000000206"
    branch1 = "ea0ea0ea-0000-4000-8000-000000000111"
    branch2 = "ea0ea0ea-0000-4000-8000-000000000112"
    contract_a = "ea0ea0ea-0000-4000-8000-000000000301"
    contract_b = "ea0ea0ea-0000-4000-8000-000000000302"
    supervisor_row = "ea0ea0ea-0000-4000-8000-000000000601"
    st_promise = "ea0ea0ea-0000-4000-8000-000000000701"
    st_noanswer = "ea0ea0ea-0000-4000-8000-000000000702"
    reach_yes = "ea0ea0ea-0000-4000-8000-000000000711"
    reach_no = "ea0ea0ea-0000-4000-8000-000000000712"
    method = "ea0ea0ea-0000-4000-8000-000000000721"
    reason1 = "ea0ea0ea-0000-4000-8000-000000000731"
    reason2 = "ea0ea0ea-0000-4000-8000-000000000732"


REPORT_KEYS = {"RPT-01": "users", "RPT-02": "clints", "RPT-03": "clints_archive",
               "RPT-04": "report_users_comments", "RPT-05": "report_collectors_currentUser_comments"}
MAY = {"created_from": "2026-05-01", "created_to": "2026-05-31"}


@pytest.fixture(scope="module", autouse=True)
def seed_a(rep_db):
    run_sql_file(MAIN_DB, SEED_A)
    from core import lookups
    lookups.invalidate()


# --------------------------------------------------------------------------- helpers
def invalidate():
    from core import lookups
    lookups.invalidate()


def make_branch(db, name=None) -> str:
    bid = str(uuid.uuid4())
    q(db, "INSERT INTO branches (id, name) VALUES (%s, %s)", (bid, name or "RA branch " + bid[:8]))
    return bid


def make_user(db, branch, role=A.role_col, contract=None, active=True, first="RA", last=None, **extra) -> str:
    uid = str(uuid.uuid4())
    q(db, """INSERT INTO users (id, email, frist_name, last_name, mobile, password, user_name, role_id,
                                nationality_id, branch_id, contract_id, "identityNumber", is_active,
                                target_is_percentage, target_amount, target_percent)
             VALUES (%s, %s, %s, %s, %s, 'not-a-real-hash', %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
      (uid, f"u{uid[:8]}@test.invalid", first, last or uid[:8], "05" + uid[:8].replace("-", "0")[:8],
       "ra." + uid[:8], role, NATIONALITY, branch, contract, "2" + uid[:9].replace("-", "0"), active,
       extra.get("pct", False), extra.get("target_amount", 0), extra.get("target_percent", 0)))
    invalidate()                      # the role/permission maps are cached for 60 s
    return uid


def supervise_user(db, supervisor_row, user):
    q(db, "INSERT INTO supervisor_users (supervisor_id, user_id) VALUES (%s, %s)", (supervisor_row, user))
    invalidate()


def supervise_contract(db, supervisor_row, contract):
    q(db, "INSERT INTO supervisor_contracts (supervisor_id, contract_id) VALUES (%s, %s)", (supervisor_row, contract))
    invalidate()


def make_supervisor(db, user) -> str:
    sid = str(uuid.uuid4())
    q(db, "INSERT INTO supervisors (id, user_id) VALUES (%s, %s)", (sid, user))
    invalidate()
    return sid


def make_note(db, case, recorder, when, status=A.st_promise, text=None, reach=A.reach_yes, callback=None,
              contact=None) -> str:
    nid = str(uuid.uuid4())
    q(db, """INSERT INTO "Comments" (id, clint_id, user_id, "contactMethod_id", comment, comment_status_id,
                                     comment_reachability_id, next_contact_date, create_at, "contactNumber")
             VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)""",
      (nid, case, recorder, A.method, text or "note " + nid[:8], status, reach, callback, when, contact))
    return nid


def retire(db, note):
    q(db, 'UPDATE "Comments" SET is_deleted = true WHERE id = %s', (note,))


def assign(db, case, user, when, deleted=False):
    q(db, """INSERT INTO "SepcifyUsers" (user_id, clint_id, create_at, is_deleted) VALUES (%s, %s, %s, %s)""",
      (user, case, when, deleted))
    q(db, "UPDATE clints SET user_id = %s WHERE id = %s", (user, case))


def archive(db, case, by=A.sup, reason=A.reason1, when="2026-06-10T09:00:00Z") -> str:
    aid = str(uuid.uuid4())
    q(db, """INSERT INTO "clintArchive" (id, clint_id, user_id, archive_reason_id, create_at)
             VALUES (%s, %s, %s, %s, %s)""", (aid, case, by, reason, when))
    q(db, 'UPDATE clints SET is_archive = true, "clintArchive_id" = %s WHERE id = %s', (aid, case))
    return aid


def restore(db, case):
    q(db, 'UPDATE clints SET is_archive = false, "clintArchive_id" = NULL WHERE id = %s', (case,))


def by_key(items, key):
    return {i[key]: i for i in items}


def keys(d, key):
    return {i[key] for i in d["items"]}


def csv_cell(col_type, v):
    from core.bulk.artifacts import safe_cell
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if col_type in ("text", "enum", "json"):
        return safe_cell(v)
    return str(v)


def assert_export_equals_preview(rc, rep_app, report, user, params, sort=None):
    """The full CSV export equals every page of the preview, cell by cell, and the totals agree."""
    items, first = rc.all_rows(report, user, {**params, **({"sort": sort} if sort else {})}, page_size=37)
    definition = rc.get(f"/v2/reports/{report}", user).get_json()["data"]
    cols = [(c["key"], c["type"]) for c in definition["columns"] if c["export"]]
    exp = rc.export(report, user, params, "csv", sort=sort)
    assert run_export(rep_app, exp["id"]) == "succeeded"
    done = rc.export_get(exp["id"], user)
    header, rows = read_csv(rc.download(exp["id"], user).data)
    assert header == [k for k, _ in cols]
    assert len(rows) == len(items) == first["total"] == done["row_count"]
    for item, row in zip(items, rows):
        assert row == [csv_cell(t, item[k]) for k, t in cols]
    assert done["totals"] == first["totals"]
    return items, first, done


# =========================================================================== definitions / keys
def test_definitions_keys_formats_and_filters(rc):
    lst = rc.get("/v2/reports", A.super).get_json()["data"]
    mine = {i["id"]: i for i in lst["items"] if i["id"] in REPORT_KEYS}
    assert set(mine) == set(REPORT_KEYS)
    for rid, key in REPORT_KEYS.items():
        assert mine[rid]["permissions"] == [key]
        assert mine[rid]["formats"] == ["csv", "xlsx", "xlsx_ar"]          # every legacy screen exported Arabic
    fams = {f["family"]: f["reports"] for f in lst["families"]}
    assert fams["users"] == ["RPT-01"] and fams["cases"] == ["RPT-02", "RPT-03"]
    assert fams["notes"] == ["RPT-04", "RPT-05"]

    d1 = rc.get("/v2/reports/rpt-01", A.super).get_json()["data"]
    prot = {c["key"]: c["protected"] for c in d1["columns"] if c["protected"]}
    assert prot == {"email": "identity", "mobile": "identity", "identity_number": "identity",
                    "target_is_percentage": "target", "target_amount": "target", "target_percent": "target"}
    assert d1["legacy"]["key"] == "users" and d1["default_sort"] == "created_at"
    sorts = {s["key"] for s in d1["sorts"]}
    assert not sorts & set(prot)                          # a hidden value can never order the rows
    q_filter = next(f for f in d1["filters"] if f["key"] == "q")
    assert q_filter["type"] == "text"
    assert {t["key"] for t in d1["totals"]} == {"count", "active_count", "inactive_count", "with_cases_count",
                                                "case_count", "original_debt"}

    d2 = rc.get("/v2/reports/RPT-02", A.super).get_json()["data"]
    d3 = rc.get("/v2/reports/RPT-03", A.super).get_json()["data"]
    c2, c3 = [c["key"] for c in d2["columns"]], [c["key"] for c in d3["columns"]]
    assert {"archived_by_name", "archive_reason", "archived_at"} <= set(c3) and "archived_at" not in c2
    assert {"last_note_at", "last_note_status", "assignment_days", "outstanding_amount", "original_debt"} <= set(c2)
    assert {f["key"] for f in d3["filters"]} - {f["key"] for f in d2["filters"]} == \
        {"archived", "archived_by", "archive_reason"}

    d4 = rc.get("/v2/reports/RPT-04", A.super).get_json()["data"]
    d5 = rc.get("/v2/reports/RPT-05", A.super).get_json()["data"]
    f4 = {f["key"]: f for f in d4["filters"]}
    f5 = {f["key"]: f for f in d5["filters"]}
    assert f4["created"]["required"] and f4["created"]["max_days"] == 366
    assert f4["created"]["semantics"] == "riyadh_day_bounds" and f4["created"]["timezone"] == "Asia/Riyadh"
    assert f4["note_selection"]["choices"] == [{"value": "latest_per_case", "labels": {
        "en": "Latest matching note per case", "ar": "آخر متابعة مطابقة لكل عميل"}}]
    assert "note_selection" not in f5 and "recorder" not in f5 and "recorder_branch" not in f5
    assert f5["created"]["required"] and f5["created"]["max_days"] is None
    money4 = [c["key"] for c in d4["columns"] if c["money"]]
    assert money4 == ["original_debt"]
    assert "original_debt" not in {t["key"] for t in d4["totals"] + d5["totals"]}       # never summed
    assert {c["role"] for c in d4["columns"] if c["role"]} == {"recorder", "owner"}


def test_screen_keys_gate_every_report(rc):
    visible = lambda u: {i["id"] for i in rc.get("/v2/reports", u).get_json()["data"]["items"]} & set(REPORT_KEYS)
    assert visible(A.super) == visible(A.sup) == set(REPORT_KEYS)
    assert visible(A.col1) == {"RPT-02", "RPT-05"}                  # the collector role on cawp-db: clints + my notes
    assert visible(A.tgt) == visible(A.idadmin) == {"RPT-01"}
    assert visible(A.none) == set()
    for rid, key in REPORT_KEYS.items():
        for path in (f"/v2/reports/{rid}", f"/v2/reports/{rid}/rows"):
            r = rc.get(path, A.none, MAY if rid in ("RPT-04", "RPT-05") else None)
            assert r.status_code == 403 and r.get_json()["error"] == {"required_any": [key]}
        r = rc.post(f"/v2/reports/{rid}/exports", A.none, {"format": "csv"})
        assert r.status_code == 403
    for rid in ("RPT-01", "RPT-03", "RPT-04"):
        assert rc.get(f"/v2/reports/{rid}/rows", A.col1, MAY).status_code == 403


def test_validation(rc):
    r = rc.rows("RPT-04", A.sup, {}, expect=422)
    assert r["error"] == {"created_from": "required"}
    r = rc.rows("RPT-04", A.sup, {"created_from": "2025-01-01", "created_to": "2026-01-02"}, expect=422)
    assert r["error"] == {"created_to": "range_too_long"}
    r = rc.rows("RPT-05", A.col1, {"created_from": "2025-01-01"}, expect=422)
    assert r["error"] == {"created_to": "required"}
    assert rc.rows("RPT-05", A.col1, {"created_from": "2024-01-01", "created_to": "2026-12-31"})["total"] >= 0
    r = rc.rows("RPT-04", A.sup, {**MAY, "note_selection": "latest_per_case,all"}, expect=422)
    assert r["error"] == {"note_selection": "unknown_value"}
    r = rc.rows("RPT-05", A.col1, {**MAY, "recorder": A.sup, "note_selection": "latest_per_case"}, expect=422)
    assert r["error"] == {"recorder": "unknown_parameter", "note_selection": "unknown_parameter"}
    r = rc.rows("RPT-01", A.super, {"case_state": "archived,unarchived"}, expect=422)
    assert r["error"] == {"case_state": "too_many"}
    r = rc.rows("RPT-01", A.super, {"sort": "email"}, expect=422)          # protected fields are not sortable
    assert r["error"] == {"sort": "unknown_value"}
    r = rc.rows("RPT-02", A.super, {"debt_from": "10.005", "archived_from": "2026-01-01"}, expect=422)
    assert r["error"] == {"debt_from": "too_precise", "archived_from": "unknown_parameter"}


# =========================================================================== RPT-01
def test_rpt01_employee_scope(rc, db):
    br = make_branch(db)
    k = make_contract(db, "RA roster contract K")
    other = make_contract(db, "RA roster contract L")
    supervise_contract(db, A.supervisor_row, k)
    linked = make_user(db, br, last="linked")
    supervise_user(db, A.supervisor_row, linked)
    by_contract = make_user(db, br, contract=k, last="bycontract", active=False)      # no case at all
    owner = make_user(db, br, last="owner")
    outside = make_user(db, br, contract=other, last="outside")
    historical = make_user(db, br, last="historical")
    pk, pl = make_portfolio(db, k), make_portfolio(db, other)
    make_case(db, pk, owner=owner)
    make_case(db, pl, owner=outside)
    moved = make_case(db, pk, owner=historical)
    make_note(db, moved, historical, "2026-05-02T10:00:00Z")
    assign(db, moved, owner, "2026-05-03T10:00:00Z")                  # historical owner + recorder: not a link
    names = lambda d: {i["user_id"] for i in d["items"]}

    d = rc.rows("RPT-01", A.sup, {"branch": br})
    assert names(d) == {linked, by_contract, owner}
    assert d["totals"]["count"] == 3 and d["totals"]["inactive_count"] == 1
    # identity administrators (users_edit / users_details) and SuperAdmin see every employee
    everyone = {linked, by_contract, owner, outside, historical}
    assert names(rc.rows("RPT-01", A.idadmin, {"branch": br})) == everyone
    assert names(rc.rows("RPT-01", A.super, {"branch": br})) == everyone
    # a key holder with neither supervision nor identity administration: themselves only
    lone_branch = make_branch(db)
    lone = make_user(db, lone_branch, role=A.role_tgt, last="lone")
    assert names(rc.rows("RPT-01", lone, {})) == {lone}
    # the selector candidates follow the same scope
    opts = rc.get("/v2/report-options/branch", A.sup, {"report": "RPT-01", "selected": f"{br},{lone_branch}"})
    sel = {i["id"] for i in opts.get_json()["data"]["selected"]}
    assert sel == {br}


def test_rpt01_assignment_totals_follow_the_case_scope_and_case_state(rc, db):
    br = make_branch(db)
    k = make_contract(db)
    supervise_contract(db, A.supervisor_row, k)
    worker = make_user(db, br, contract=k, last="worker")
    pk = make_portfolio(db, k)
    pb = make_portfolio(db, make_contract(db))                          # not supervised by ra.sup
    make_case(db, pk, owner=worker, amount="100.10")
    make_case(db, pk, owner=worker, amount="200.20")
    archived = make_case(db, pk, owner=worker, amount="50.05")
    archive(db, archived)
    make_case(db, pb, owner=worker, amount="1000.01")                   # outside ra.sup's case scope

    def row(user, **params):
        d = rc.rows("RPT-01", user, {"branch": br, **params})
        return by_key(d["items"], "user_id")[worker], d["totals"]

    r, t = row(A.super)
    assert (r["case_count"], r["original_debt"]) == (4, "1350.36")
    assert (r["unarchived_case_count"], r["archived_case_count"]) == (3, 1)
    assert t["case_count"] == 4 and t["original_debt"] == "1350.36" and t["with_cases_count"] == 1
    r, t = row(A.sup)
    assert (r["case_count"], r["original_debt"]) == (3, "350.35")
    assert (r["unarchived_case_count"], r["archived_case_count"]) == (2, 1)
    r, t = row(A.sup, case_state="unarchived")
    assert (r["case_count"], r["original_debt"]) == (2, "300.30") and t["original_debt"] == "300.30"
    r, t = row(A.sup, case_state="archived")
    assert (r["case_count"], r["original_debt"]) == (1, "50.05") and t["case_count"] == 1
    r, t = row(A.super, case_state="all")
    assert r["case_count"] == 4
    # sort by assigned cases, then the id tie-breaker
    idle = make_user(db, br, contract=k, last="idle")
    d = rc.rows("RPT-01", A.sup, {"branch": br, "sort": "-case_count"})
    assert [i["user_id"] for i in d["items"]] == [worker, idle]
    assert d["totals"]["with_cases_count"] == 1


def test_rpt01_protected_fields_follow_user_privacy(rc, rep_app, db):
    br = make_branch(db)
    k = make_contract(db)
    supervise_contract(db, A.supervisor_row, k)
    staff = make_user(db, br, contract=k, last="staff", target_amount="1500.50")
    pct = make_user(db, br, contract=k, last="pct", pct=True, target_percent="12.5")
    # a target editor who supervises the two (targets visible, identity not)
    editor = make_user(db, br, role=A.role_tgt, last="editor")
    sid = make_supervisor(db, editor)
    supervise_user(db, sid, staff)
    supervise_user(db, sid, pct)
    # ra.sup is in the roster of its own branch: its own row is never redacted
    own = by_key(rc.rows("RPT-01", A.sup, {"branch": A.branch1})["items"], "user_id")[A.sup]
    assert own["email"] == "ra.sup@test.invalid" and own["mobile"] == "0500006002"
    assert own["identity_number"] == "2000006002" and own["target_amount"] == "90000.25"

    rows = lambda u: by_key(rc.rows("RPT-01", u, {"branch": br})["items"], "user_id")
    plain = rows(A.sup)                                           # no users_details / users_edit / target key
    assert editor not in plain                                     # not linked to ra.sup
    for u in (staff, pct):
        r = plain[u]
        assert r["email"] is r["mobile"] is r["identity_number"] is None
        assert r["target_amount"] is r["target_percent"] is r["target_is_percentage"] is None
        assert r["name"] and r["user_name"] and r["role"] and r["status"]          # not protected
    tg = rows(editor)
    assert tg[staff]["target_amount"] == "1500.50" and tg[staff]["target_is_percentage"] is False
    assert tg[pct]["target_percent"] == "12.50" and tg[pct]["target_is_percentage"] is True
    assert tg[staff]["email"] is None and tg[staff]["identity_number"] is None
    assert tg[editor]["email"] is not None                                       # own row
    full = rows(A.idadmin)
    assert full[staff]["email"].endswith("@test.invalid") and full[staff]["mobile"]
    assert full[staff]["target_amount"] == "1500.50"
    # hidden values are not searchable: q covers name, login and extension only
    email = full[staff]["email"]
    assert rc.rows("RPT-01", A.sup, {"branch": br, "q": email})["total"] == 0
    assert rc.rows("RPT-01", A.idadmin, {"branch": br, "q": email})["total"] == 0
    assert rc.rows("RPT-01", A.sup, {"branch": br, "q": "staff"})["total"] == 1
    # the export redacts exactly like the preview (the worker reads the rights from the database)
    items, first, done = assert_export_equals_preview(rc, rep_app, "RPT-01", A.sup, {"branch": br}, sort="name")
    assert all(i["email"] is None for i in items)


# =========================================================================== RPT-02 / RPT-03
def test_rpt02_scope_unarchived_notes_assignment_and_money(rc, db):
    pa = make_portfolio(db, A.contract_a)
    pb = make_portfolio(db, A.contract_b)
    c1 = make_case(db, pa, owner=A.col1, amount="1000.00")
    c2 = make_case(db, pb, owner=A.col2, amount="500.55")             # supervised user, contract B: in scope
    c3 = make_case(db, pb, owner=A.col3, amount="700.00")             # contract B, not linked: out of scope
    un = make_case(db, pa, owner=None, amount="10.00")               # unassigned, contract A: in scope
    gone = make_case(db, pa, owner=A.col1, amount="99.00")
    archive(db, gone)
    params = {"contract": f"{A.contract_a},{A.contract_b}", "portfolio": f"{pa},{pb}"}
    d = rc.rows("RPT-02", A.sup, params)
    assert keys(d, "case_id") == {c1, c2, un}
    assert d["totals"]["unassigned_count"] == 1 and d["totals"]["original_debt"] == "1510.55"
    assert keys(rc.rows("RPT-02", A.col1, params), "case_id") == {c1}
    assert keys(rc.rows("RPT-02", A.super, params), "case_id") == {c1, c2, c3, un}
    # owner filters only narrow
    assert rc.rows("RPT-02", A.sup, {**params, "owner": A.col3})["total"] == 0
    assert keys(rc.rows("RPT-02", A.sup, {**params, "assignment": "unassigned"}), "case_id") == {un}

    # latest effective note (a retired note moves it back), status and reachability of that note
    n1 = make_note(db, c1, A.col1, "2026-05-01T08:00:00Z", status=A.st_noanswer, reach=A.reach_no)
    n2 = make_note(db, c1, A.col3, "2026-05-02T08:00:00Z", status=A.st_promise)
    row = by_key(rc.rows("RPT-02", A.sup, params)["items"], "case_id")[c1]
    assert row["last_note_at"] == "2026-05-02T08:00:00Z" and row["last_note_status"] == "RA وعد بالسداد"
    retire(db, n2)
    row = by_key(rc.rows("RPT-02", A.sup, params)["items"], "case_id")[c1]
    assert row["last_note_at"] == "2026-05-01T08:00:00Z" and row["last_note_status"] == "RA لا يرد"
    assert row["last_note_reachability"] == "RA لم يتم التوصل"
    assert keys(rc.rows("RPT-02", A.sup, {**params, "has_note": "no"}), "case_id") == {c2, un}
    assert keys(rc.rows("RPT-02", A.sup, {**params, "note_status": A.st_noanswer}), "case_id") == {c1}
    retire(db, n1)
    assert by_key(rc.rows("RPT-02", A.sup, params)["items"], "case_id")[c1]["last_note_at"] is None

    # assignment: the latest assignment counts only when it went to the current owner
    assign(db, c1, A.col1, "2026-04-01T09:00:00Z")
    d = rc.rows("RPT-02", A.sup, params)
    row = by_key(d["items"], "case_id")[c1]
    snap = dt.datetime.fromisoformat(d["snapshot_at"].replace("Z", "+00:00"))
    at = dt.datetime(2026, 4, 1, 9, tzinfo=dt.timezone.utc)
    assert row["assigned_at"] == "2026-04-01T09:00:00Z"
    assert row["assignment_days"] == math.floor((snap - at).total_seconds() / 86400)
    q(db, 'UPDATE clints SET user_id = %s WHERE id = %s', (A.col2, c1))    # moved without an assignment row
    row = by_key(rc.rows("RPT-02", A.sup, params)["items"], "case_id")[c1]
    assert row["assigned_at"] is None and row["assignment_days"] is None       # unknown, never 0
    assert by_key(rc.rows("RPT-02", A.sup, params)["items"], "case_id")[un]["assignment_days"] is None

    # exact money: approved excludes pending, rejected and reversed
    make_payment(db, c2, "100.10", "approved")
    make_payment(db, c2, "0.20", "approved")
    make_payment(db, c2, "50.05", None)
    make_payment(db, c2, "25.25", "rejected")
    make_payment(db, c2, "300.30", "reversed")
    d = rc.rows("RPT-02", A.sup, params)
    row = by_key(d["items"], "case_id")[c2]
    assert (row["approved_amount"], row["pending_amount"], row["outstanding_amount"], row["payment_count"]) == \
        ("100.30", "50.05", "400.25", 5)
    t = d["totals"]
    assert (t["original_debt"], t["approved_amount"], t["pending_amount"], t["outstanding_amount"]) == \
        ("1510.55", "100.30", "50.05", "1410.25")
    assert t["with_payments_count"] == 1
    assert d["items"][0]["case_id"] == un          # default sort: newest case first


def test_rpt02_filters_debt_account_and_attributes(rc, db):
    pa = make_portfolio(db, A.contract_a)
    a = make_case(db, pa, owner=A.col1, amount="10.00", account="RA-%_\\-" + uuid.uuid4().hex[:6],
                  attribute={"البروفجن": 125.5, "المنتج": "بطاقة"})
    b = make_case(db, pa, owner=A.col1, amount="10.01", name="عميل أحمد")
    c = make_case(db, pa, owner=A.col1, amount="9.99")
    q(db, "UPDATE clints SET is_locked = true WHERE id = %s", (c,))
    p = {"portfolio": pa}
    assert keys(rc.rows("RPT-02", A.sup, {**p, "debt_from": "10.00", "debt_to": "10.01"}), "case_id") == {a, b}
    assert keys(rc.rows("RPT-02", A.sup, {**p, "account": "%_\\"}), "case_id") == {a}
    assert keys(rc.rows("RPT-02", A.sup, {**p, "q": "أحمد"}), "case_id") == {b}
    d = rc.rows("RPT-02", A.sup, {**p, "locked": "yes"})
    assert keys(d, "case_id") == {c} and d["totals"]["locked_count"] == 1
    row = by_key(rc.rows("RPT-02", A.sup, p)["items"], "case_id")[a]
    assert row["provision"] == "125.5" and row["product"] == "بطاقة"
    assert by_key(rc.rows("RPT-02", A.sup, p)["items"], "case_id")[b]["provision"] is None   # unknown, not 0
    assert [i["case_id"] for i in rc.rows("RPT-02", A.sup, {**p, "sort": "original_debt"})["items"]] == [c, a, b]


def test_rpt03_archive_actor_reason_date_restore_and_scope(rc, db):
    pa = make_portfolio(db, A.contract_a)
    pb = make_portfolio(db, A.contract_b)
    x = make_case(db, pa, owner=A.col1, amount="300.00")
    y = make_case(db, pb, owner=A.col3, amount="400.00")              # out of ra.sup's scope
    z = make_case(db, pa, owner=None, amount="50.50")
    live = make_case(db, pa, owner=A.col1, amount="1.00")
    archive(db, x, by=A.col1, reason=A.reason1, when="2026-06-10T09:00:00Z")
    archive(db, y, by=A.sup, when="2026-06-11T09:00:00Z")
    q(db, 'UPDATE clints SET is_archive = true WHERE id = %s', (z,))           # archived without an archive record
    make_payment(db, x, "100.00", "approved")
    make_payment(db, x, "20.00", None)
    params = {"portfolio": f"{pa},{pb}"}
    d = rc.rows("RPT-03", A.sup, params)
    assert keys(d, "case_id") == {x, z}                                   # scoped (legacy was not), live excluded
    rows = by_key(d["items"], "case_id")
    assert rows[x]["archived_by_name"] == "محصل ر-أ" and rows[x]["archive_reason"] == "RA منتهي"
    assert rows[x]["archived_at"] == "2026-06-10T09:00:00Z" and rows[x]["owner_name"] == "محصل ر-أ"
    assert rows[z]["archived_by_name"] is rows[z]["archived_at"] is rows[z]["archive_reason"] is None
    t = d["totals"]
    assert (t["count"], t["original_debt"], t["approved_amount"], t["pending_amount"], t["outstanding_amount"],
            t["unassigned_count"]) == (2, "350.50", "100.00", "20.00", "250.50", 1)
    assert keys(rc.rows("RPT-03", A.super, params), "case_id") == {x, y, z}
    # restore removes it; re-archiving replaces the attribution with the newest archive record
    restore(db, x)
    assert keys(rc.rows("RPT-03", A.sup, params), "case_id") == {z}
    assert x in keys(rc.rows("RPT-02", A.sup, {"portfolio": pa}), "case_id")
    archive(db, x, by=A.sup, reason=A.reason2, when="2026-07-01T10:00:00Z")
    row = by_key(rc.rows("RPT-03", A.sup, params)["items"], "case_id")[x]
    assert (row["archived_by_name"], row["archive_reason"], row["archived_at"]) == \
        ("RA Supervisor", "RA سداد كامل", "2026-07-01T10:00:00Z")
    assert live not in keys(rc.rows("RPT-03", A.super, params), "case_id")
    # filters: actor, reason, Riyadh archive days
    assert keys(rc.rows("RPT-03", A.super, {**params, "archived_by": A.sup}), "case_id") == {x, y}
    assert keys(rc.rows("RPT-03", A.super, {**params, "archive_reason": A.reason1}), "case_id") == {y}
    assert keys(rc.rows("RPT-03", A.super, {**params, "archive_reason": A.reason2}), "case_id") == {x}
    w = make_case(db, pa, owner=A.col1)
    archive(db, w, when="2026-06-30T21:00:00Z")                           # Riyadh 1 July 00:00
    v = make_case(db, pa, owner=A.col1)
    archive(db, v, when="2026-06-30T20:59:59Z")                           # Riyadh 30 June 23:59:59
    one = rc.rows("RPT-03", A.super, {**params, "archived_from": "2026-07-01", "archived_to": "2026-07-01"})
    assert keys(one, "case_id") == {w, x}
    assert keys(rc.rows("RPT-03", A.super, {**params, "archived_to": "2026-06-30"}), "case_id") == {v, y}
    assert [i["case_id"] for i in rc.rows("RPT-03", A.super, {**params, "sort": "-archived_at",
                                                               "archive_reason": A.reason1})["items"]][:2] == [w, v]
    # the archive-actor selector lists the actors on scoped archived cases only
    u2 = make_case(db, pa, owner=A.col1)
    archive(db, u2, by=A.col1)
    opts = rc.get("/v2/report-options/archived_by", A.sup, {"report": "RPT-03"}).get_json()["data"]
    sel = rc.get("/v2/report-options/archived_by", A.sup,
                 {"report": "RPT-03", "selected": f"{A.sup},{A.col1},{A.col3}"}).get_json()["data"]["selected"]
    assert {i["id"] for i in sel} == {A.sup, A.col1} and opts["entity"] == "user"


# =========================================================================== RPT-04 / RPT-05
def test_note_dates_are_riyadh_days(rc, db):
    pa = make_portfolio(db, A.contract_a)
    case = make_case(db, pa, owner=A.col1)
    stored = {
        "apr30_last": "2026-04-30T20:59:59Z",     # Riyadh 30 April 23:59:59 -> out
        "may01_first": "2026-04-30T21:00:00Z",    # Riyadh 1 May 00:00 -> in (legacy UTC days miss it)
        "may31_last": "2026-05-31T20:59:59Z",     # Riyadh 31 May 23:59:59 -> in
        "jun01_first": "2026-05-31T21:00:00Z",    # Riyadh 1 June 00:00 -> out (legacy UTC days include it)
    }
    nid = {k: make_note(db, case, A.col1, v) for k, v in stored.items()}
    for report, user in (("RPT-04", A.sup), ("RPT-05", A.col1)):
        d = rc.rows(report, user, {**MAY, "portfolio": pa})
        assert keys(d, "note_id") == {nid["may01_first"], nid["may31_last"]}, report
        hours = {i["note_id"]: i["note_hour"] for i in d["items"]}
        assert hours[nid["may01_first"]] == 0 and hours[nid["may31_last"]] == 23
        one = rc.rows(report, user, {"created_from": "2026-06-01", "created_to": "2026-06-01", "portfolio": pa})
        assert keys(one, "note_id") == {nid["jun01_first"]}
        one = rc.rows(report, user, {"created_from": "2026-04-30", "created_to": "2026-04-30", "portfolio": pa})
        assert keys(one, "note_id") == {nid["apr30_last"]}


def test_rpt04_scope_retired_filters_and_totals(rc, db):
    pa = make_portfolio(db, A.contract_a)
    pb = make_portfolio(db, A.contract_b)
    mine = make_case(db, pa, owner=A.col1, amount="777.77")
    theirs = make_case(db, pb, owner=A.col3)                           # outside ra.sup's case scope
    n_other = make_note(db, mine, A.col3, "2026-05-05T10:00:00Z", callback="2026-05-06T09:00:00Z")
    n_col1 = make_note(db, mine, A.col1, "2026-05-06T10:00:00Z", status=A.st_noanswer, reach=A.reach_no,
                       text="اتصال بدون رد")
    n_sup_out = make_note(db, theirs, A.sup, "2026-05-07T10:00:00Z")
    n_ret = make_note(db, mine, A.col1, "2026-05-08T10:00:00Z")
    retire(db, n_ret)
    params = {**MAY, "portfolio": f"{pa},{pb}"}
    d = rc.rows("RPT-04", A.sup, params)
    assert keys(d, "note_id") == {n_other, n_col1}                    # recorder-independent; out-of-scope case excluded
    assert d["totals"] == {"count": 2, "case_count": 1, "recorder_count": 2, "callback_count": 1}
    rows = by_key(d["items"], "note_id")
    r = rows[n_col1]
    assert (r["recorder_name"], r["owner_name"], r["status"], r["reachability"], r["contact_method"]) == \
        ("محصل ر-أ", "محصل ر-أ", "RA لا يرد", "RA لم يتم التوصل", "RA جوال")
    assert r["original_debt"] == "777.77" and r["case_archived"] is False and r["corrected"] is False
    assert rows[n_other]["recorder_user_name"] == "ra.col3" and rows[n_other]["callback_at"] == "2026-05-06T09:00:00Z"
    assert keys(rc.rows("RPT-04", A.super, params), "note_id") == {n_other, n_col1, n_sup_out}
    # filters
    assert keys(rc.rows("RPT-04", A.sup, {**params, "recorder": A.col1}), "note_id") == {n_col1}
    assert keys(rc.rows("RPT-04", A.sup, {**params, "recorder_branch": A.branch2}), "note_id") == {n_other}
    assert keys(rc.rows("RPT-04", A.sup, {**params, "status": A.st_noanswer}), "note_id") == {n_col1}
    assert keys(rc.rows("RPT-04", A.sup, {**params, "reachability": A.reach_yes}), "note_id") == {n_other}
    assert keys(rc.rows("RPT-04", A.sup, {**params, "callback": "yes"}), "note_id") == {n_other}
    assert keys(rc.rows("RPT-04", A.sup, {**params, "text": "بدون رد"}), "note_id") == {n_col1}
    archive(db, mine)
    assert keys(rc.rows("RPT-04", A.sup, {**params, "case_state": "archived"}), "note_id") == {n_other, n_col1}
    assert rc.rows("RPT-04", A.sup, {**params, "case_state": "unarchived"})["total"] == 0
    # a corrected note shows the effective text and the corrected flag, recorder and date unchanged
    q(db, """INSERT INTO note_revisions (comment_id, case_id, revision_no, kind, text)
             VALUES (%s, %s, 0, 'original', 'اتصال بدون رد')""", (n_col1, mine))
    q(db, """INSERT INTO note_revisions (comment_id, case_id, revision_no, kind, text, reason, corrected_by, corrected_at)
             VALUES (%s, %s, 1, 'correction', 'تم الاتصال', 'typo', %s, now())""", (n_col1, mine, A.sup))
    q(db, 'UPDATE "Comments" SET comment = %s WHERE id = %s', ("تم الاتصال", n_col1))
    r = by_key(rc.rows("RPT-04", A.sup, params)["items"], "note_id")[n_col1]
    assert r["corrected"] is True and r["note"] == "تم الاتصال" and r["created_at"] == "2026-05-06T10:00:00Z"
    # selectors: recorders within the window and the caller's scope
    opts = rc.get("/v2/report-options/recorder", A.sup,
                  {"report": "RPT-04", **MAY, "selected": f"{A.col1},{A.col3},{A.sup}"}).get_json()["data"]
    assert {i["id"] for i in opts["selected"]} == {A.col1, A.col3}           # ra.sup's note is out of scope
    assert opts["depends_on"]["created"] == {"from": "2026-05-01", "to": "2026-05-31"}


def test_rpt04_latest_matching_note_across_the_full_filtered_set(rc, rep_app, db):
    pa = make_portfolio(db, A.contract_a)
    cases = [make_case(db, pa, owner=A.col1) for _ in range(35)]
    tie_winner, newest, text_only = {}, {}, {}
    for i, case in enumerate(cases):
        t = f"2026-05-{10 + i % 5:02d}T08:00:00Z"
        a = make_note(db, case, A.col1, t, text="tie match")
        b = make_note(db, case, A.col1, t, text="tie match")
        c = make_note(db, case, A.col1, f"2026-05-{10 + i % 5:02d}T07:00:00Z", text="older special")
        n = make_note(db, case, A.col3, f"2026-05-{10 + i % 5:02d}T09:00:00Z", text="newer other")
        tie_winner[case], newest[case], text_only[case] = max(a, b), n, c
    base = {**MAY, "portfolio": pa, "note_selection": "latest_per_case"}
    everything = rc.rows("RPT-04", A.sup, {**MAY, "portfolio": pa})
    assert everything["total"] == 140
    d = rc.rows("RPT-04", A.sup, {**base, "page_size": 200})
    assert keys(d, "note_id") == set(newest.values())
    assert d["totals"] == {"count": 35, "case_count": 35, "recorder_count": 1, "callback_count": 0}
    # the filters act BEFORE the selection: latest matching note, not "latest note if it matches"
    d = rc.rows("RPT-04", A.sup, {**base, "recorder": A.col1, "page_size": 200})
    assert keys(d, "note_id") == set(tie_winner.values())                 # equal timestamps: highest id
    d = rc.rows("RPT-04", A.sup, {**base, "text": "special", "page_size": 200})
    assert keys(d, "note_id") == set(text_only.values())
    d = rc.rows("RPT-04", A.sup, {**base, "recorder": A.col3, "text": "special"})
    assert d["total"] == 0 and d["totals"]["case_count"] == 0 and d["items"] == []
    # the display sort never changes which note is selected; pages are stable and complete
    one = rc.rows("RPT-04", A.sup, {**base, "recorder": A.col1, "page_size": 200, "sort": "created_at"})
    paged, page = [], 1
    while True:
        p = rc.rows("RPT-04", A.sup, {**base, "recorder": A.col1, "page_size": 10, "page": page, "sort": "created_at"})
        assert p["total"] == 35 and p["pages"] == 4
        paged += [i["note_id"] for i in p["items"]]
        if page == p["pages"]:
            break
        page += 1
    assert paged == [i["note_id"] for i in one["items"]] and len(set(paged)) == 35
    desc = rc.rows("RPT-04", A.sup, {**base, "recorder": A.col1, "page_size": 200, "sort": "-account_number"})
    assert keys(desc, "note_id") == keys(one, "note_id")
    beyond = rc.rows("RPT-04", A.sup, {**base, "recorder": A.col1, "page_size": 10, "page": 9})
    assert beyond["items"] == [] and beyond["totals"]["count"] == 35
    # a later window that excludes the newer notes' days selects within the window only
    d = rc.rows("RPT-04", A.sup, {**base, "created_from": "2026-05-10", "created_to": "2026-05-10"})
    assert d["total"] == 7 and keys(d, "note_id") == {newest[c] for i, c in enumerate(cases) if i % 5 == 0}
    # the full export of the selection equals the full preview
    assert_export_equals_preview(rc, rep_app, "RPT-04", A.sup, {**base, "recorder": A.col1}, sort="created_at")


def test_rpt05_recorder_rule_regardless_of_current_owner(rc, rep_app, db):
    pa = make_portfolio(db, A.contract_a)
    pb = make_portfolio(db, A.contract_b)
    case = make_case(db, pa, owner=A.col1, amount="1234.56")
    foreign = make_case(db, pb, owner=A.col3)
    mine1 = make_note(db, case, A.col1, "2026-05-03T10:00:00Z")
    mine2 = make_note(db, case, A.col1, "2026-05-04T10:00:00Z")
    by_other = make_note(db, case, A.col3, "2026-05-05T10:00:00Z")
    sup_note = make_note(db, foreign, A.sup, "2026-05-06T10:00:00Z")
    retired = make_note(db, case, A.col1, "2026-05-07T10:00:00Z")
    retire(db, retired)
    p = {**MAY, "portfolio": f"{pa},{pb}"}
    d = rc.rows("RPT-05", A.col1, p)
    assert keys(d, "note_id") == {mine1, mine2}
    assert d["totals"] == {"count": 2, "case_count": 1, "callback_count": 0}          # no repeated-debt sum
    # the case moves to another collector: col1's recorded notes stay in RPT-05, the case leaves RPT-02
    assign(db, case, A.col3, "2026-05-10T10:00:00Z")
    assert keys(rc.rows("RPT-05", A.col1, p), "note_id") == {mine1, mine2}
    assert rc.rows("RPT-02", A.col1, {"portfolio": pa})["total"] == 0
    row = by_key(rc.rows("RPT-05", A.col1, p)["items"], "note_id")[mine1]
    assert row["owner_name"] == "RA Collector3" and row["recorder_name"] == "محصل ر-أ"
    # the worker (fresh scope) exports them too
    items, _, done = assert_export_equals_preview(rc, rep_app, "RPT-05", A.col1, p)
    assert {i["note_id"] for i in items} == {mine1, mine2} and done["row_count"] == 2
    # the supervisor's own note on a case outside its scope: in RPT-05, not in RPT-04
    assert keys(rc.rows("RPT-05", A.sup, p), "note_id") == {sup_note}
    assert sup_note not in keys(rc.rows("RPT-04", A.sup, p), "note_id")
    assert by_other not in keys(rc.rows("RPT-05", A.col1, p), "note_id")
    # a recorder override is not a parameter of RPT-05
    assert rc.rows("RPT-05", A.col1, {**p, "recorder": A.col3}, expect=422)["error"] == {
        "recorder": "unknown_parameter"}


# =========================================================================== paging, exports, XLSX
def test_rpt02_paging_is_stable_with_full_ties(rc, rep_app, db):
    pa = make_portfolio(db, A.contract_a)
    same = "2026-02-02T02:02:02Z"
    ids_ = []
    for _ in range(110):
        c = make_case(db, pa, owner=A.col1, amount="5.00")
        q(db, "UPDATE clints SET create_at = %s WHERE id = %s", (same, c))
        ids_.append(c)
    for sort in ("case_created_at", "-case_created_at", "original_debt", "-original_debt", "owner_name",
                 "-last_note_at"):
        one = rc.rows("RPT-02", A.sup, {"portfolio": pa, "sort": sort, "page_size": 200})
        pages = []
        for page in range(1, 7):
            d = rc.rows("RPT-02", A.sup, {"portfolio": pa, "sort": sort, "page_size": 20, "page": page})
            assert d["total"] == 110 and d["totals"]["original_debt"] == "550.00"
            pages += [i["case_id"] for i in d["items"]]
        assert pages == [i["case_id"] for i in one["items"]], sort
        assert len(set(pages)) == 110
        assert pages == sorted(ids_, reverse=sort.startswith("-")), sort                                    # ties ordered by the case id
    assert_export_equals_preview(rc, rep_app, "RPT-02", A.sup, {"portfolio": pa}, sort="owner_name")


def test_full_exports_equal_previews_rpt03_rpt05(rc, rep_app, db):
    pa = make_portfolio(db, A.contract_a)
    for i in range(23):
        c = make_case(db, pa, owner=A.col1 if i % 3 else None, amount=f"{i}.{i:02d}")
        make_note(db, c, A.col1, f"2026-05-{1 + i % 28:02d}T1{i % 10}:00:00Z", callback="2026-06-01T09:00:00Z")
        archive(db, c, by=A.sup if i % 2 else A.col1, when=f"2026-06-{1 + i % 28:02d}T09:00:00Z")
        make_payment(db, c, "1.00", "approved" if i % 2 else None)
    _, first, _ = assert_export_equals_preview(rc, rep_app, "RPT-03", A.sup, {"portfolio": pa}, sort="-archived_at")
    assert first["totals"]["count"] == 23
    _, first, _ = assert_export_equals_preview(rc, rep_app, "RPT-05", A.col1, {**MAY, "portfolio": pa})
    assert first["totals"]["count"] == 23


def test_formula_safety_in_notes_and_names(rc, rep_app, db):
    pa = make_portfolio(db, A.contract_a)
    bad = ["=HYPERLINK(\"http://x\",\"y\")", "+1+1", "-2+3", "@SUM(A1)", "\tTAB", "=cmd|' /C calc'!A0"]
    for i, text in enumerate(bad):
        c = make_case(db, pa, owner=A.col1, name=text)
        make_note(db, c, A.col1, f"2026-05-1{i}T10:00:00Z", text=text, contact="=1+1")
    exp = rc.export("RPT-05", A.col1, {**MAY, "portfolio": pa}, "csv", sort="created_at")
    assert run_export(rep_app, exp["id"]) == "succeeded"
    header, rows = read_csv(rc.download(exp["id"], A.col1).data)
    ni, ci, ki = header.index("note"), header.index("customer_name"), header.index("contact_number")
    assert [r[ni] for r in rows] == ["'" + b for b in bad]
    assert [r[ci] for r in rows] == ["'" + b for b in bad]
    assert {r[ki] for r in rows} == {"'=1+1"}
    x = rc.export("RPT-05", A.col1, {**MAY, "portfolio": pa}, "xlsx", sort="created_at")
    assert run_export(rep_app, x["id"]) == "succeeded"
    import zipfile
    xml = zipfile.ZipFile(artifact_path(db, x["id"])).read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in xml


def test_arabic_xlsx_rpt01_and_rpt03(rc, rep_app, db):
    from openpyxl import load_workbook
    br = make_branch(db, "فرع اختبار")
    u = make_user(db, br, first="موظف", last="نشط", target_amount="2500.75")
    make_user(db, br, first="موظف", last="موقوف", active=False)
    x = rc.export("RPT-01", A.super, {"branch": br}, "xlsx_ar", sort="-name")
    assert run_export(rep_app, x["id"]) == "succeeded"
    wb = load_workbook(artifact_path(db, x["id"]))
    assert wb.sheetnames == ["المستخدمين", "الملخص"]
    assert all(ws.sheet_view.rightToLeft is True for ws in wb.worksheets)
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    assert hdr[:5] == ["معرف المستخدم", "الاسم", "اسم المستخدم", "الصلاحية", "الحالة"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert [r[hdr.index("الحالة")] for r in rows] == ["نشط", "غير نشط"]          # "-name": نشط (ن) before موقوف (م)
    active = next(r for r in rows if r[0] == u)
    assert active[hdr.index("المستهدف نسبة")] == "لا" and active[hdr.index("المستهدف (مبلغ)")] == 2500.75
    summary = {r[0].value: r[1].value for r in wb.worksheets[1].iter_rows()}
    assert summary["عدد السجلات"] == 2 and summary["النشطون"] == 1 and summary["التقرير"].startswith("RPT-01")
    assert "الفرع" in summary

    pa = make_portfolio(db, A.contract_a)
    c = make_case(db, pa, owner=A.col1, amount="321.00", name="عميل مؤرشف")
    archive(db, c, by=A.col1, reason=A.reason2, when="2026-06-15T12:00:00Z")
    x = rc.export("RPT-03", A.sup, {"portfolio": pa}, "xlsx_ar")
    assert run_export(rep_app, x["id"]) == "succeeded"
    wb = load_workbook(artifact_path(db, x["id"]))
    assert wb.sheetnames == ["عملاء الأرشيف", "الملخص"]
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    assert row[hdr.index("قام بالنقل")] == "محصل ر-أ" and row[hdr.index("سبب النقل")] == "RA سداد كامل"
    assert row[hdr.index("تاريخ النقل")] == dt.datetime(2026, 6, 15, 15, 0)            # Riyadh local time
    assert row[hdr.index("مبلغ المديونية")] == 321 and row[hdr.index("مقفل")] == "لا"
    assert row[hdr.index("المحفظة")] and row[hdr.index("اسم العميل")] == "عميل مؤرشف"


# =========================================================================== legacy vs v2 (test DB)
@pytest.fixture(scope="module")
def legacy_app(rep_app):
    from flask import Flask
    from sqlalchemy.pool import NullPool

    from conftest import db_url
    from extensions import db, jwt

    app = Flask("cawp_reports_a_legacy")
    app.config.update(TESTING=True, SQLALCHEMY_DATABASE_URI=db_url(MAIN_DB, "psycopg2"),
                      SQLALCHEMY_ENGINE_OPTIONS={"poolclass": NullPool}, SQLALCHEMY_TRACK_MODIFICATIONS=False,
                      JWT_SECRET_KEY=os.environ["JWT_SECRET_KEY"])
    db.init_app(app)
    jwt.init_app(app)
    import api.clients_reports
    import api.reports
    app.register_blueprint(api.reports.bp)
    app.register_blueprint(api.clients_reports.bp)
    return app


def _legacy(app, path, user, params):
    from flask_jwt_extended import create_access_token
    with app.app_context():
        token = create_access_token(identity=user)
    r = app.test_client().get(path, headers={"Authorization": f"Bearer {token}"}, query_string=params)
    assert r.status_code == 200, r.get_data(as_text=True)[:500]
    return r.get_json()


def test_documented_differences_from_the_legacy_endpoints(rc, db, legacy_app):
    """The reconciliation categories of review/phase4-reports-a.md, reproduced on synthetic rows."""
    pa = make_portfolio(db, A.contract_a)
    pb = make_portfolio(db, A.contract_b)
    inside = make_case(db, pa, owner=A.col1, amount="100.10")
    outside = make_case(db, pb, owner=A.col3, amount="200.20")
    win = [("customeFilterDateValue[]", "2025-02-01"), ("customeFilterDateValue[]", "2025-02-28"),
           ("page", "1"), ("pageSize", "0")]
    same = make_note(db, inside, A.col1, "2025-02-10T10:00:00Z")
    riyadh_first = make_note(db, inside, A.col1, "2025-01-31T21:30:00Z")     # Riyadh 1 Feb, UTC 31 Jan
    utc_last = make_note(db, inside, A.col1, "2025-02-28T22:00:00Z")         # Riyadh 1 Mar, UTC 28 Feb
    foreign = make_note(db, outside, A.col1, "2025-02-11T10:00:00Z")         # recorded by col1 on col3's case
    v2p = {"created_from": "2025-02-01", "created_to": "2025-02-28", "portfolio": f"{pa},{pb}"}

    # RPT-04 vs /comments: legacy is unscoped and uses UTC days
    leg = {c["id"] for c in _legacy(legacy_app, "/comments", A.sup, win)["comments"]}
    v2 = keys(rc.rows("RPT-04", A.sup, v2p), "note_id")
    assert same in leg and same in v2
    assert riyadh_first not in leg and riyadh_first in v2                    # +riyadh_first_day
    assert utc_last in leg and utc_last not in v2                            # -utc_after_end
    assert foreign in leg and foreign not in v2                              # -case_scope
    # RPT-05 vs /comments/current_user: recorder = me on both sides; only the day bounds differ
    body = _legacy(legacy_app, "/comments/current_user", A.col1, win)
    leg5 = {c["id"] for c in body["comments"]}
    v25 = keys(rc.rows("RPT-05", A.col1, v2p), "note_id")
    assert {same, foreign} <= leg5 and {same, foreign} <= v25
    assert riyadh_first in v25 - leg5 and utc_last in leg5 - v25
    assert abs(body["totalAmount"] - (2 * 100.10 + 200.20)) < 1e-6         # legacy sums the case debt once per note

    # RPT-03 vs /clints/archive/: legacy is unscoped and drops archived cases without an archive record
    archive(db, outside, by=A.col3)
    bare = make_case(db, pa, owner=A.col1)
    q(db, 'UPDATE clints SET is_archive = true WHERE id = %s', (bare,))
    for pf, legacy_n, v2_ids in ((pb, 1, set()), (pa, 0, {bare})):
        body = _legacy(legacy_app, "/clints/archive/", A.sup, [("SelectedPortfolios_Qkey[]", pf), ("pageSize", "50")])
        assert body["countRow"] == legacy_n
        assert keys(rc.rows("RPT-03", A.sup, {"portfolio": pf}), "case_id") == v2_ids

    # RPT-02 vs /clints: the same scope and the same rows; legacy totalAmount is a float
    make_case(db, pa, owner=None, amount="0.20")
    body = _legacy(legacy_app, "/clints", A.sup, [("customeFilterPortfoliosValue[]", pa), ("pageSize", "50")])
    d = rc.rows("RPT-02", A.sup, {"portfolio": pa})
    assert body["countRow"] == d["total"] == 2
    assert {c["id"] for c in body["clints"]} == keys(d, "case_id")
    assert Decimal(str(body["totalAmount"])).quantize(Decimal("0.01")) == Decimal(d["totals"]["original_debt"])

    # RPT-01 vs /uesrs/export: legacy lists every user with counts over every case
    br = make_branch(db)
    stranger = make_user(db, br, contract=make_contract(db))
    make_case(db, make_portfolio(db, make_contract(db)), owner=stranger, amount="5.00")
    users = {u["id"]: u for u in _legacy(legacy_app, "/uesrs/export", A.sup, [])["users"]}
    assert users[stranger]["count_clints"] == 1 and users[stranger]["total_amount"] == 5.0
    assert users[stranger]["email"] is None                                  # legacy redaction (user_privacy)
    assert rc.rows("RPT-01", A.sup, {"branch": br})["total"] == 0            # -employee_scope
    assert rc.rows("RPT-01", A.super, {"branch": br})["items"][0]["case_count"] == 1
